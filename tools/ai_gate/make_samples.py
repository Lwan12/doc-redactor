"""生成出厂门禁用的编造样本：普通 PDF 里的各种花样、PDF 隐藏处、增量保存、带标签的 PDF，
Word 正文和隐藏处、图片替代文字、嵌入对象、不显眼处。全部是编造数据。

用法：python tools/ai_gate/make_samples.py <目录>
P4_scan 是整页图片的扫描件（要 OCR），找得到中文字体才生成；Word 文档属性归 #19，不在这里。
"""

from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

import pymupdf
from docx import Document

EN = "Sponsor: Acme Pharma; Drug: Fictizumab; Investigator: Johnathan Quill; Subject SUBJ-0042-117; Contract HT-2026-0815; Protocol ABC-2026-001"
ZH = "申办方星辰海制药，药物菲替珠单抗，研究者王小明，受试者SUBJ-0042-117，合同HT-2026-0815。"
PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde"
    b"\x00\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe\x02\xfe\xa7\x9a\xa0\xa1\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _font(text: str) -> str:
    return "china-s" if re.search(r"[\u4e00-\u9fff]", text) else "helv"


def _lines(page, y, items, size=9):
    for text in items:
        page.insert_text((50, y), text, fontname=_font(text), fontsize=size)
        y += 14
    return y


def _save(document, path: Path) -> None:
    data = document.tobytes()
    document.close()
    path.write_bytes(data)


def pdf_plain(path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    _lines(page, 60, [EN[:70], EN[70:], ZH[:24], ZH[24:],
                      "Case variants: ACME PHARMA / acme pharma / FICTIZUMAB",
                      "Line break inside a term: Acme", "Pharma continues here.",
                      "Hyphenated: Fictizu-", "mab continues."])
    page.insert_text((50, 300), "Invisible layer: Johnathan Quill SUBJ-0042-117", fontsize=9, render_mode=3)
    page.insert_text((50, 320), "White text: Acme Pharma HT-2026-0815", fontsize=9, color=(1, 1, 1))
    page.insert_text((50, 340), "Tiny text: Fictizumab ABC-2026-001", fontsize=1)
    page.insert_text((300, 500), "Rotated: Johnathan Quill", fontsize=9, rotate=90)
    page.insert_text((50, -20), "Off page: Acme Pharma SUBJ-0042-117", fontsize=9)
    page.insert_text((50, 800), "Below crop: Fictizumab HT-2026-0815", fontsize=9)
    source = pymupdf.open()
    form = source.new_page(width=300, height=60)
    form.insert_text((5, 30), "In form XObject: 星辰海制药 Fictizumab", fontname="china-s", fontsize=9)
    page.show_pdf_page(pymupdf.Rect(50, 400, 350, 460), source, 0)
    page.set_cropbox(pymupdf.Rect(0, 0, 595, 760))
    _save(document, path)


def pdf_hidden(path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    _lines(page, 60, [EN[:70], "See the attachment and the form below."])
    document.set_metadata({"title": "Acme Pharma Fictizumab CSR", "author": "Johnathan Quill", "subject": "SUBJ-0042-117",
                           "keywords": "HT-2026-0815, ABC-2026-001, 星辰海制药", "creator": "王小明", "producer": "菲替珠单抗 tool"})
    info = int(document.xref_get_key(-1, "Info")[1].split()[0])
    document.xref_set_key(info, "Sponsor", pymupdf.get_pdf_str("Acme Pharma 星辰海制药"))
    document.xref_set_key(info, "ContractNo", "(HT-2026-0815)")
    document.set_toc([[1, "1 Acme Pharma overview", 1], [2, "1.1 Fictizumab 菲替珠单抗", 1]])
    page.add_text_annot((400, 100), "Reviewer note: Johnathan Quill checked SUBJ-0042-117")
    page.add_freetext_annot(pymupdf.Rect(300, 200, 550, 240), "FreeText: Acme Pharma HT-2026-0815", fontsize=9)
    page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(50, 300, 200, 315),
                      "uri": "https://example.org/trial/ABC-2026-001?subj=SUBJ-0042-117"})
    widget = pymupdf.Widget()
    widget.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    widget.field_name = "investigator"
    widget.field_value = "王小明 SUBJ-0042-117"
    widget.rect = pymupdf.Rect(50, 400, 300, 420)
    page.add_widget(widget)
    document.embfile_add("subjects.txt", "SUBJ-0042-117 Johnathan Quill Acme Pharma 王小明\n".encode(), desc="Fictizumab listing")
    document.set_page_labels([{"startpage": 0, "prefix": "Acme Pharma-", "style": "D", "firstpagenum": 1}])
    document.add_ocg("Fictizumab layer 菲替珠单抗", on=True)
    _save(document, path)


def pdf_incremental(path: Path) -> None:
    document = pymupdf.open()
    _lines(document.new_page(), 60, [EN[:70], EN[70:], ZH[:24], ZH[24:]])
    document.save(path)
    document.close()
    document = pymupdf.open(path)
    page = document[0]
    for term in ("Acme Pharma", "Fictizumab", "Johnathan Quill"):
        for rect in page.search_for(term):
            page.add_redact_annot(rect, fill=(1, 1, 1))
    page.apply_redactions()
    page.insert_text((50, 200), "v2 kept: Subject SUBJ-0042-117.", fontsize=9)
    document.saveIncr()
    document.close()


def pdf_tagged(path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    _lines(page, 60, ["Figure 1 below.", EN[:40]])
    tmp = document.tobytes()
    document.close()
    document = pymupdf.open("pdf", tmp)
    page = document[0]
    extra = document.get_new_xref()
    document.update_object(extra, "<<>>")
    actual = pymupdf.get_pdf_str("Acme Pharma 星辰海制药 SUBJ-0042-117")
    document.update_stream(
        extra,
        ("q /Span <</ActualText " + actual + ">> BDC BT /helv 9 Tf 50 520 Td (Shown as: name withheld) Tj ET EMC Q\n"
         "q /Span <</Alt (Fictizumab HT-2026-0815) /E (Johnathan Quill)>> BDC BT /helv 9 Tf 50 500 Td (abbr J.Q.) Tj ET EMC Q\n").encode("latin1"),
    )
    contents = page.get_contents()
    document.xref_set_key(page.xref, "Contents", "[" + " ".join(f"{x} 0 R" for x in contents + [extra]) + "]")
    root = document.get_new_xref()
    figure = document.get_new_xref()
    alt = pymupdf.get_pdf_str("Chart of Fictizumab for Johnathan Quill, 王小明")
    document.update_object(figure, f"<</Type/StructElem/S/Figure/P {root} 0 R/Alt{alt}/ActualText(ABC-2026-001)>>")
    document.update_object(root, f"<</Type/StructTreeRoot/K[{figure} 0 R]>>")
    catalog = document.pdf_catalog()
    document.xref_set_key(catalog, "StructTreeRoot", f"{root} 0 R")
    document.xref_set_key(catalog, "MarkInfo", "<</Marked true>>")
    _save(document, path)


# --- Word ---------------------------------------------------------------------

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
NSALL = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)
RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"


def R(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _patch(path: Path, edits=None, add=None, ct="", rels=None) -> None:
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for name, fn in (edits or {}).items():
        parts[name] = fn(parts[name].decode("utf-8")).encode("utf-8")
    for name, data in (add or {}).items():
        parts[name] = data if isinstance(data, bytes) else data.encode("utf-8")
    if ct:
        types = parts["[Content_Types].xml"].decode("utf-8")
        for item in re.findall(r"<Default [^>]*/>|<Override [^>]*/>", ct):
            ext = re.search(r'Extension="([^"]+)"', item)
            if ext and f'Extension="{ext.group(1)}"' in types:
                continue
            types = types.replace("</Types>", item + "</Types>")
        parts["[Content_Types].xml"] = types.encode("utf-8")
    for rel_part, xml in (rels or {}).items():
        text = parts[rel_part].decode("utf-8")
        parts[rel_part] = text.replace("</Relationships>", xml + "</Relationships>").encode("utf-8")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    path.write_bytes(buffer.getvalue())


def _body(snippet: str):
    def edit(xml: str) -> str:
        xml = re.sub(r"<w:document [^>]*>", f'<w:document {NSALL} mc:Ignorable="wps">', xml, count=1)
        return xml.replace("<w:sectPr", snippet + "<w:sectPr", 1)

    return edit


def _notes(kind: str, text: str) -> str:
    tag = kind[:-1]
    return (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:{kind} {W}>'
        f'<w:{tag} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{tag}>'
        f'<w:{tag} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:{tag}>'
        f'<w:{tag} w:id="1"><w:p>{R(text)}</w:p></w:{tag}></w:{kind}>'
    )


def docx_body(path: Path) -> None:
    document = Document()
    document.add_paragraph(EN)
    document.add_paragraph(ZH)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cell: Acme Pharma SUBJ-0042-117"
    table.cell(0, 1).text = "单元格：菲替珠单抗 王小明"
    document.sections[0].header.paragraphs[0].text = "Header: Acme Pharma Protocol ABC-2026-001"
    document.sections[0].footer.paragraphs[0].text = "Footer: Contract HT-2026-0815 星辰海制药"
    paragraph = document.add_paragraph("Commented text here.")
    document.add_comment(paragraph.runs, text="Comment: Johnathan Quill verify SUBJ-0042-117 for Acme Pharma",
                         author="Johnathan Quill", initials="JQ")
    document.save(str(path))
    rev = 'w:author="王小明" w:date="2026-01-01T00:00:00Z"'
    snippet = (
        f'<w:p>{R("Footnote ref")}<w:r><w:footnoteReference w:id="1"/></w:r>{R(" endnote ref")}<w:r><w:endnoteReference w:id="1"/></w:r></w:p>'
        f'<w:p>{R("Tracked: ")}<w:ins w:id="501" {rev}>{R("inserted Fictizumab HT-2026-0815")}</w:ins>'
        f'<w:del w:id="502" {rev}><w:r><w:delText xml:space="preserve"> deleted Johnathan Quill 星辰海制药</w:delText></w:r></w:del></w:p>'
        f'<w:p>{R("Visible. ")}<w:r><w:rPr><w:vanish/></w:rPr><w:t xml:space="preserve">Hidden: Fictizumab SUBJ-0042-117 王小明</w:t></w:r></w:p>'
        f'<w:p>{R("Box: ")}<w:r><w:pict><v:rect style="width:200pt;height:30pt"><v:textbox><w:txbxContent><w:p>{R("Textbox: Acme Pharma Johnathan Quill")}</w:p></w:txbxContent></v:textbox></v:rect></w:pict></w:r></w:p>'
        f'<w:p>{R("Link field: ")}<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> HYPERLINK "https://example.org/subject/SUBJ-0042-117" </w:instrText></w:r>'
        f'<w:r><w:fldChar w:fldCharType="separate"/></w:r>{R("trial portal")}<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        f'{R(" quote: ")}<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> QUOTE "Johnathan Quill 菲替珠单抗" </w:instrText></w:r>'
        f'<w:r><w:fldChar w:fldCharType="separate"/></w:r>{R("[see sponsor]")}<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        f'{R(" simple: ")}<w:fldSimple w:instr=" DOCPROPERTY Company  HT-2026-0815 ">{R("company")}</w:fldSimple></w:p>'
    )
    _patch(
        path,
        {"word/document.xml": _body(snippet)},
        add={"word/footnotes.xml": _notes("footnotes", "Footnote: Acme Pharma funded; Johnathan Quill, SUBJ-0042-117"),
             "word/endnotes.xml": _notes("endnotes", "Endnote: 星辰海制药 菲替珠单抗 ABC-2026-001")},
        ct='<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
           '<Override PartName="/word/endnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"/>',
        rels={"word/_rels/document.xml.rels": f'<Relationship Id="rIdFn1" Type="{RT}footnotes" Target="footnotes.xml"/>'
                                              f'<Relationship Id="rIdEn1" Type="{RT}endnotes" Target="endnotes.xml"/>'},
    )


def docx_alt(path: Path) -> None:
    document = Document()
    document.add_paragraph("Figure below. Protocol ABC-2026-001.")
    document.add_picture(io.BytesIO(PNG))
    document.save(str(path))

    def edit(xml: str) -> str:
        xml = re.sub(r"<wp:docPr [^>]*?/>", '<wp:docPr id="5" name="Acme Pharma logo" descr="Alt: Acme Pharma Fictizumab SUBJ-0042-117" title="王小明 星辰海制药 HT-2026-0815"/>', xml)
        return re.sub(r"<pic:cNvPr [^>]*?/>", '<pic:cNvPr id="0" name="Fictizumab.png" descr="Johnathan Quill chart"/>', xml)

    _patch(path, {"word/document.xml": edit})


def docx_embed(path: Path) -> None:
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as book:
        book.writestr("[Content_Types].xml", "<Types/>")
        book.writestr("xl/sharedStrings.xml", "<sst><si><t>SUBJ-0042-117 Fictizumab Johnathan Quill 星辰海制药</t></si></sst>")
    document = Document()
    document.add_paragraph("Embedded workbook below. Protocol ABC-2026-001.")
    document.save(str(path))
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504 + b"Ole10Native Johnathan Quill Acme Pharma " + "王小明 菲替珠单抗".encode("utf-16-le")
    obj = (
        '<w:p><w:r><w:object w:dxaOrig="4000" w:dyaOrig="1200"><v:shape id="_x0000_i10{n}" type="#_x0000_t75" style="width:200pt;height:60pt" o:ole="">'
        '<v:imagedata r:id="rIdPv1" o:title="Acme Pharma sheet"/></v:shape>'
        '<o:OLEObject Type="Embed" ProgID="{prog}" ShapeID="_x0000_i10{n}" DrawAspect="Content" ObjectID="_170000000{n}" r:id="{rid}"/></w:object></w:r></w:p>'
    )
    _patch(
        path,
        {"word/document.xml": _body(obj.format(n=1, prog="Excel.Sheet.12", rid="rIdEm1") + obj.format(n=2, prog="Package", rid="rIdEm2"))},
        add={"word/embeddings/Microsoft_Excel_Worksheet1.xlsx": inner.getvalue(), "word/embeddings/oleObject1.bin": ole, "word/media/preview1.png": PNG},
        ct='<Default Extension="xlsx" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"/>'
           '<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.oleObject"/>'
           '<Default Extension="png" ContentType="image/png"/>',
        rels={"word/_rels/document.xml.rels": f'<Relationship Id="rIdEm1" Type="{RT}package" Target="embeddings/Microsoft_Excel_Worksheet1.xlsx"/>'
                                              f'<Relationship Id="rIdEm2" Type="{RT}oleObject" Target="embeddings/oleObject1.bin"/>'
                                              f'<Relationship Id="rIdPv1" Type="{RT}image" Target="media/preview1.png"/>'},
    )


def docx_obscure(path: Path) -> None:
    document = Document()
    style = document.styles.add_style("Fictizumab 正文 Acme Pharma", 1)
    document.add_paragraph("Styled paragraph.", style=style)
    reviewed = document.add_paragraph("Reviewed.")
    document.add_comment(reviewed.runs, text="ok", author="王小明 (Acme Pharma)", initials="WXM")
    document.save(str(path))
    snippet = (
        '<w:p><w:bookmarkStart w:id="9" w:name="Acme_Pharma_SUBJ_0042_117"/>' + R("Portal: ")
        + '<w:hyperlink r:id="rIdHl1"><w:r><w:t>click here</w:t></w:r></w:hyperlink><w:bookmarkEnd w:id="9"/>'
        '<w:hyperlink w:anchor="Acme_Pharma_SUBJ_0042_117" w:tooltip="Tooltip: Johnathan Quill HT-2026-0815"><w:r><w:t xml:space="preserve"> internal link</w:t></w:r></w:hyperlink></w:p>'
        f'<w:p><w:ins w:id="601" w:author="Johnathan Quill" w:date="2026-01-01T00:00:00Z">{R("plain inserted words")}</w:ins></w:p>'
        '<w:sdt><w:sdtPr><w:alias w:val="Sponsor: Acme Pharma"/><w:tag w:val="SUBJ-0042-117"/><w:id w:val="42"/></w:sdtPr>'
        f'<w:sdtContent><w:p>{R("Click to enter")}</w:p></w:sdtContent></w:sdt>'
        '<w:altChunk r:id="rIdAc1"/>'
    )

    def settings(xml: str) -> str:
        xml = xml.replace("</w:settings>", '<w:attachedTemplate r:id="rIdTpl1"/><w:docVars><w:docVar w:name="Sponsor" w:val="Acme Pharma 星辰海制药"/>'
                                           '<w:docVar w:name="SubjectId" w:val="SUBJ-0042-117"/></w:docVars></w:settings>')
        if "xmlns:r=" not in xml.split(">", 2)[1]:
            xml = xml.replace("<w:settings ", '<w:settings xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
        return xml

    _patch(
        path,
        {"word/document.xml": _body(snippet), "word/settings.xml": settings},
        add={"word/afchunk.htm": "<html><body><p>altChunk: Acme Pharma Fictizumab Johnathan Quill SUBJ-0042-117</p></body></html>",
             "word/_rels/settings.xml.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                                             f'<Relationship Id="rIdTpl1" Type="{RT}attachedTemplate" Target="file:///C:/Users/JohnathanQuill/Templates/Acme%20Pharma_HT-2026-0815.dotm" TargetMode="External"/></Relationships>'},
        ct='<Default Extension="htm" ContentType="application/xhtml+xml"/>',
        rels={"word/_rels/document.xml.rels": f'<Relationship Id="rIdHl1" Type="{RT}hyperlink" Target="https://example.org/trials/SUBJ-0042-117?pi=Johnathan%20Quill" TargetMode="External"/>'
                                              f'<Relationship Id="rIdAc1" Type="{RT}aFChunk" Target="afchunk.htm"/>'},
    )


SAMPLES = {
    "P1_plain.pdf": pdf_plain,
    "P2_hidden.pdf": pdf_hidden,
    "P3_incremental.pdf": pdf_incremental,
    "P5_tagged.pdf": pdf_tagged,
    "W1_body_hidden.docx": docx_body,
    "W2_alt.docx": docx_alt,
    "W4_embed.docx": docx_embed,
    "W6_obscure.docx": docx_obscure,
}


CJK_FONTS = (
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
    "C:/Windows/Fonts/msyh.ttc",
    "C:/Windows/Fonts/simsun.ttc",
)


def cjk_font() -> str | None:
    return next((path for path in CJK_FONTS if Path(path).is_file()), None)


def pdf_scan(path: Path, font_path: str) -> None:
    """整页图片的扫描件：白底黑字 34 像素，tesseract 会把药名认错一个字。"""
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1654, 600), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(font_path, 34)
    rows = [
        "Sponsor: Acme Pharma   Drug: Fictizumab",
        "Investigator: Johnathan Quill",
        "Subject SUBJ-0042-117  Contract HT-2026-0815",
        "Protocol ABC-2026-001",
        "申办方星辰海制药  药物菲替珠单抗  研究者王小明",
    ]
    for row, text in enumerate(rows):
        draw.text((60, 40 + row * 60), text, font=font, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(pymupdf.Rect(20, 40, 575, 242), stream=buffer.getvalue())
    _save(document, path)


def make_all(folder: Path, scan: bool = True) -> list[Path]:
    """``scan=True`` 时找得到中文字体就多生成一份扫描件 P4_scan.pdf。"""
    folder.mkdir(parents=True, exist_ok=True)
    made = []
    for name, build in SAMPLES.items():
        target = folder / name
        build(target)
        made.append(target)
    font = cjk_font() if scan else None
    if font is not None:
        target = folder / "P4_scan.pdf"
        pdf_scan(target, font)
        made.append(target)
    return made


if __name__ == "__main__":
    for path in make_all(Path(sys.argv[1])):
        print(path.name)
