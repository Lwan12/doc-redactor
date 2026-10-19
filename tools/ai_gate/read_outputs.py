"""出厂门禁的读法：只用 AI 程序最常用的几种解析器，读出原词就不放行。

口径照审阅的「AI 读取复核」流程：
- PDF 三种：PyMuPDF get_text、pymupdf4llm.to_markdown（遇到扫描图会用 RapidOCR 看图）、
  pdftotext（poppler）。
- Word 两种：python-docx（正文、表格、页眉页脚、core 属性）和 pandoc（-t markdown）。

匹配：NFKC、不分大小写、还原 URL 编码和 HTML 实体、去掉行内标签，再去掉空白、
连字符、下划线、星号和反斜杠后找整词（断行或下划线拼回来的也算）。

用法：python tools/ai_gate/read_outputs.py <输出目录> [--words words.txt] [--strict]
任何一个文件、任何一种读法读出原词，或解析器出错，退出码 1。
--strict（或环境变量 AI_GATE_STRICT=1）时，缺任何一种读法也算失败（CI 门禁用）。
只打印文件名、读法和原词序号，不打印原词。
"""

from __future__ import annotations

import argparse
import html
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from pathlib import Path
from urllib.parse import unquote

HERE = Path(__file__).resolve().parent
PDF_READERS = ("PyMuPDF", "pymupdf4llm", "pdftotext")
DOCX_READERS = ("python-docx", "pandoc")
FAILED = 0  # 解析失败时的「原词序号」


def strict_mode() -> bool:
    return os.environ.get("AI_GATE_STRICT", "") not in ("", "0")


def load_words(path: Path = HERE / "words.txt") -> list[str]:
    lines = path.read_text(encoding="utf-8").splitlines()
    return [line.strip() for line in lines if line.strip() and not line.startswith("#")]


def _flat(text: str) -> str:
    text = re.sub(r"(?i)<br\s*/?>", "\n", text or "")
    text = re.sub(r"</?[A-Za-z][^>\n]{0,80}>", "", text)
    text = html.unescape(unquote(text))
    text = unicodedata.normalize("NFKC", text).casefold()
    return re.sub(r"[\s\-_*\\\u00ad\u2010]+", "", text)


def _poppler_pdftotext() -> str | None:
    exe = shutil.which("pdftotext")
    if exe is None:
        return None
    try:
        probe = subprocess.run([exe, "-v"], capture_output=True, timeout=30)
    except (OSError, subprocess.TimeoutExpired):
        return None
    # Windows 上的 xpdf 版没有中文映射，只认 poppler。
    return exe if b"oppler" in probe.stdout + probe.stderr else None


def _has_pymupdf4llm() -> bool:
    try:
        import pymupdf4llm  # noqa: F401
    except Exception:
        return False
    return True


def _pandoc() -> str | None:
    try:
        import pypandoc

        return pypandoc.get_pandoc_path()
    except Exception:
        return shutil.which("pandoc")


def available_readers() -> list[str]:
    names = ["PyMuPDF"]
    if _has_pymupdf4llm():
        names.append("pymupdf4llm")
    if _poppler_pdftotext():
        names.append("pdftotext")
    names.append("python-docx")
    if _pandoc():
        names.append("pandoc")
    return names


def missing_readers() -> list[str]:
    have = set(available_readers())
    return [name for name in PDF_READERS + DOCX_READERS if name not in have]


def _pymupdf(path: Path) -> str:
    import pymupdf

    with pymupdf.open(path) as document:
        return "\n".join(page.get_text() for page in document)


def _pymupdf4llm(path: Path) -> str:
    import contextlib
    import io

    import pymupdf4llm

    # 它会往标准输出打 OCR 进度，收起来，免得混进门禁的结果行。
    with contextlib.redirect_stdout(io.StringIO()):
        return pymupdf4llm.to_markdown(str(path))


def _pdftotext(path: Path) -> str:
    result = subprocess.run(
        [_poppler_pdftotext(), "-enc", "UTF-8", str(path), "-"], capture_output=True, timeout=120, check=True
    )
    return result.stdout.decode("utf-8", "replace")


def _python_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        for part in (section.header, section.footer):
            parts.extend(p.text for p in part.paragraphs)
    core = document.core_properties
    for name in ("title", "subject", "author", "keywords", "comments", "last_modified_by", "category"):
        parts.append(getattr(core, name) or "")
    return "\n".join(parts)


def _pandoc_md(path: Path) -> str:
    result = subprocess.run(
        [_pandoc(), "-f", "docx", "-t", "markdown", str(path)], capture_output=True, timeout=120, check=True
    )
    return result.stdout.decode("utf-8", "replace")


_READ = {
    "PyMuPDF": _pymupdf,
    "pymupdf4llm": _pymupdf4llm,
    "pdftotext": _pdftotext,
    "python-docx": _python_docx,
    "pandoc": _pandoc_md,
}


def readers_for(path: Path) -> tuple[str, ...]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        names = PDF_READERS
    elif suffix == ".docx":
        names = DOCX_READERS
    else:
        return ()
    have = set(available_readers())
    return tuple(name for name in names if name in have)


def findings(path: Path, words: list[str]) -> list[tuple[str, int]]:
    """(读法, 原词序号) 列表；序号从 1 开始，对应 words.txt 的顺序；0 表示这种读法出错。"""
    needles = [(number, _flat(word)) for number, word in enumerate(words, start=1)]
    found = []
    for method in readers_for(path):
        try:
            flat = _flat(_READ[method](path))
        except Exception:
            found.append((method, FAILED))
            continue
        for number, needle in needles:
            if needle and needle in flat:
                found.append((method, number))
    return found


def describe(path: Path, method: str, number: int) -> str:
    if number == FAILED:
        return f"{path.name}: {method} 解析失败"
    return f"{path.name}: {method} 读出原词 #{number}"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("outdir")
    parser.add_argument("--words", default=str(HERE / "words.txt"))
    parser.add_argument("--strict", action="store_true")
    args = parser.parse_args(argv)
    words = load_words(Path(args.words))
    print("读法：" + "、".join(available_readers()))
    missing = missing_readers()
    if missing and (args.strict or strict_mode()):
        print("缺少读法：" + "、".join(missing))
        return 1
    bad = 0
    for path in sorted(Path(args.outdir).iterdir()):
        for method, number in findings(path, words):
            print(describe(path, method, number))
            bad += 1
    print("通过" if not bad else f"不放行：{bad} 处")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
