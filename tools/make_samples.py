"""Write fictional sample.docx, sample.pdf, and sample.png.

Nothing in these files is a real protocol number, sponsor, or drug.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"

PROJECT_ID = "ABC-2026-001"
SPONSOR = "Acme Pharma"
DRUG = "Fictizumab"

_FONT_CANDIDATES = (
    Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
    Path("C:/Windows/Fonts/arial.ttf"),
)


def font_path() -> Path:
    for candidate in _FONT_CANDIDATES:
        if candidate.is_file():
            return candidate
    raise FileNotFoundError("没有找到可用于样例图片的 TrueType 字体")


def render_png(path: Path) -> None:
    font = ImageFont.truetype(str(font_path()), 72)
    image = Image.new("RGB", (1800, 640), "white")
    draw = ImageDraw.Draw(image)
    lines = [
        f"Project {PROJECT_ID}",
        f"Sponsor {SPONSOR}",
        f"Drug {DRUG}",
    ]
    y = 48
    for line in lines:
        draw.text((48, y), line, fill="black", font=font)
        y += 180
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG", dpi=(300, 300))


_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)
_REV = 'w:author="Sample" w:date="2026-01-01T00:00:00Z"'


def _run_xml(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _el(fragment: str):
    head, _sep, rest = fragment.partition(">")
    return parse_xml(f"{head} {_NS}>{rest}")


def _add_hidden_containers(document) -> None:
    """Text box, simple field, content control, and tracked changes (all fictional)."""
    box = f"<w:txbxContent><w:p>{_run_xml('BOX ' + SPONSOR)}</w:p></w:txbxContent>"
    header = document.sections[0].header.paragraphs[0]._p
    header.append(
        _el(
            "<w:r><mc:AlternateContent>"
            '<mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic>'
            '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            f"<wps:wsp><wps:txbx>{box}</wps:txbx></wps:wsp>"
            "</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
            f"<mc:Fallback><w:pict><v:shape><v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback>"
            "</mc:AlternateContent></w:r>"
        )
    )
    p = document.add_paragraph("Field ")._p
    p.append(_el(f'<w:fldSimple w:instr=" DOCPROPERTY Company ">{_run_xml(SPONSOR)}</w:fldSimple>'))
    p = document.add_paragraph("Control ")._p
    p.append(_el(f"<w:sdt><w:sdtPr/><w:sdtContent>{_run_xml(DRUG)}</w:sdtContent></w:sdt>"))
    p = document.add_paragraph("Revised ")._p
    p.append(_el(f'<w:ins w:id="901" {_REV}>{_run_xml(PROJECT_ID)}</w:ins>'))
    p.append(
        _el(
            f'<w:del w:id="902" {_REV}><w:r><w:delText xml:space="preserve"> {DRUG}</w:delText></w:r></w:del>'
        )
    )


def render_docx(path: Path, png: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = True

    section.first_page_header.paragraphs[0].text = f"FIRST {PROJECT_ID}"
    section.header.paragraphs[0].text = f"ODD {SPONSOR}"
    section.even_page_header.paragraphs[0].text = f"EVEN {DRUG}"
    section.first_page_footer.paragraphs[0].text = f"FF {DRUG}"
    section.footer.paragraphs[0].text = f"DF {PROJECT_ID}"
    section.even_page_footer.paragraphs[0].text = f"EF {SPONSOR}"

    document.add_paragraph(f"Project {PROJECT_ID}.")
    document.add_paragraph("Keep fictizumab in lowercase.")
    document.add_paragraph(f"Seen {SPONSOR.lower()} today.")

    split = document.add_paragraph()
    split.add_run("ID ")
    bold = split.add_run("ABC-20")
    bold.bold = True
    italic = split.add_run("26")
    italic.italic = True
    underlined = split.add_run("-001")
    underlined.underline = True
    split.add_run(" end")

    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = f"Site {SPONSOR}"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = f"Drug {DRUG}"

    _add_hidden_containers(document)

    document.add_picture(str(png))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def render_pdf(path: Path, png: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.setFont("Helvetica", 14)
    y = 740
    for line in (
        f"Project {PROJECT_ID}",
        f"Sponsor {SPONSOR}",
        f"Drug {DRUG}",
        "fictizumab stays",
    ):
        pdf.drawString(72, y, line)
        y -= 22
    pdf.drawImage(str(png), 72, 360, width=450, height=160, preserveAspectRatio=True, mask="auto")
    pdf.save()


def main() -> None:
    png = SAMPLES / "sample.png"
    render_png(png)
    render_docx(SAMPLES / "sample.docx", png)
    render_pdf(SAMPLES / "sample.pdf", png)
    print(f"wrote {png}")
    print(f"wrote {SAMPLES / 'sample.docx'}")
    print(f"wrote {SAMPLES / 'sample.pdf'}")


if __name__ == "__main__":
    main()
