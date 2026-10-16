"""Fail if packaged redaction left fictional source terms in the output.

Docx and PDF checks read extracted text. The png check OCRs with the
tesseract executable passed on the command line (the copy inside the zip),
not with PATH or TESSERACT_CMD.
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path

SECRETS = ("ABC-2026-001", "Acme Pharma", "acme pharma", "Fictizumab")
# The sample project id is split across runs in the docx. Leftover pieces count.
FRAGMENTS = SECRETS + ("ABC-20", "26-001")
OCR_TOKENS = ("abc-2026-001", "acmepharma", "fictizumab")
REPORT_KEYS = ("DRUG", "PROJECT_ID", "SPONSOR")


def fail(message: str) -> None:
    print(message, file=sys.stderr)
    raise SystemExit(1)


def _container_text(container) -> list[str]:
    chunks = [paragraph.text for paragraph in container.paragraphs]
    for table in container.tables:
        seen: set[int] = set()
        for row in table.rows:
            for cell in row.cells:
                marker = id(cell._tc)
                if marker in seen:
                    continue
                seen.add(marker)
                chunks.extend(_container_text(cell))
    return chunks


def docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    chunks = _container_text(document)
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            chunks.extend(_container_text(part))
    return "\n".join(chunks)


def pdf_text(path: Path) -> str:
    import pymupdf

    document = pymupdf.open(path)
    try:
        return "\n".join(page.get_text() for page in document)
    finally:
        document.close()


def ocr_png(tesseract: Path, image: Path) -> str:
    # Sibling of the given path, not of a symlink target. The zip lays tessdata
    # next to tesseract.exe; resolving the exe would look in the wrong place.
    tessdata = tesseract.parent / "tessdata"
    if not tessdata.is_dir():
        fail(f"bundled tessdata directory not found beside {tesseract}")
    completed = subprocess.run(
        [
            str(tesseract),
            str(image),
            "stdout",
            "-l",
            "chi_sim+eng",
            "--tessdata-dir",
            str(tessdata),
        ],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=_ocr_env(tessdata),
    )
    if completed.returncode != 0:
        detail = (completed.stderr or completed.stdout or "").strip()
        fail(f"bundled tesseract failed on {image.name}: {detail}")
    return completed.stdout


def _ocr_env(tessdata: Path) -> dict[str, str]:
    import os

    env = os.environ.copy()
    env.pop("TESSERACT_CMD", None)
    env["TESSDATA_PREFIX"] = str(tessdata)
    return env


def _reject_terms(label: str, text: str, terms: tuple[str, ...]) -> None:
    for term in terms:
        if term in text:
            fail(f"{label} still contains {term!r}")


def check_docx(path: Path) -> None:
    if not path.is_file():
        fail(f"missing {path.name}")
    text = docx_text(path)
    _reject_terms(path.name, text, FRAGMENTS)
    if "[REDACTED:" not in text:
        fail(f"{path.name} has no redaction marker")
    if "fictizumab" not in text:
        fail(f"{path.name} lost the lowercase fictizumab that the sample keeps")
    print(f"ok {path.name}")


def check_pdf(path: Path) -> None:
    if not path.is_file():
        fail(f"missing {path.name}")
    text = pdf_text(path)
    _reject_terms(path.name, text, FRAGMENTS)
    if "fictizumab stays" not in text:
        fail(f"{path.name} is missing the unredacted line 'fictizumab stays'")
    print(f"ok {path.name}")


def check_png(path: Path, tesseract: Path) -> None:
    if not path.is_file():
        fail(f"missing {path.name}")
    folded = "".join(ocr_png(tesseract, path).split()).lower()
    for token in OCR_TOKENS:
        if token in folded:
            fail(f"{path.name} OCR still contains {token!r}")
    print(f"ok {path.name}")


def check_report(path: Path) -> None:
    if not path.is_file():
        fail(f"missing {path.name}")
    raw = path.read_text(encoding="utf-8")
    _reject_terms(path.name, raw, SECRETS)
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError as exc:
        fail(f"{path.name} is not json: {exc}")
    if not isinstance(payload, dict):
        fail(f"{path.name} is not an object")
    for key in REPORT_KEYS:
        value = payload.get(key)
        if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
            fail(f"{path.name} missing a positive count for {key}")
    for key, value in payload.items():
        if not isinstance(key, str):
            fail(f"{path.name} has a non-string category")
        if isinstance(value, bool) or not isinstance(value, int):
            fail(f"{path.name} count for {key} is not an int")
    print(f"ok {path.name}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Check redacted sample output for leftover source terms.")
    parser.add_argument("--output", required=True, type=Path, help="Directory written by doc-redactor.exe")
    parser.add_argument("--tesseract", required=True, type=Path, help="Packaged tesseract executable")
    args = parser.parse_args(argv)
    output = args.output
    tesseract = args.tesseract
    if not output.is_dir():
        fail(f"output directory not found: {output}")
    if not tesseract.is_file():
        fail(f"tesseract executable not found: {tesseract}")
    check_docx(output / "sample.redacted.docx")
    check_pdf(output / "sample.redacted.pdf")
    check_png(output / "sample.redacted.png", tesseract)
    check_report(output / "sample.report.json")
    for extra in sorted(output.glob("*.report.json")):
        if extra.name != "sample.report.json":
            check_report(extra)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
