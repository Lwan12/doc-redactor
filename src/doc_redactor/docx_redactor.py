"""Redact docx body, tables, headers/footers, notes, comments, text boxes, fields,
content controls, revisions, alt text, link targets, bookmarks, embedded objects, and images."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from urllib.parse import quote, unquote

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.part import PartFactory, XmlPart
from docx.table import Table, _Cell

# python-docx loads footnotes and endnotes as opaque bytes. As XmlPart they get
# an element tree that is written back on save, so the same walk can edit them.
PartFactory.part_type_for.setdefault(CT.WML_FOOTNOTES, XmlPart)
PartFactory.part_type_for.setdefault(CT.WML_ENDNOTES, XmlPart)

from doc_redactor.engine import Hit, RedactionEngine, merge_counts, redaction_label
from doc_redactor.image_redactor import is_raster_suffix, redact_image_bytes


def redact_docx(src: Path, dest: Path, engine: RedactionEngine) -> dict[str, int]:
    if src.resolve() == dest.resolve():
        raise ValueError("输出路径与原文件相同，已拒绝覆盖原文件")
    document = Document(str(src))
    counts: dict[str, int] = {}
    for unit in iter_document_paragraphs(document):
        merge_counts(counts, _counts(_redact_paragraph(unit, engine)))
    finish_document(document)
    dest.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(dest))
    merge_counts(counts, _redact_media(dest, engine))
    return counts


def iter_table_cells(table: Table):
    """Yield each physical ``w:tc`` once, in XML order.

    ``row.cells`` repeats merged cells, and deduping by ``id(cell._tc)`` breaks
    because lxml proxies are freed and ids reused. Walking ``w:tc`` directly
    visits every cell exactly once (gridSpan and vMerge included).
    """
    for tr in table._tbl.tr_lst:
        for tc in tr.tc_lst:
            yield _Cell(tc, table)


def _owned_headers_footers(section):
    """Yield header/footer parts this section defines.

    Reading ``.paragraphs`` on a missing part creates an empty header. Skip
    parts that are only linked to the previous section.
    """
    parts = (
        section.header,
        section.footer,
        section.first_page_header,
        section.first_page_footer,
        section.even_page_header,
        section.even_page_footer,
    )
    for part in parts:
        if not part.is_linked_to_previous:
            yield part


def _redact_paragraph(unit: TextUnit, engine: RedactionEngine) -> list[str]:
    """Replace each hit inside one text unit, keeping the first involved run's formatting.

    Matches are applied from the right. The placeholder is written into the
    first segment that overlaps the hit; text inside the hit is removed from
    the later segments. Every stored copy of the unit gets the same change.
    """
    categories: list[str] = []
    for _ in range(10000):
        hits = unit.find_hits(engine)
        if not hits:
            break
        hit = hits[-1]
        if not unit.apply(hit, engine.label(hit.category)):
            break
        categories.append(hit.category)
    else:
        raise RuntimeError("同一段落命中过多，已停止")
    categories.reverse()
    return categories


# ---------------------------------------------------------------------------
# One walk over the XML for scanning and for writing.
#
# python-docx's ``.paragraphs`` only sees direct ``w:p`` children and
# ``.runs`` only direct ``w:r`` children, so text inside text boxes, content
# controls (``w:sdt``), simple fields (``w:fldSimple``), hyperlinks and
# tracked insertions (``w:ins``) was never read. This walk visits every
# ``w:p`` under a part and every ``w:r`` a paragraph owns, at any depth.

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
_W_P = _W + "p"
_W_R = _W + "r"
_W_DELTEXT = _W + "delText"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_INSERTED = (_W + "ins", _W + "moveTo")
_MC_ALT = _MC + "AlternateContent"
_MC_BRANCHES = (_MC + "Choice", _MC + "Fallback")

KIND_TEXT = "text"
KIND_DELETED = "deleted"


class _RunSegment:
    """Visible text of one ``w:r`` (``w:t``, tabs, breaks). Not ``w:delText``."""

    __slots__ = ("_r", "writable")

    def __init__(self, r, writable: bool = True) -> None:
        self._r = r
        self.writable = writable

    @property
    def text(self) -> str:
        return self._r.text

    @text.setter
    def text(self, value: str) -> None:
        # CT_R.text keeps w:rPr and rebuilds w:t / w:tab / w:br.
        self._r.text = value


class _DelTextSegment:
    """One ``w:delText``: text removed under tracked changes, still in the file."""

    __slots__ = ("_el", "writable")

    def __init__(self, element) -> None:
        self._el = element
        self.writable = True

    @property
    def text(self) -> str:
        return self._el.text or ""

    @text.setter
    def text(self, value: str) -> None:
        self._el.text = value
        self._el.set(_XML_SPACE, "preserve")


class TextUnit:
    """One paragraph's text as the scanner sees it, plus where to write it back.

    ``kind`` is ``"text"`` for what the reader sees (plain runs, hyperlinks,
    fields, content controls, insertions), or ``"deleted"`` for the paragraph
    as it read before tracked changes: kept runs plus ``w:delText``, without
    insertions. A deleted unit only reports hits that touch deleted text, and
    only rewrites deleted text, so the kept runs stay for the text unit.

    ``copies`` holds the segment lists of every stored copy: a text box saved
    as both a DrawingML and a VML shape is one unit with two copies, so one
    decision covers both.
    """

    __slots__ = ("kind", "copies", "where")

    def __init__(self, kind: str, copies: list[list], where: str = "") -> None:
        self.kind = kind
        self.copies = copies
        # "" for a paragraph of the body or a header/footer; otherwise the
        # corner it comes from (WHERE_*), used for the location label.
        self.where = where

    @property
    def text(self) -> str:
        return "".join(segment.text for segment in self.copies[0])

    @property
    def is_deletion(self) -> bool:
        return self.kind == KIND_DELETED

    def find_hits(self, engine: RedactionEngine) -> list[Hit]:
        """Hits in this unit. The one place that decides what a unit reports."""
        hits = engine.find_hits(self.text)
        if not self.is_deletion:
            return hits
        writable = [span for span in _spans(self.copies[0]) if span[0].writable]
        return [hit for hit in hits if any(s[1] < hit.end and s[2] > hit.start for s in writable)]

    def apply(self, hit: Hit, replacement: str) -> bool:
        applied = False
        for segments in self.copies:
            if _apply_hit(_spans(segments), hit, replacement=replacement):
                applied = True
        return applied


def iter_text_units(root, where: str = ""):
    """Yield every text unit under ``root`` (a body, header, footer or notes element), in XML order.

    A paragraph with tracked deletions yields its deleted unit first. That
    unit only rewrites ``w:delText``, so the offsets of the text unit that
    follows stay valid.
    """
    for group in _iter_paragraph_groups(root):
        owned = [list(_own_runs(p)) for p in group]
        if any(r.find(_W_DELTEXT) is not None for runs in owned for r, _inserted in runs):
            yield TextUnit(KIND_DELETED, [_before_changes(runs) for runs in owned], where)
        yield TextUnit(KIND_TEXT, [[_RunSegment(r) for r, _inserted in runs] for runs in owned], where)


def _before_changes(runs) -> list:
    """Segments of the paragraph as it read before tracked changes."""
    segments: list = []
    for r, inserted in runs:
        if inserted:
            continue
        deleted = list(r.iter(_W_DELTEXT))
        if deleted:
            segments.extend(_DelTextSegment(el) for el in deleted)
        else:
            segments.append(_RunSegment(r, writable=False))
    return segments


def _iter_paragraph_groups(element):
    """Yield lists of ``w:p``: one paragraph, or the mirrored copies of one."""
    for child in element:
        tag = child.tag
        if not isinstance(tag, str):
            continue
        if tag == _W_P:
            yield [child]
            # Text boxes anchored in this paragraph hold their own paragraphs.
            yield from _iter_paragraph_groups(child)
        elif tag == _MC_ALT:
            yield from _iter_alternate(child)
        else:
            yield from _iter_paragraph_groups(child)


def _iter_alternate(alternate):
    """mc:AlternateContent stores the same content once per branch.

    When every branch holds the same paragraphs with the same text, pair them
    into one group per paragraph. Otherwise every branch is walked on its own,
    so nothing is skipped.
    """
    branches = [
        list(_iter_paragraph_groups(child))
        for child in alternate
        if isinstance(child.tag, str) and child.tag in _MC_BRANCHES
    ]
    if len(branches) >= 2 and _mirrored(branches):
        for groups in zip(*branches):
            yield [p for group in groups for p in group]
        return
    for groups in branches:
        yield from groups


def _mirrored(branches) -> bool:
    first = branches[0]
    for other in branches[1:]:
        if len(other) != len(first):
            return False
        for a, b in zip(first, other):
            if len(a) != len(b):
                return False
            if any(_paragraph_key(x) != _paragraph_key(y) for x, y in zip(a, b)):
                return False
    return True


def _paragraph_key(p) -> tuple[str, str]:
    runs = [r for r, _inserted in _own_runs(p)]
    live = "".join(r.text for r in runs)
    deleted = "".join(el.text or "" for r in runs for el in r.iter(_W_DELTEXT))
    return live, deleted


def _own_runs(element, inserted: bool = False):
    """``(w:r, inside a tracked insertion)`` this paragraph owns, at any depth, in XML order.

    Stops at nested ``w:p`` (text box paragraphs are their own units) and does
    not look inside a run.
    """
    for child in element:
        tag = child.tag
        if not isinstance(tag, str) or tag == _W_P:
            continue
        if tag == _W_R:
            yield child, inserted
        else:
            yield from _own_runs(child, inserted or tag in _INSERTED)


def _spans(segments):
    spans = []
    position = 0
    for segment in segments:
        segment_text = segment.text
        if segment_text:
            # Empty runs (a drawing, a field mark) never take part in a hit,
            # so a hit around them never clears them.
            spans.append((segment, position, position + len(segment_text), segment_text))
        position += len(segment_text)
    return spans


WHERE_FOOTNOTE = "脚注"
WHERE_ENDNOTE = "尾注"
WHERE_COMMENT = "批注"
WHERE_FIELD = "域代码"
WHERE_ALT = "图片替代文字"
WHERE_LINK = "超链接地址"
WHERE_ATTR = "超链接提示和内容控件名称"
WHERE_AUTHOR = "修订和批注作者"
WHERE_STYLE = "样式名"


def iter_document_paragraphs(document):
    """Yield text units in a fixed order: body, headers/footers, footnotes,
    endnotes, comments, then field codes, alt text and other attributes, then
    external link targets.

    Scanning, writing and the after-write check all use this one walk, and a
    stored hit is matched to its unit by position, so the order must not
    depend on earlier edits. Each unit has ``.text``; the name is kept for
    existing callers.
    """
    roots = [(document.part, document._element)]
    yield from iter_text_units(document._element)
    for section in document.sections:
        for header_footer in _owned_headers_footers(section):
            roots.append((header_footer.part, header_footer._element))
            yield from iter_text_units(header_footer._element)
    for part, where in _note_parts(document):
        roots.append((part, part.element))
        yield from iter_text_units(part.element, where)
    for _part, root in roots:
        yield from _iter_field_units(root)
        yield from _iter_attribute_units(root)
    styles = _styles_element(document)
    if styles is not None:
        for name in styles.iter(_W + "name"):
            if name.get(_W + "val"):
                yield TextUnit(KIND_TEXT, [[_AttrSegment(name, _W + "val")]], WHERE_STYLE)
    seen: set[int] = set()
    for part, _root in roots:
        if id(part) in seen:
            continue
        seen.add(id(part))
        for rId in sorted(part.rels):
            rel = part.rels[rId]
            if rel.is_external and rel.reltype == RT.HYPERLINK:
                yield TextUnit(KIND_TEXT, [[_LinkSegment(rel)]], WHERE_LINK)


def _note_parts(document):
    """Footnotes, endnotes and comments parts that have an element tree, in that order."""
    wanted = ((RT.FOOTNOTES, WHERE_FOOTNOTE), (RT.ENDNOTES, WHERE_ENDNOTE), (RT.COMMENTS, WHERE_COMMENT))
    rels = document.part.rels
    for reltype, where in wanted:
        for rId in sorted(rels):
            rel = rels[rId]
            if rel.is_external or rel.reltype != reltype:
                continue
            element = getattr(rel.target_part, "element", None)
            if element is not None:
                yield rel.target_part, where


_O = "{urn:schemas-microsoft-com:office:office}"
_V = "{urn:schemas-microsoft-com:vml}"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_W_INSTR = (_W + "instrText", _W + "delInstrText")
_W_FLDCHAR = _W + "fldChar"
_W_FLDTYPE = _W + "fldCharType"
_W_FLDSIMPLE = _W + "fldSimple"
_W_BOOKMARK = _W + "bookmarkStart"
_W_HYPERLINK = _W + "hyperlink"

# (local tag name, namespace or None for any, attributes, where)
_ATTRIBUTES = (
    ("docPr", None, ("descr", "title", "name"), WHERE_ALT),
    ("cNvPr", None, ("descr", "title", "name"), WHERE_ALT),
    ("imagedata", _V, (_O + "title",), WHERE_ALT),
    ("shape", _V, ("alt", "title", _O + "title"), WHERE_ALT),
    ("hyperlink", _W, (_W + "tooltip",), WHERE_ATTR),
    ("alias", _W, (_W + "val",), WHERE_ATTR),
    ("tag", _W, (_W + "val",), WHERE_ATTR),
    # Revision marks and comments: who made them.
    (None, _W, (_W + "author", _W + "initials"), WHERE_AUTHOR),
)


def _iter_attribute_units(root):
    """One unit per attribute value a parser shows as text: image alt text,
    VML titles, link tooltips, content control names."""
    for element in root.iter():
        tag = element.tag
        if not isinstance(tag, str):
            continue
        namespace, _, local = tag[1:].partition("}") if tag.startswith("{") else ("", "", tag)
        for name, ns, attributes, where in _ATTRIBUTES:
            if (name is not None and local != name) or (ns is not None and "{" + namespace + "}" != ns):
                continue
            for attribute in attributes:
                if element.get(attribute):
                    yield TextUnit(KIND_TEXT, [[_AttrSegment(element, attribute)]], where)


def _iter_field_units(root):
    """One unit per field instruction: ``w:fldSimple/@w:instr``, or the
    ``w:instrText`` runs between a field's begin and separate marks."""
    stack: list[list] = []
    for element in root.iter(_W_FLDCHAR, _W_FLDSIMPLE, *_W_INSTR):
        tag = element.tag
        if tag == _W_FLDSIMPLE:
            if element.get(_W + "instr"):
                yield TextUnit(KIND_TEXT, [[_AttrSegment(element, _W + "instr")]], WHERE_FIELD)
            continue
        if tag == _W_FLDCHAR:
            kind = element.get(_W_FLDTYPE)
            if kind == "begin":
                stack.append([])
            elif kind in ("separate", "end") and stack:
                segments = stack.pop()
                if segments:
                    yield TextUnit(KIND_TEXT, [segments], WHERE_FIELD)
            continue
        segment = _ElementTextSegment(element)
        if stack:
            stack[-1].append(segment)
        else:
            yield TextUnit(KIND_TEXT, [[segment]], WHERE_FIELD)
    for segments in stack:
        if segments:
            yield TextUnit(KIND_TEXT, [segments], WHERE_FIELD)


class _AttrSegment:
    """One attribute value, read and written as text."""

    __slots__ = ("_el", "_name", "writable")

    def __init__(self, element, name: str) -> None:
        self._el = element
        self._name = name
        self.writable = True

    @property
    def text(self) -> str:
        return self._el.get(self._name) or ""

    @text.setter
    def text(self, value: str) -> None:
        self._el.set(self._name, value)


class _ElementTextSegment:
    """The text of one element such as ``w:instrText``."""

    __slots__ = ("_el", "writable")

    def __init__(self, element) -> None:
        self._el = element
        self.writable = True

    @property
    def text(self) -> str:
        return self._el.text or ""

    @text.setter
    def text(self, value: str) -> None:
        self._el.text = value
        self._el.set(_XML_SPACE, "preserve")


# Characters left as they are when a link target is written back.
_URL_SAFE = ":/?#[]@!$&'()*+,;=%-._~"


class _LinkSegment:
    """The target of an external hyperlink, percent escapes decoded for matching."""

    __slots__ = ("_rel", "writable")

    def __init__(self, rel) -> None:
        self._rel = rel
        self.writable = True

    @property
    def text(self) -> str:
        return unquote(str(self._rel._target or ""))

    @text.setter
    def text(self, value: str) -> None:
        self._rel._target = quote(value, safe=_URL_SAFE)


def _styles_element(document):
    try:
        return document.styles.element
    except Exception:
        return None


def finish_document(document) -> None:
    """Changes made on every write, after the hits: rename bookmarks, drop
    embedded objects, and drop parts no rule can be applied to."""
    _rename_bookmarks(document)
    _rename_custom_styles(document)
    _drop_embedded_objects(document)
    _drop_unscannable(document)


_STYLE_REFS = ("pStyle", "rStyle", "tblStyle", "basedOn", "next", "link", "numStyleLink", "styleLink")


def _rename_custom_styles(document) -> None:
    """Give styles the user created a neutral id (``UserStyle1``...).

    Word builds a style id from its name without spaces (``Acme Pharma
    Heading`` becomes ``AcmePharmaHeading``), so a rule that needs a word
    edge never matches it. The display name is a scanned unit; the id is
    renamed here, and every reference to it follows.
    """
    styles = _styles_element(document)
    if styles is None:
        return
    ids: dict[str, str] = {}
    for style in styles.iter(_W + "style"):
        if style.get(_W + "customStyle") not in ("1", "true", "on"):
            continue
        old = style.get(_W + "styleId")
        if old and old not in ids:
            ids[old] = f"UserStyle{len(ids) + 1}"
            style.set(_W + "styleId", ids[old])
    if not ids:
        return
    for _part, element in _xml_parts(document):
        for local in _STYLE_REFS:
            for node in element.iter(_W + local):
                value = node.get(_W + "val")
                if value in ids:
                    node.set(_W + "val", ids[value])


def _drop_unscannable(document) -> None:
    """Remove what holds text the tool cannot redact in place.

    - ``w:altChunk``: an HTML or other file Word merges into the body when it
      opens the document. Its content is not Word XML, so it is removed.
    - ``w:docVars``: hidden name/value pairs in settings.
    - ``w:attachedTemplate``: the template path, often with a folder or user name.
    """
    for part, element in list(_xml_parts(document)):
        doomed: set[str] = set()
        for tag in (_W + "altChunk", _W + "attachedTemplate"):
            for node in list(element.iter(tag)):
                rId = node.get(_R_ID)
                if rId:
                    doomed.add(rId)
                node.getparent().remove(node)
        for node in list(element.iter(_W + "docVars")):
            node.getparent().remove(node)
        for rId in doomed:
            if rId in part.rels and not any(n.get(_R_ID) == rId for n in element.iter()):
                part.rels.pop(rId)


def _xml_parts(document):
    for part in document.part.package.iter_parts():
        element = getattr(part, "element", None)
        if element is not None:
            yield part, element


def _rename_bookmarks(document) -> None:
    """Give every bookmark the user named a neutral name (``bm1``, ``bm2``...).

    A bookmark name such as ``Acme_Pharma_ABC_2026_001`` is not split by
    whitespace, so a rule never matches it, and pandoc prints it as an anchor.
    Names Word makes itself start with ``_`` (``_Toc…``, ``_Ref…``, ``_GoBack``)
    and are kept. Internal links (``w:anchor``) and field instructions that
    name a renamed bookmark follow the new name.
    """
    parts = list(_xml_parts(document))
    names: dict[str, str] = {}
    for _part, element in parts:
        for bookmark in element.iter(_W_BOOKMARK):
            name = bookmark.get(_W + "name") or ""
            if not name or name.startswith("_"):
                continue
            if name not in names:
                names[name] = f"bm{len(names) + 1}"
            bookmark.set(_W + "name", names[name])
    if not names:
        return
    pattern = re.compile(r"(?<![\w])(" + "|".join(re.escape(name) for name in sorted(names, key=len, reverse=True)) + r")(?![\w])")
    for _part, element in parts:
        for link in element.iter(_W_HYPERLINK):
            anchor = link.get(_W + "anchor")
            if anchor in names:
                link.set(_W + "anchor", names[anchor])
        for instr in element.iter(*_W_INSTR):
            if instr.text:
                instr.text = pattern.sub(lambda m: names[m.group(1)], instr.text)
        for simple in element.iter(_W_FLDSIMPLE):
            value = simple.get(_W + "instr")
            if value:
                simple.set(_W + "instr", pattern.sub(lambda m: names[m.group(1)], value))


_EMBED_RELTYPES = (RT.OLE_OBJECT, RT.PACKAGE)


def _drop_embedded_objects(document) -> None:
    """Remove embedded objects (an Excel sheet, an OLE ``.bin``); keep their preview image.

    LibreOffice renders an embedded sheet in full when it converts the file,
    and the object bytes carry the original text. The ``o:OLEObject`` that
    points at the object goes, the ``v:shape`` with the preview image stays,
    so the page still shows the picture. Once no relationship points at the
    object part, python-docx does not write it.
    """
    for part, element in list(_xml_parts(document)):
        doomed = [rId for rId, rel in part.rels.items() if not rel.is_external and rel.reltype in _EMBED_RELTYPES]
        if not doomed:
            continue
        for node in list(element.iter()):
            # o:OLEObject points at an embedded object; c:externalData is a
            # chart's embedded workbook (the chart keeps its cached values).
            if node.get(_R_ID) in doomed and isinstance(node.tag, str) and node.tag.endswith(("}OLEObject", "}externalData")):
                parent = node.getparent()
                if parent is not None:
                    parent.remove(node)
        for rId in doomed:
            still_used = any(node.get(_R_ID) == rId for node in element.iter())
            if not still_used:
                part.rels.pop(rId)


def apply_hits_to_paragraph(unit: TextUnit, hits: list[Hit], replacement_for) -> list[str]:
    """Apply stored hits to one text unit from the right. replacement_for(category) -> str."""
    ordered = sorted(hits, key=lambda item: item.start, reverse=True)
    categories: list[str] = []
    for hit in ordered:
        text = unit.text
        rebuilt = Hit(
            start=hit.start,
            end=hit.end,
            category=hit.category,
            text=text[hit.start : hit.end] if 0 <= hit.start <= hit.end <= len(text) else hit.text,
        )
        if not unit.apply(rebuilt, replacement_for(hit.category)):
            continue
        categories.append(hit.category)
    categories.reverse()
    return categories


def _apply_hit(spans, hit: Hit, replacement: str | None = None) -> bool:
    """Write the placeholder into the first writable span the hit touches.

    The hit's text is removed from the other writable spans it touches.
    Read-only spans (kept runs seen from a deleted unit) are left alone.
    """
    involved = [span for span in spans if span[1] < hit.end and span[2] > hit.start and span[0].writable]
    if not involved:
        return False
    if replacement is None:
        replacement = redaction_label(hit.category)
    for number, (segment, start, end, text) in enumerate(involved):
        local_start = max(hit.start, start) - start
        local_end = min(hit.end, end) - start
        middle = replacement if number == 0 else ""
        segment.text = text[:local_start] + middle + text[local_end:]
    return True


def _counts(categories: list[str]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for category in categories:
        counts[category] = counts.get(category, 0) + 1
    return counts


def _redact_media(path: Path, engine: RedactionEngine) -> dict[str, int]:
    with zipfile.ZipFile(path, "r") as archive:
        names = [info.filename for info in archive.infolist() if _is_media_image(info.filename)]
        if not names:
            return {}
        payloads = {name: archive.read(name) for name in names}
        infos = archive.infolist()

    counts: dict[str, int] = {}
    replaced: dict[str, bytes] = {}
    for name, payload in payloads.items():
        suffix = Path(name).suffix
        new_bytes, image_counts = redact_image_bytes(payload, suffix, engine)
        replaced[name] = new_bytes
        merge_counts(counts, image_counts)

    buffer_path = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(buffer_path, "w") as target:
        for info in infos:
            data = replaced.get(info.filename, source.read(info.filename))
            cloned = zipfile.ZipInfo(filename=info.filename, date_time=info.date_time)
            cloned.compress_type = info.compress_type
            cloned.external_attr = info.external_attr
            cloned.flag_bits = info.flag_bits
            cloned.create_system = info.create_system
            target.writestr(cloned, data)
    buffer_path.replace(path)
    return counts


def _is_media_image(name: str) -> bool:
    lowered = name.lower()
    if not lowered.startswith("word/media/"):
        return False
    return is_raster_suffix(Path(lowered).suffix)
