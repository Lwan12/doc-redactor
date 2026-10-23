from __future__ import annotations

import csv
import io
import os
import re
import struct
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import unescape

from openpyxl import load_workbook

from doc_redactor.engine import RedactionEngine
from doc_redactor.scan import (
    MSG_OPEN_FAIL,
    CellPlace,
    CommentPlace,
    FileScan,
    HeaderFooterPlace,
    IdGen,
    LocatedHit,
    SheetCellView,
    SheetNamePlace,
    SheetView,
    _finish,
)
from doc_redactor.verify_output import count_not_allowed

MSG_ENCRYPTED_SHEET = "加密文件，没法检查，未写出。"
MSG_PICTURE_SHEET = "图片里的字没法检查，未写出。"

OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
_ILLEGAL_TITLE = str.maketrans("", "", "\\/?*[]:")
_CELL_RE = re.compile(r"<c\b([^>/]*?)(?:/>|>(.*?)</c>)", re.DOTALL)
_REL_RE = re.compile(r"<Relationship\b[^>]*/?>", re.IGNORECASE)
_SHEET_TAG_RE = re.compile(r"<sheet\b[^>]*/?>", re.IGNORECASE)

BIFF_FORMULA = 0x0006
BIFF_EOF = 0x000A
BIFF_HEADER = 0x0014
BIFF_FOOTER = 0x0015
BIFF_FILEPASS = 0x002F
BIFF_CONTINUE = 0x003C
BIFF_OBJ = 0x005D
BIFF_COLINFO = 0x007D
BIFF_BOUNDSHEET = 0x0085
BIFF_SST = 0x00FC
BIFF_EXTSST = 0x00FF
BIFF_LABELSST = 0x00FD
BIFF_DBCELL = 0x00D7
BIFF_BOF = 0x0809
BIFF_INDEX = 0x020B
BIFF_LABEL = 0x0204
BIFF_ROW = 0x0208
BIFF_MAX = 8224


class SheetRefused(Exception):
    def __init__(self, sentence: str) -> None:
        self.sentence = sentence
        super().__init__(sentence)


class _HeaderPicture(Exception):
    pass


class InScope:
    def __init__(self, parts: list[str]) -> None:
        self.parts = parts

    def blob(self) -> str:
        return "\n".join(self.parts)


@dataclass
class _Cell:
    column: int
    row: int
    text: str
    formula: bool
    hidden_row: bool
    hidden_col: bool


@dataclass
class _Comment:
    column: int
    row: int
    text: str
    author: str


@dataclass
class _Header:
    slot: str
    raw: str
    literals: list[str]


@dataclass
class _Sheet:
    title: str
    cells: list[_Cell]
    comments: list[_Comment]
    headers: list[_Header]
    name_is_hit: bool = True


@dataclass
class _Book:
    sheets: list[_Sheet]
    formulas: list[str] = field(default_factory=list)


def extract_in_scope(path: Path) -> InScope:
    book = _load(Path(path), refuse=False)
    parts: list[str] = []
    for sheet in book.sheets:
        parts.append(sheet.title)
        for cell in sheet.cells:
            if not cell.formula:
                parts.append(cell.text)
        for comment in sheet.comments:
            parts.append(comment.text)
            parts.append(comment.author)
        for header in sheet.headers:
            parts.extend(header.literals)
    return InScope(parts)


def extract_formulas(path: Path) -> list[str]:
    return list(_load(Path(path), refuse=False).formulas)


def fill_scan(path: Path, engine: RedactionEngine, ids: IdGen, base: FileScan) -> FileScan:
    book = _load(Path(path), refuse=True)
    hits: list[LocatedHit] = []
    views: list[SheetView] = []
    for index, sheet in enumerate(book.sheets):
        by_cell: dict[tuple[int, int], list[int]] = {}
        for cell in sheet.cells:
            if cell.formula:
                continue
            before = len(hits)
            _collect(
                hits,
                ids,
                engine,
                cell.text,
                CellPlace(index, cell.column, cell.row, cell.hidden_row, cell.hidden_col),
                _cell_location(cell),
                f"s:{index}:c:{cell.column}:{cell.row}",
            )
            by_cell[(cell.column, cell.row)] = [hit.hit_id for hit in hits[before:]]
        for comment in sheet.comments:
            where = f"{_column_letter(comment.column)}{comment.row + 1}"
            _collect(
                hits,
                ids,
                engine,
                comment.text,
                CommentPlace(index, comment.column, comment.row, "text"),
                f"{where} 批注",
                f"s:{index}:note:{comment.column}:{comment.row}:text",
            )
            _collect(
                hits,
                ids,
                engine,
                comment.author,
                CommentPlace(index, comment.column, comment.row, "author"),
                f"{where} 批注作者",
                f"s:{index}:note:{comment.column}:{comment.row}:author",
            )
        for header in sheet.headers:
            kind = "页眉" if header.slot.startswith("h:") else "页脚"
            for literal_index, literal in enumerate(header.literals):
                _collect(
                    hits,
                    ids,
                    engine,
                    literal,
                    HeaderFooterPlace(index, header.slot, literal_index),
                    kind,
                    f"s:{index}:{header.slot}:{literal_index}",
                )
        if sheet.name_is_hit:
            _collect(
                hits,
                ids,
                engine,
                sheet.title,
                SheetNamePlace(index),
                "工作表名",
                f"s:{index}:name",
            )
        views.append(
            SheetView(
                title=sheet.title,
                index=index,
                cells=tuple(
                    SheetCellView(
                        column=cell.column,
                        row=cell.row,
                        text=cell.text,
                        formula=cell.formula,
                        hidden_row=cell.hidden_row,
                        hidden_col=cell.hidden_col,
                        hit_ids=tuple(by_cell.get((cell.column, cell.row), ())),
                    )
                    for cell in sheet.cells
                ),
            )
        )
    base.hits = hits
    base.sheets = tuple(views)
    return _finish(base)


def write_sheet(scan: FileScan, dest: Path, label) -> dict[str, int]:
    dest = Path(dest)
    suffix = Path(scan.path).suffix.lower()
    if suffix == ".xlsx":
        blob = _write_xlsx(scan, label)
    elif suffix == ".xls":
        blob = _write_xls(scan, label)
    elif suffix == ".csv":
        blob = _write_csv(scan, label)
    else:
        raise ValueError(MSG_OPEN_FAIL)
    tmp = dest.with_name(f".{dest.name}.writing")
    try:
        tmp.write_bytes(blob)
        os.replace(tmp, dest)
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)
    counts: dict[str, int] = {}
    for hit in scan.enabled_hits:
        counts[hit.category] = counts.get(hit.category, 0) + 1
    return counts


def sheet_leftover_count(dest: Path, enabled_texts: list[str], disabled_texts: list[str]) -> int:
    blob = extract_in_scope(Path(dest)).blob()
    found: Counter[str] = Counter()
    for text in set(enabled_texts) | set(disabled_texts):
        if text:
            found[text] = blob.count(text)
    return count_not_allowed(found, disabled_texts)


def _load(path: Path, *, refuse: bool) -> _Book:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _load_xlsx(path, refuse=refuse)
    if suffix == ".xls":
        return _load_xls(path, refuse=refuse)
    if suffix == ".csv":
        return _load_csv(path)
    raise ValueError(MSG_OPEN_FAIL)


def _collect(hits, ids, engine, text: str, place, location: str, unit: str) -> None:
    for found in engine.find_hits(text or ""):
        hits.append(
            LocatedHit(
                hit_id=ids.next(),
                category=found.category,
                text=found.text,
                location=location,
                enabled=True,
                unit=unit,
                start=found.start,
                end=found.end,
                page=0,
                place=place,
            )
        )


def _column_letter(index: int) -> str:
    n = index + 1
    letters = []
    while n:
        n, rem = divmod(n - 1, 26)
        letters.append(chr(65 + rem))
    return "".join(reversed(letters))


def _cell_location(cell: _Cell) -> str:
    text = f"{_column_letter(cell.column)}{cell.row + 1}"
    extra = []
    if cell.hidden_row:
        extra.append("隐藏行")
    if cell.hidden_col:
        extra.append("隐藏列")
    if extra:
        text += "，" + "，".join(extra)
    return text


def _coord(column: int, row: int) -> str:
    return f"{_column_letter(column)}{row + 1}"


def _split_ref(ref: str) -> tuple[int, int]:
    match = re.match(r"([A-Z]+)(\d+)$", ref or "")
    if match is None:
        raise ValueError(MSG_OPEN_FAIL)
    col = 0
    for char in match.group(1):
        col = col * 26 + (ord(char) - 64)
    return col - 1, int(match.group(2)) - 1


def _apply(text: str, hits: list[LocatedHit], label) -> str:
    chosen = [hit for hit in hits if hit.enabled]
    out = text
    for hit in sorted(chosen, key=lambda item: item.start, reverse=True):
        if out[hit.start : hit.end] != hit.text:
            raise ValueError(MSG_OPEN_FAIL)
        out = out[: hit.start] + label(hit.category) + out[hit.end :]
    return out


def _hits_at(scan: FileScan, predicate) -> list[LocatedHit]:
    return [hit for hit in scan.hits if predicate(hit.place)]


def _xml_text(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _xml_attr(text: str) -> str:
    return _xml_text(text).replace('"', "&quot;")


def _attr(tag: str, name: str) -> str | None:
    match = re.search(rf'\b{name}="([^"]*)"', tag)
    return match.group(1) if match else None


def tokenize_header(raw: str) -> list[tuple[str, str]]:
    """Excel header codes stay. Text pieces are what the engine sees. ``&&`` is one ampersand."""
    tokens: list[tuple[str, str]] = []
    buf: list[str] = []

    def flush() -> None:
        if buf:
            tokens.append(("text", "".join(buf)))
            buf.clear()

    i = 0
    while i < len(raw):
        if raw[i] != "&" or i + 1 >= len(raw):
            buf.append(raw[i])
            i += 1
            continue
        nxt = raw[i + 1]
        if nxt == "&":
            buf.append("&")
            i += 2
            continue
        if nxt == "G":
            raise _HeaderPicture()
        if nxt in "LCR":
            flush()
            tokens.append(("code", "&" + nxt))
            i += 2
            continue
        if nxt == '"':
            end = raw.find('"', i + 2)
            end = len(raw) if end < 0 else end + 1
            tokens.append(("code", raw[i:end]))
            i = end
            continue
        if nxt == "K":
            j = i + 2
            if j < len(raw) and raw[j] in "+-":
                j += 1
            taken = 0
            while j < len(raw) and taken < 6 and raw[j] in "0123456789ABCDEFabcdef":
                j += 1
                taken += 1
            tokens.append(("code", raw[i:j]))
            i = j
            continue
        if nxt.isdigit():
            j = i + 2
            if j < len(raw) and raw[j].isdigit():
                j += 1
            tokens.append(("code", raw[i:j]))
            i = j
            continue
        tokens.append(("code", raw[i : i + 2]))
        i += 2
    flush()
    return tokens


def _literals(raw: str, *, refuse: bool) -> list[str]:
    try:
        tokens = tokenize_header(raw)
    except _HeaderPicture:
        if refuse:
            raise SheetRefused(MSG_PICTURE_SHEET) from None
        return []
    return [text for kind, text in tokens if kind == "text"]


def _rewrite_header(raw: str, hits_by_literal: dict[int, list[LocatedHit]], label) -> str:
    try:
        tokens = tokenize_header(raw)
    except _HeaderPicture:
        raise SheetRefused(MSG_PICTURE_SHEET) from None
    seen = 0
    rewritten: list[tuple[str, str]] = []
    for kind, text in tokens:
        if kind != "text":
            rewritten.append((kind, text))
            continue
        group = hits_by_literal.get(seen, [])
        seen += 1
        rewritten.append((kind, _apply(text, group, label) if group else text))
    return "".join(piece for _, piece in rewritten)


def sanitize_sheet_title(name: str, used: set[str]) -> str:
    cleaned = (name or "").translate(_ILLEGAL_TITLE).strip()
    cleaned = cleaned[:31]
    if not cleaned:
        cleaned = "表"
    base = cleaned
    n = 2
    while cleaned.casefold() in used:
        suffix = f" {n}"
        room = 31 - len(suffix)
        cleaned = (base[:room] if room > 0 else "") + suffix
        cleaned = cleaned[:31] or "表"
        n += 1
    used.add(cleaned.casefold())
    return cleaned


def _final_title(title: str, hits: list[LocatedHit], label, used: set[str]) -> str:
    if not any(hit.enabled for hit in hits):
        used.add(title.casefold())
        return title
    return sanitize_sheet_title(_apply(title, hits, label), used)



def _zip_names(path: Path) -> list[str]:
    try:
        with zipfile.ZipFile(path) as package:
            return package.namelist()
    except zipfile.BadZipFile as exc:
        raise ValueError(MSG_OPEN_FAIL) from exc


def _refuse_xlsx(path: Path) -> None:
    try:
        head = path.read_bytes()[:8]
    except OSError as exc:
        raise ValueError(MSG_OPEN_FAIL) from exc
    if head == OLE_MAGIC:
        raise SheetRefused(MSG_ENCRYPTED_SHEET)
    names = _zip_names(path)
    if "xl/cellimages.xml" in names:
        raise SheetRefused(MSG_PICTURE_SHEET)
    with zipfile.ZipFile(path) as package:
        for name in names:
            if not name.endswith(".rels"):
                continue
            if name.startswith("xl/charts/") or name.startswith("docProps/"):
                continue
            try:
                xml = package.read(name).decode("utf-8")
            except UnicodeDecodeError:
                continue
            for tag in _REL_RE.findall(xml):
                kind = _attr(tag, "Type") or ""
                if kind.endswith("/image"):
                    raise SheetRefused(MSG_PICTURE_SHEET)
        for name in names:
            if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
                continue
            xml = package.read(name).decode("utf-8")
            for match in re.finditer(r"<headerFooter\b[^>]*>(.*?)</headerFooter>", xml, re.DOTALL):
                _literals(unescape(match.group(1)), refuse=True)


def _worksheet_paths(package: zipfile.ZipFile) -> list[str]:
    workbook = package.read("xl/workbook.xml").decode("utf-8")
    rels = package.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets: dict[str, str] = {}
    for tag in _REL_RE.findall(rels):
        ident = _attr(tag, "Id")
        target = _attr(tag, "Target")
        kind = _attr(tag, "Type") or ""
        if ident and target and kind.endswith("/worksheet"):
            cleaned = target.split("?")[0].lstrip("/")
            if not cleaned.startswith("xl/"):
                cleaned = "xl/" + cleaned
            targets[ident] = cleaned
    paths = []
    for tag in _SHEET_TAG_RE.findall(workbook):
        match = re.search(r'\br:id="([^"]+)"', tag)
        if match is None or match.group(1) not in targets:
            raise ValueError(MSG_OPEN_FAIL)
        paths.append(targets[match.group(1)])
    return paths


def _hidden_maps(xml: str) -> tuple[set[int], set[int]]:
    rows: set[int] = set()
    cols: set[int] = set()
    for tag in re.findall(r"<row\b[^>]*>", xml):
        if _attr(tag, "hidden") == "1" and _attr(tag, "r"):
            rows.add(int(_attr(tag, "r")) - 1)
    for tag in re.findall(r"<col\b[^>]*/?>", xml):
        if _attr(tag, "hidden") != "1":
            continue
        first = int(_attr(tag, "min") or "1")
        last = int(_attr(tag, "max") or str(first))
        for column in range(first - 1, last):
            cols.add(column)
    return rows, cols


def _header_slots(xml: str, *, refuse: bool) -> list[_Header]:
    found = []
    pairs = (
        ("oddHeader", "h:odd"),
        ("evenHeader", "h:even"),
        ("firstHeader", "h:first"),
        ("oddFooter", "f:odd"),
        ("evenFooter", "f:even"),
        ("firstFooter", "f:first"),
    )
    for tag, slot in pairs:
        match = re.search(rf"<{tag}\b[^>]*>(.*?)</{tag}>", xml, re.DOTALL)
        if match is None:
            continue
        raw = unescape(match.group(1))
        literals = _literals(raw, refuse=refuse)
        if raw:
            found.append(_Header(slot=slot, raw=raw, literals=literals))
    return found


def _comments_in(xml: str) -> list[_Comment]:
    authors = [unescape(text) for text in re.findall(r"<author>(.*?)</author>", xml, re.DOTALL)]
    found = []
    for block in re.findall(r"<comment\b[^>]*>.*?</comment>", xml, re.DOTALL):
        ref = _attr(block.split(">", 1)[0] + ">", "ref")
        if not ref:
            continue
        column, row = _split_ref(ref)
        author_id = int(_attr(block.split(">", 1)[0] + ">", "authorId") or "0")
        parts = [unescape(text) for text in re.findall(r"<t\b[^>]*>(.*?)</t>", block, re.DOTALL)]
        author = authors[author_id] if 0 <= author_id < len(authors) else ""
        found.append(_Comment(column, row, "".join(parts), author))
    return found


def _load_xlsx(path: Path, *, refuse: bool) -> _Book:
    if refuse:
        _refuse_xlsx(path)
    elif path.read_bytes()[:8] == OLE_MAGIC:
        raise SheetRefused(MSG_ENCRYPTED_SHEET)
    try:
        package = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise ValueError(MSG_OPEN_FAIL) from exc
    try:
        paths = _worksheet_paths(package)
        sheet_xml = [package.read(name).decode("utf-8") for name in paths]
        comments: list[list[_Comment]] = [[] for _ in paths]
        for name in package.namelist():
            if not name.endswith(".xml") or "comment" not in name.lower():
                continue
            xml = package.read(name).decode("utf-8")
            target = 0
            for index, sheet_name in enumerate(paths):
                rel_name = sheet_name.replace("xl/worksheets/", "xl/worksheets/_rels/") + ".rels"
                if rel_name in package.namelist() and name.split("/")[-1] in package.read(rel_name).decode("utf-8"):
                    target = index
                    break
            comments[target].extend(_comments_in(xml))
    finally:
        package.close()
    try:
        workbook = load_workbook(path, data_only=False)
    except Exception as exc:
        raise ValueError(MSG_OPEN_FAIL) from exc
    try:
        sheets: list[_Sheet] = []
        formulas: list[str] = []
        for index, worksheet in enumerate(workbook.worksheets):
            xml = sheet_xml[index] if index < len(sheet_xml) else ""
            hidden_rows, hidden_cols = _hidden_maps(xml)
            cells: list[_Cell] = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None and cell.data_type != "f":
                        continue
                    formula = cell.data_type == "f"
                    text = "" if cell.value is None else str(cell.value)
                    if formula:
                        formulas.append(text if text.startswith("=") else "=" + text)
                    column = cell.column - 1
                    row_index = cell.row - 1
                    if formula or isinstance(cell.value, str):
                        cells.append(
                            _Cell(
                                column,
                                row_index,
                                text,
                                formula,
                                row_index in hidden_rows or bool(worksheet.row_dimensions[cell.row].hidden),
                                column in hidden_cols or bool(worksheet.column_dimensions[cell.column_letter].hidden),
                            )
                        )
            sheets.append(
                _Sheet(
                    title=worksheet.title,
                    cells=cells,
                    comments=comments[index] if index < len(comments) else [],
                    headers=_header_slots(xml, refuse=refuse),
                    name_is_hit=True,
                )
            )
    finally:
        workbook.close()
    return _Book(sheets, formulas)


def _write_xlsx(scan: FileScan, label) -> bytes:
    path = Path(scan.path)
    _refuse_xlsx(path)
    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        files = {name: package.read(name) for name in names}
    paths = _worksheet_paths(zipfile.ZipFile(io.BytesIO(_pack_zip(files))))
    strings: list[str] = []
    sst_name = "xl/sharedStrings.xml" if "xl/sharedStrings.xml" in files else None
    if sst_name:
        strings = _sst_texts(files[sst_name].decode("utf-8"))
    appended: list[str] = []
    workbook_xml = files["xl/workbook.xml"].decode("utf-8")
    used: set[str] = set()
    sheet_tags = list(_SHEET_TAG_RE.finditer(workbook_xml))
    for index, sheet_path in enumerate(paths):
        xml = files[sheet_path].decode("utf-8")
        xml = _rewrite_sheet_xml(xml, scan, index, strings, appended, label)
        files[sheet_path] = xml.encode("utf-8")
        title = _sheet_title(workbook_xml, index)
        hits = _hits_at(scan, lambda place, i=index: isinstance(place, SheetNamePlace) and place.sheet == i)
        new_title = _final_title(title, hits, label, used)
        if new_title != title and index < len(sheet_tags):
            tag = sheet_tags[index].group(0)
            workbook_xml = workbook_xml.replace(tag, _replace_attr(tag, "name", new_title), 1)
            sheet_tags = list(_SHEET_TAG_RE.finditer(workbook_xml))
    files["xl/workbook.xml"] = workbook_xml.encode("utf-8")
    if sst_name and appended:
        files[sst_name] = _append_sst(files[sst_name].decode("utf-8"), appended).encode("utf-8")
    for name in list(files):
        if not name.endswith(".xml") or "comment" not in name.lower():
            continue
        files[name] = _rewrite_comments(files[name].decode("utf-8"), scan, label).encode("utf-8")
    return _pack_zip(files)


def _pack_zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        for name, data in files.items():
            info = zipfile.ZipInfo(filename=name)
            info.compress_type = zipfile.ZIP_DEFLATED
            package.writestr(info, data)
    return buffer.getvalue()


def _sheet_title(workbook_xml: str, index: int) -> str:
    tags = _SHEET_TAG_RE.findall(workbook_xml)
    if index >= len(tags):
        raise ValueError(MSG_OPEN_FAIL)
    return unescape(_attr(tags[index], "name") or "")


def _replace_attr(tag: str, name: str, value: str) -> str:
    return re.sub(rf'\b{name}="[^"]*"', f'{name}="{_xml_attr(value)}"', tag, count=1)


def _sst_texts(xml: str) -> list[str]:
    texts = []
    for block in re.findall(r"<si\b[^>]*>.*?</si>", xml, re.DOTALL):
        parts = [unescape(text) for text in re.findall(r"<t\b[^>]*>(.*?)</t>", block, re.DOTALL)]
        texts.append("".join(parts))
    return texts


def _append_sst(xml: str, extra: list[str]) -> str:
    pieces = []
    for text in extra:
        pieces.append(f"<si><t>{_xml_text(text)}</t></si>")
    if "</sst>" not in xml:
        raise ValueError(MSG_OPEN_FAIL)
    xml = xml.replace("</sst>", "".join(pieces) + "</sst>", 1)
    unique = len(_sst_texts(xml))
    if "uniqueCount=" in xml:
        xml = re.sub(r'uniqueCount="\d+"', f'uniqueCount="{unique}"', xml, count=1)
    return xml


def _intern(strings: list[str], appended: list[str], text: str) -> int:
    for index, existing in enumerate(strings):
        if existing == text:
            return index
    strings.append(text)
    appended.append(text)
    return len(strings) - 1


def _rewrite_sheet_xml(xml: str, scan: FileScan, index: int, strings: list[str], appended: list[str], label) -> str:
    headers = _header_slots(xml, refuse=True)

    def header_hits(slot: str, literal: int) -> list[LocatedHit]:
        return _hits_at(
            scan,
            lambda place, s=slot, n=literal: isinstance(place, HeaderFooterPlace)
            and place.sheet == index
            and place.slot == s
            and place.literal == n,
        )

    for header in headers:
        grouped = {n: header_hits(header.slot, n) for n in range(len(header.literals))}
        if not any(grouped.values()):
            continue
        updated = _rewrite_header(header.raw, grouped, label)
        tag = {"h:odd": "oddHeader", "h:even": "evenHeader", "h:first": "firstHeader", "f:odd": "oddFooter", "f:even": "evenFooter", "f:first": "firstFooter"}[header.slot]
        xml = re.sub(
            rf"(<{tag}\b[^>]*>)(.*?)(</{tag}>)",
            lambda match, text=updated: match.group(1) + _xml_text(text) + match.group(3),
            xml,
            count=1,
            flags=re.DOTALL,
        )

    def rewrite_cell(match: re.Match[str]) -> str:
        attrs = match.group(1)
        inner = match.group(2)
        if inner is None or "<f" in inner or "<f " in inner or "<f>" in inner:
            return match.group(0)
        ref = _attr("<c" + attrs + ">", "r")
        if not ref:
            return match.group(0)
        try:
            column, row = _split_ref(ref)
        except ValueError:
            return match.group(0)
        hits = _hits_at(
            scan,
            lambda place, c=column, r=row: isinstance(place, CellPlace)
            and place.sheet == index
            and place.column == c
            and place.row == r,
        )
        if not any(hit.enabled for hit in hits):
            return match.group(0)
        kind = _attr("<c" + attrs + ">", "t")
        if kind == "s":
            value = re.search(r"<v\b[^>]*>(\d+)</v>", inner)
            if value is None or not strings:
                raise ValueError(MSG_OPEN_FAIL)
            current = strings[int(value.group(1))]
            updated = _apply(current, hits, label)
            if updated == current:
                return match.group(0)
            new_index = _intern(strings, appended, updated)
            inner = re.sub(r"(<v\b[^>]*>)(.*?)(</v>)", lambda m: m.group(1) + str(new_index) + m.group(3), inner, count=1, flags=re.DOTALL)
            return f"<c{attrs}>{inner}</c>"
        if kind == "inlineStr" or "<is" in inner:
            parts = [unescape(text) for text in re.findall(r"<t\b[^>]*>(.*?)</t>", inner, re.DOTALL)]
            current = "".join(parts)
            updated = _apply(current, hits, label)
            if updated == current:
                return match.group(0)
            new_inner = re.sub(
                r"<is\b[^>]*>.*?</is>",
                f"<is><t>{_xml_text(updated)}</t></is>",
                inner,
                count=1,
                flags=re.DOTALL,
            )
            return f"<c{attrs}>{new_inner}</c>"
        return match.group(0)

    return _CELL_RE.sub(rewrite_cell, xml)


def _rewrite_comments(xml: str, scan: FileScan, label) -> str:
    if "<comment" not in xml:
        return xml
    authors = [unescape(text) for text in re.findall(r"<author>(.*?)</author>", xml, re.DOTALL)]
    extra: list[str] = []

    def one(match: re.Match[str]) -> str:
        block = match.group(0)
        head = block.split(">", 1)[0] + ">"
        ref = _attr(head, "ref")
        if not ref:
            return block
        column, row = _split_ref(ref)
        author_id = int(_attr(head, "authorId") or "0")
        text_hits = _hits_at(
            scan,
            lambda place, c=column, r=row: isinstance(place, CommentPlace)
            and place.column == c
            and place.row == r
            and place.part == "text",
        )
        author_hits = _hits_at(
            scan,
            lambda place, c=column, r=row: isinstance(place, CommentPlace)
            and place.column == c
            and place.row == r
            and place.part == "author",
        )
        parts = [unescape(text) for text in re.findall(r"<t\b[^>]*>(.*?)</t>", block, re.DOTALL)]
        current = "".join(parts)
        if any(hit.enabled for hit in text_hits):
            updated = _apply(current, text_hits, label)
            block = re.sub(
                r"<text\b[^>]*>.*?</text>",
                f"<text><t>{_xml_text(updated)}</t></text>",
                block,
                count=1,
                flags=re.DOTALL,
            )
        if any(hit.enabled for hit in author_hits) and 0 <= author_id < len(authors):
            updated_author = _apply(authors[author_id], author_hits, label)
            if updated_author != authors[author_id]:
                extra.append(updated_author)
                new_id = len(authors) + len(extra) - 1
                block = _replace_attr(block, "authorId", str(new_id)) if "authorId=" in block.split(">", 1)[0] else block
                block = re.sub(r'authorId="\d+"', f'authorId="{new_id}"', block, count=1)
        return block

    xml = re.sub(r"<comment\b[^>]*>.*?</comment>", one, xml, flags=re.DOTALL)
    if extra:
        added = "".join(f"<author>{_xml_text(text)}</author>" for text in extra)
        xml = xml.replace("</authors>", added + "</authors>", 1)
    return xml



def _csv_dialect(text: str) -> csv.Dialect:
    try:
        sniffed = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        sniffed = csv.excel
    class _Dialect(csv.excel):
        delimiter = getattr(sniffed, "delimiter", ",") or ","
        quotechar = getattr(sniffed, "quotechar", '"') or '"'
        doublequote = True
        quoting = csv.QUOTE_MINIMAL
        escapechar = None

    return _Dialect()


def _load_csv(path: Path) -> _Book:
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig")
    dialect = _csv_dialect(text)
    rows = list(csv.reader(io.StringIO(text), dialect))
    cells: list[_Cell] = []
    formulas: list[str] = []
    for row_index, row in enumerate(rows):
        for column, value in enumerate(row):
            if value.strip().startswith("="):
                formulas.append(value)
                cells.append(_Cell(column, row_index, value, True, False, False))
            elif value:
                cells.append(_Cell(column, row_index, value, False, False, False))
    return _Book([_Sheet(title=path.stem, cells=cells, comments=[], headers=[], name_is_hit=False)], formulas)


def _write_csv(scan: FileScan, label) -> bytes:
    path = Path(scan.path)
    raw = path.read_bytes()
    bom = raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    newline = "\r\n" if "\r\n" in text else "\n"
    dialect = _csv_dialect(text)
    rows = list(csv.reader(io.StringIO(text), dialect))
    delimiter = dialect.delimiter
    for row_index, row in enumerate(rows):
        for column, value in enumerate(row):
            if value.lstrip().startswith("="):
                continue
            hits = _hits_at(
                scan,
                lambda place, c=column, r=row_index: isinstance(place, CellPlace)
                and place.sheet == 0
                and place.column == c
                and place.row == r,
            )
            if any(hit.enabled for hit in hits):
                row[column] = _apply(value, hits, label)
    rendered = newline.join(_csv_line(row, delimiter) for row in rows)
    if text.endswith(("\n", "\r")):
        rendered += newline
    out = rendered.encode("utf-8")
    return (b"\xef\xbb\xbf" + out) if bom else out


def _csv_line(row: list[str], delimiter: str) -> str:
    return delimiter.join(_csv_field(value, delimiter) for value in row)


def _csv_field(value: str, delimiter: str) -> str:
    # A formula is copied as stored. Quoting it would turn ="ABC" into "=""ABC""".
    if value.lstrip().startswith("="):
        return value
    if any(char in value for char in (delimiter, '"', "\n", "\r")):
        return '"' + value.replace('"', '""') + '"'
    return value



@dataclass
class _Rec:
    typ: int
    payload: bytes
    raw: bytes | None = None
    dirty: bool = False
    skip: bool = False


def _load_xls(path: Path, *, refuse: bool) -> _Book:
    streams = _read_ole(path.read_bytes())
    if "Workbook" not in streams:
        raise ValueError(MSG_OPEN_FAIL)
    records = _parse_records(streams["Workbook"])
    _guard_biff(records, refuse=refuse)
    return _book_from_records(records, formulas_only=False)


def _write_xls(scan: FileScan, label) -> bytes:
    source = Path(scan.path).read_bytes()
    streams = _read_ole(source)
    if "Workbook" not in streams:
        raise ValueError(MSG_OPEN_FAIL)
    records = _parse_records(streams["Workbook"])
    _guard_biff(records, refuse=True)
    _redact_records(records, scan, label)
    streams["Workbook"] = _emit_records(records)
    ordered = [("Workbook", streams.pop("Workbook"))]
    for name in sorted(streams):
        ordered.append((name, streams[name]))
    return _build_ole(ordered)


def _book_from_records(records: list[_Rec], *, formulas_only: bool) -> _Book:
    groups = _groups(records)
    if not groups:
        raise ValueError(MSG_OPEN_FAIL)
    sheets: list[_Sheet] = []
    formulas: list[str] = []
    bounds = [rec for rec in groups[0] if rec.typ == BIFF_BOUNDSHEET]
    for index, group in enumerate(groups[1:]):
        title = _boundsheet_name(bounds[index].payload) if index < len(bounds) else ""
        hidden_rows, hidden_cols = _xls_hidden(group)
        cells: list[_Cell] = []
        strings = _sst_of(groups[0])
        for rec in group:
            if rec.typ == BIFF_FORMULA and len(rec.payload) >= 4:
                row, col = struct.unpack_from("<HH", rec.payload, 0)
                formulas.extend(_formula_strings(rec.payload))
                cells.append(_Cell(col, row, "", True, row in hidden_rows, col in hidden_cols))
            elif rec.typ == BIFF_LABELSST and len(rec.payload) >= 10:
                row, col, _xf, sst_index = struct.unpack_from("<HHHI", rec.payload, 0)
                text = strings[sst_index] if 0 <= sst_index < len(strings) else ""
                cells.append(_Cell(col, row, text, False, row in hidden_rows, col in hidden_cols))
            elif rec.typ == BIFF_LABEL and len(rec.payload) >= 6:
                row, col = struct.unpack_from("<HH", rec.payload, 0)
                text, _pos = _read_biff_string(rec.payload, 6, short=False)
                cells.append(_Cell(col, row, text, False, row in hidden_rows, col in hidden_cols))
        headers = []
        for rec in group:
            if rec.typ not in (BIFF_HEADER, BIFF_FOOTER) or not rec.payload:
                continue
            raw, _pos = _read_biff_string(rec.payload, 0, short=False)
            slot = "h:odd" if rec.typ == BIFF_HEADER else "f:odd"
            headers.append(_Header(slot, raw, _literals(raw, refuse=False)))
        if not formulas_only:
            sheets.append(_Sheet(title, cells, [], headers, True))
    return _Book(sheets, formulas)


def _guard_biff(records: list[_Rec], *, refuse: bool) -> None:
    if not records or records[0].typ != BIFF_BOF or len(records[0].payload) < 4:
        raise ValueError(MSG_OPEN_FAIL)
    version = struct.unpack_from("<H", records[0].payload, 0)[0]
    if version != 0x0600:
        raise ValueError(MSG_OPEN_FAIL)
    for rec in records:
        if rec.typ == BIFF_FILEPASS and refuse:
            raise SheetRefused(MSG_ENCRYPTED_SHEET)
        if rec.typ == BIFF_OBJ and _obj_is_picture(rec.payload) and refuse:
            raise SheetRefused(MSG_PICTURE_SHEET)
        if rec.typ in (BIFF_HEADER, BIFF_FOOTER) and rec.payload and refuse:
            raw, _pos = _read_biff_string(rec.payload, 0, short=False)
            _literals(raw, refuse=True)


def _redact_records(records: list[_Rec], scan: FileScan, label) -> None:
    groups = _groups(records)
    strings = _sst_of(groups[0])
    uses: list[tuple[_Rec, int, int, int, int]] = []
    for index, group in enumerate(groups[1:]):
        for rec in group:
            if rec.typ == BIFF_LABELSST and len(rec.payload) >= 10:
                row, col, _xf, sst_index = struct.unpack_from("<HHHI", rec.payload, 0)
                uses.append((rec, index, row, col, sst_index))
            elif rec.typ == BIFF_LABEL and len(rec.payload) >= 6:
                row, col = struct.unpack_from("<HH", rec.payload, 0)
                text, _pos = _read_biff_string(rec.payload, 6, short=False)
                hits = _cell_hit_list(scan, index, col, row)
                if any(hit.enabled for hit in hits):
                    updated = _apply(text, hits, label)
                    rec.payload = rec.payload[:6] + _encode_biff_string(updated, short=False)
                    rec.dirty = True
                    rec.raw = None
            elif rec.typ in (BIFF_HEADER, BIFF_FOOTER) and rec.payload:
                raw, _pos = _read_biff_string(rec.payload, 0, short=False)
                slot = "h:odd" if rec.typ == BIFF_HEADER else "f:odd"
                literals = _literals(raw, refuse=True)
                grouped = {}
                for literal_index in range(len(literals)):
                    grouped[literal_index] = _hits_at(
                        scan,
                        lambda place, i=index, s=slot, n=literal_index: isinstance(place, HeaderFooterPlace)
                        and place.sheet == i
                        and place.slot == s
                        and place.literal == n,
                    )
                if any(grouped.values()):
                    updated = _rewrite_header(raw, grouped, label)
                    rec.payload = _encode_biff_string(updated, short=False)
                    rec.dirty = True
                    rec.raw = None
    original = list(strings)
    plans: list[tuple[int, int, int, int, str | None]] = []
    for _rec, index, row, col, sst_index in uses:
        text = original[sst_index] if 0 <= sst_index < len(original) else ""
        hits = _cell_hit_list(scan, index, col, row)
        updated = _apply(text, hits, label) if any(hit.enabled for hit in hits) else None
        plans.append((index, row, col, sst_index, updated))
    by_index: dict[int, list[tuple[int, int, int, str | None]]] = {}
    for index, row, col, sst_index, updated in plans:
        by_index.setdefault(sst_index, []).append((index, row, col, updated))
    assigned: dict[tuple[int, int, int], int] = {}
    for sst_index, group in by_index.items():
        current = original[sst_index] if 0 <= sst_index < len(original) else ""
        targets = {updated for _i, _r, _c, updated in group if updated is not None}
        keeps_original = any(updated is None or updated == current for _i, _r, _c, updated in group)
        if targets == {current} or not targets:
            for index, row, col, _updated in group:
                assigned[(index, row, col)] = sst_index
            continue
        if not keeps_original and len(targets) == 1:
            strings[sst_index] = next(iter(targets))
            for index, row, col, _updated in group:
                assigned[(index, row, col)] = sst_index
            continue
        for index, row, col, updated in group:
            if updated is None or updated == current:
                assigned[(index, row, col)] = sst_index
                continue
            if updated in strings:
                assigned[(index, row, col)] = strings.index(updated)
                continue
            strings.append(updated)
            assigned[(index, row, col)] = len(strings) - 1
    for rec, index, row, col, sst_index in uses:
        new_index = assigned.get((index, row, col), sst_index)
        if new_index != sst_index:
            row_v, col_v, xf, _old = struct.unpack_from("<HHHI", rec.payload, 0)
            rec.payload = struct.pack("<HHHI", row_v, col_v, xf, new_index)
            rec.dirty = True
            rec.raw = None
    _write_sst(groups[0], strings, len(uses))
    _rename_sheets(groups, scan, label)
    _retarget_positions(groups)


def _cell_hit_list(scan: FileScan, sheet: int, column: int, row: int) -> list[LocatedHit]:
    return _hits_at(
        scan,
        lambda place, s=sheet, c=column, r=row: isinstance(place, CellPlace)
        and place.sheet == s
        and place.column == c
        and place.row == r,
    )


def _rename_sheets(groups: list[list[_Rec]], scan: FileScan, label) -> None:
    bounds = [rec for rec in groups[0] if rec.typ == BIFF_BOUNDSHEET]
    used: set[str] = set()
    for index, rec in enumerate(bounds):
        pos, vis, kind, name = _parse_boundsheet(rec.payload)
        hits = _hits_at(scan, lambda place, i=index: isinstance(place, SheetNamePlace) and place.sheet == i)
        title = _final_title(name, hits, label, used)
        rec.payload = _encode_boundsheet(pos, vis, kind, title)
        rec.dirty = True
        rec.raw = None


def _retarget_positions(groups: list[list[_Rec]]) -> None:
    global_len = sum(len(_emit_one(rec)) for rec in groups[0])
    offset = global_len
    sheet_starts = []
    sheet_layouts = []
    for group in groups[1:]:
        sheet_starts.append(offset)
        starts = []
        cursor = 0
        for rec in group:
            starts.append(cursor)
            cursor += len(_emit_one(rec))
        sheet_layouts.append(starts)
        offset += cursor
    bounds = [rec for rec in groups[0] if rec.typ == BIFF_BOUNDSHEET]
    for index, rec in enumerate(bounds):
        if index >= len(sheet_starts):
            break
        pos, vis, kind, name = _parse_boundsheet(rec.payload)
        rec.payload = _encode_boundsheet(sheet_starts[index], vis, kind, name)
        rec.dirty = True
        rec.raw = None
    for group, base, starts in zip(groups[1:], sheet_starts, sheet_layouts):
        dbcells = [pos for pos, rec in zip(starts, group) if rec.typ == BIFF_DBCELL]
        for rec in group:
            if rec.typ != BIFF_INDEX or len(rec.payload) < 12:
                continue
            slots = (len(rec.payload) - 12) // 4
            if slots != len(dbcells) or (len(rec.payload) - 12) % 4:
                continue
            packed = rec.payload[:12]
            for rel in dbcells:
                packed += struct.pack("<I", base + rel)
            rec.payload = packed
            rec.dirty = True
            rec.raw = None


def _write_sst(globals_: list[_Rec], strings: list[str], total: int) -> None:
    payload = struct.pack("<II", total, len(strings))
    for text in strings:
        payload += _encode_biff_string(text, short=False)
    written = False
    for rec in globals_:
        if rec.typ == BIFF_SST and not written:
            rec.payload = payload
            rec.dirty = True
            rec.raw = None
            written = True
        elif rec.typ == BIFF_EXTSST and written:
            rec.skip = True


def _sst_of(globals_: list[_Rec]) -> list[str]:
    for rec in globals_:
        if rec.typ == BIFF_SST:
            _total, strings = _parse_sst(rec.payload)
            return strings
    return []


def _parse_sst(payload: bytes) -> tuple[int, list[str]]:
    if len(payload) < 8:
        return 0, []
    total, unique = struct.unpack_from("<II", payload, 0)
    pos = 8
    strings = []
    for _ in range(unique):
        text, pos = _read_biff_string(payload, pos, short=False)
        strings.append(text)
    if pos != len(payload):
        raise ValueError(MSG_OPEN_FAIL)
    return total, strings


def _xls_hidden(group: list[_Rec]) -> tuple[set[int], set[int]]:
    rows: set[int] = set()
    cols: set[int] = set()
    for rec in group:
        if rec.typ == BIFF_ROW and len(rec.payload) >= 16:
            row = struct.unpack_from("<H", rec.payload, 0)[0]
            options = struct.unpack_from("<I", rec.payload, 12)[0]
            if options & 0x20:
                rows.add(row)
        elif rec.typ == BIFF_COLINFO and len(rec.payload) >= 10:
            first, last, _width, _xf, options = struct.unpack_from("<5H", rec.payload, 0)
            if options & 0x0001:
                cols.update(range(first, last + 1))
    return rows, cols


def _formula_strings(payload: bytes) -> list[str]:
    found = []
    for run in re.findall(rb"[\x20-\x7e]{4,}", payload):
        found.append(run.decode("ascii"))
    if len(payload) >= 4:
        decoded = payload.decode("utf-16-le", errors="ignore")
        for run in re.findall(r"[\x20-\x7e]{4,}", decoded):
            if run not in found:
                found.append(run)
    return found


def _obj_is_picture(payload: bytes) -> bool:
    if len(payload) < 6:
        return False
    ft, cb = struct.unpack_from("<HH", payload, 0)
    if ft != 0x0015:
        return False
    ot = struct.unpack_from("<H", payload, 4)[0]
    return ot == 0x0008 and cb >= 2


def _groups(records: list[_Rec]) -> list[list[_Rec]]:
    groups: list[list[_Rec]] = []
    bucket: list[_Rec] = []
    depth = 0
    for rec in records:
        if rec.typ == BIFF_BOF and depth == 0 and bucket:
            groups.append(bucket)
            bucket = []
        bucket.append(rec)
        if rec.typ == BIFF_BOF:
            depth += 1
        elif rec.typ == BIFF_EOF:
            depth -= 1
            if depth == 0:
                groups.append(bucket)
                bucket = []
    if bucket:
        groups.append(bucket)
    return groups


def _parse_records(data: bytes) -> list[_Rec]:
    records: list[_Rec] = []
    pos = 0
    limit = len(data)
    while pos + 4 <= limit:
        typ, length = struct.unpack_from("<HH", data, pos)
        if typ == 0 and length == 0:
            break
        if pos + 4 + length > limit:
            raise ValueError(MSG_OPEN_FAIL)
        payload = data[pos + 4 : pos + 4 + length]
        raw = data[pos : pos + 4 + length]
        pos += 4 + length
        if typ == BIFF_CONTINUE and records:
            records[-1].payload += payload
            records[-1].raw = (records[-1].raw or b"") + raw
            continue
        records.append(_Rec(typ, payload, raw))
    if not records:
        raise ValueError(MSG_OPEN_FAIL)
    return records


def _emit_records(records: list[_Rec]) -> bytes:
    return b"".join(_emit_one(rec) for group in _groups(records) for rec in group)


def _emit_one(rec: _Rec) -> bytes:
    if rec.skip:
        return b""
    if rec.raw is not None and not rec.dirty:
        return rec.raw
    payload = rec.payload
    if len(payload) <= BIFF_MAX:
        return struct.pack("<HH", rec.typ, len(payload)) + payload
    out = struct.pack("<HH", rec.typ, BIFF_MAX) + payload[:BIFF_MAX]
    rest = payload[BIFF_MAX:]
    while rest:
        chunk = rest[:BIFF_MAX]
        rest = rest[BIFF_MAX:]
        out += struct.pack("<HH", BIFF_CONTINUE, len(chunk)) + chunk
    return out


def _read_biff_string(buf: bytes, offset: int, *, short: bool) -> tuple[str, int]:
    if short:
        if offset + 2 > len(buf):
            raise ValueError(MSG_OPEN_FAIL)
        count = buf[offset]
        flags = buf[offset + 1]
        pos = offset + 2
    else:
        if offset + 3 > len(buf):
            raise ValueError(MSG_OPEN_FAIL)
        count, flags = struct.unpack_from("<HB", buf, offset)
        pos = offset + 3
    rich = ext = 0
    if flags & 0x08:
        rich = struct.unpack_from("<H", buf, pos)[0]
        pos += 2
    if flags & 0x04:
        ext = struct.unpack_from("<I", buf, pos)[0]
        pos += 4
    width = 2 if flags & 0x01 else 1
    raw = buf[pos : pos + count * width]
    if len(raw) < count * width:
        raise ValueError(MSG_OPEN_FAIL)
    pos += count * width
    text = raw.decode("utf-16-le" if flags & 0x01 else "latin-1")
    pos += rich * 4 + ext
    return text, pos


def _encode_biff_string(text: str, *, short: bool) -> bytes:
    try:
        raw = text.encode("latin-1")
        flags = 0
    except UnicodeEncodeError:
        raw = text.encode("utf-16-le")
        flags = 1
    if short:
        if len(text) > 255:
            text = text[:255]
            raw = text.encode("latin-1") if flags == 0 else text.encode("utf-16-le")
        return bytes((len(text), flags)) + raw
    return struct.pack("<HB", len(text), flags) + raw


def _parse_boundsheet(payload: bytes) -> tuple[int, int, int, str]:
    if len(payload) < 8:
        raise ValueError(MSG_OPEN_FAIL)
    pos, vis, kind = struct.unpack_from("<LBB", payload, 0)
    name, _end = _read_biff_string(payload, 6, short=True)
    return pos, vis, kind, name


def _boundsheet_name(payload: bytes) -> str:
    return _parse_boundsheet(payload)[3]


def _encode_boundsheet(pos: int, vis: int, kind: int, name: str) -> bytes:
    return struct.pack("<LBB", pos, vis, kind) + _encode_biff_string(name, short=True)



def _read_ole(data: bytes) -> dict[str, bytes]:
    if len(data) < 512 or data[:8] != OLE_MAGIC:
        raise ValueError(MSG_OPEN_FAIL)
    sector_shift = struct.unpack_from("<H", data, 0x1E)[0]
    mini_shift = struct.unpack_from("<H", data, 0x20)[0]
    if sector_shift != 9 or mini_shift != 6:
        raise ValueError(MSG_OPEN_FAIL)
    sector = 512
    mini = 64
    fat_count = struct.unpack_from("<I", data, 0x2C)[0]
    dir_start = struct.unpack_from("<i", data, 0x30)[0]
    cutoff = struct.unpack_from("<I", data, 0x38)[0] or 4096
    mini_fat_start = struct.unpack_from("<i", data, 0x3C)[0]
    mini_fat_count = struct.unpack_from("<I", data, 0x40)[0]
    difat_start = struct.unpack_from("<i", data, 0x44)[0]
    difat_count = struct.unpack_from("<I", data, 0x48)[0]
    difat = list(struct.unpack_from("<109i", data, 0x4C))
    sid = difat_start
    seen = set()
    while difat_count and sid >= 0:
        if sid in seen or (sid + 1) * sector + sector > len(data):
            break
        seen.add(sid)
        block = data[(sid + 1) * sector : (sid + 2) * sector]
        nums = list(struct.unpack("<128i", block))
        difat.extend(nums[:-1])
        sid = nums[-1]
        difat_count -= 1
    fat: list[int] = []
    for fat_sid in difat[:fat_count]:
        if fat_sid < 0:
            continue
        block = data[(fat_sid + 1) * sector : (fat_sid + 2) * sector]
        fat.extend(struct.unpack("<128i", block))

    def chain(start: int, size: int) -> bytes:
        if start < 0 or size <= 0:
            return b""
        out = bytearray()
        sid_now = start
        guard = set()
        while sid_now >= 0 and len(out) < size + sector:
            if sid_now in guard or sid_now >= len(fat):
                break
            guard.add(sid_now)
            out += data[(sid_now + 1) * sector : (sid_now + 2) * sector]
            sid_now = fat[sid_now]
        return bytes(out[:size])

    directory = chain(dir_start, 1_048_576)
    entries = []
    for offset in range(0, len(directory), 128):
        chunk = directory[offset : offset + 128]
        if len(chunk) < 128:
            break
        name_len, typ = struct.unpack_from("<HB", chunk, 64)
        if typ == 0 or name_len < 2:
            if typ == 0:
                continue
        name = chunk[: max(0, name_len - 2)].decode("utf-16-le", errors="ignore")
        start, size = struct.unpack_from("<iQ", chunk, 116)
        entries.append((name, typ, start, size))
    root = next((entry for entry in entries if entry[1] == 5), None)
    mini_stream = b""
    mini_fat: list[int] = []
    if root is not None and mini_fat_start >= 0 and mini_fat_count:
        mini_stream = chain(root[2], root[3])
        raw_fat = chain(mini_fat_start, mini_fat_count * sector)
        mini_fat = list(struct.unpack(f"<{len(raw_fat) // 4}i", raw_fat)) if raw_fat else []

    def mini_chain(start: int, size: int) -> bytes:
        if start < 0 or size <= 0:
            return b""
        out = bytearray()
        sid_now = start
        guard = set()
        while sid_now >= 0 and len(out) < size + mini:
            if sid_now in guard or sid_now >= len(mini_fat):
                break
            guard.add(sid_now)
            off = sid_now * mini
            out += mini_stream[off : off + mini]
            sid_now = mini_fat[sid_now]
        return bytes(out[:size])

    found: dict[str, bytes] = {}
    saw_book = False
    for name, typ, start, size in entries:
        if typ != 2 or not name:
            continue
        if name == "Book":
            saw_book = True
        blob = chain(start, size) if size >= cutoff or typ == 5 else mini_chain(start, size)
        found[name] = blob
    if "Workbook" not in found and saw_book:
        raise ValueError(MSG_OPEN_FAIL)
    if "Workbook" not in found:
        raise ValueError(MSG_OPEN_FAIL)
    return found


def _build_ole(named: list[tuple[str, bytes]]) -> bytes:
    sector = 512
    workbook = b""
    small: list[tuple[str, bytes]] = []
    large: list[tuple[str, bytes]] = []
    for name, blob in named:
        if name == "Workbook":
            rem = len(blob) % 4096
            workbook = blob + b"\x00" * (4096 - rem)
        elif len(blob) < 4096:
            small.append((name, blob))
        else:
            large.append((name, blob))
    if not workbook:
        raise ValueError(MSG_OPEN_FAIL)

    def sectors_of(blob: bytes) -> list[bytes]:
        if not blob:
            return []
        padded = blob + b"\x00" * ((-len(blob)) % sector)
        return [padded[i : i + sector] for i in range(0, len(padded), sector)]

    mini_fat: list[int] = []
    mini_stream = bytearray()
    mini_place: dict[str, tuple[int, int]] = {}
    for name, blob in small:
        if not blob:
            mini_place[name] = (-2, 0)
            continue
        start = len(mini_fat)
        padded = blob + b"\x00" * ((-len(blob)) % 64)
        mini_stream += padded
        count = len(padded) // 64
        for step in range(count - 1):
            mini_fat.append(start + step + 1)
        mini_fat.append(-2)
        mini_place[name] = (start, len(blob))
    if mini_fat:
        slots = ((len(mini_fat) + 127) // 128) * 128
        packed = mini_fat + [-1] * (slots - len(mini_fat))
        minifat_bytes = struct.pack(f"<{len(packed)}i", *packed)
    else:
        minifat_bytes = b""

    stream_meta: list[tuple[str, int, int]] = []
    pieces: list[bytes] = []
    pieces.extend(sectors_of(workbook))
    stream_meta.append(("Workbook", 0, len(workbook)))
    cursor = len(pieces)
    for name, blob in large:
        start = cursor
        parts = sectors_of(blob)
        pieces.extend(parts)
        stream_meta.append((name, start, len(blob)))
        cursor = len(pieces)
    mini_sector_start = -2
    if mini_stream:
        mini_sector_start = cursor
        pieces.extend(sectors_of(bytes(mini_stream)))
        cursor = len(pieces)
    minifat_start = -2
    minifat_count = 0
    if minifat_bytes:
        minifat_start = cursor
        parts = sectors_of(minifat_bytes)
        minifat_count = len(parts)
        pieces.extend(parts)
        cursor = len(pieces)

    nodes = [("Workbook", 0, len(workbook))]
    for name, blob in large:
        meta = next(item for item in stream_meta if item[0] == name)
        nodes.append((name, meta[1], meta[2]))
    for name, blob in small:
        start, size = mini_place[name]
        nodes.append((name, start, size))
    order = sorted(range(len(nodes)), key=lambda i: nodes[i][0].upper())
    left = [-1] * len(nodes)
    right = [-1] * len(nodes)

    def balance(idxs: list[int]) -> int:
        if not idxs:
            return -1
        mid = len(idxs) // 2
        sid = idxs[mid]
        left[sid] = balance(idxs[:mid])
        right[sid] = balance(idxs[mid + 1 :])
        return sid

    root_child = balance(order)
    dir_slots = ((1 + len(nodes) + 3) // 4) * 4
    directory = bytearray()

    def _sid(node_index: int) -> int:
        return node_index + 1 if node_index >= 0 else -1

    directory += _dir_entry("Root Entry", 5, 1, -1, -1, _sid(root_child), mini_sector_start, len(mini_stream))
    for index, (name, start, size) in enumerate(nodes):
        directory += _dir_entry(name, 2, 1, _sid(left[index]), _sid(right[index]), -1, start, size)
    while len(directory) // 128 < dir_slots:
        directory += _dir_entry("", 0, 1, -1, -1, -1, -2, 0)
    dir_start = cursor
    dir_parts = sectors_of(bytes(directory))
    pieces.extend(dir_parts)
    data_count = len(pieces)
    fat_count = 1
    while (data_count + fat_count + 127) // 128 > fat_count:
        fat_count += 1
    if fat_count > 109:
        raise ValueError(MSG_OPEN_FAIL)
    fat = [-1] * (fat_count * 128)

    def link(start: int, count: int) -> None:
        for step in range(count - 1):
            fat[start + step] = start + step + 1
        if count:
            fat[start + count - 1] = -2

    link(0, len(sectors_of(workbook)))
    running = len(sectors_of(workbook))
    for _name, blob in large:
        count = len(sectors_of(blob))
        link(running, count)
        running += count
    if mini_stream:
        count = len(sectors_of(bytes(mini_stream)))
        link(running, count)
        running += count
    if minifat_bytes:
        link(running, minifat_count)
        running += minifat_count
    link(dir_start, len(dir_parts))
    fat_start = data_count
    for step in range(fat_count):
        fat[fat_start + step] = -3
    fat_bytes = struct.pack(f"<{len(fat)}i", *fat)
    pieces.extend(sectors_of(fat_bytes))
    header = bytearray(512)
    header[0:8] = OLE_MAGIC
    struct.pack_into("<H", header, 0x18, 0x003E)
    struct.pack_into("<H", header, 0x1A, 0x0003)
    struct.pack_into("<H", header, 0x1C, 0xFFFE)
    struct.pack_into("<H", header, 0x1E, 9)
    struct.pack_into("<H", header, 0x20, 6)
    struct.pack_into("<I", header, 0x2C, fat_count)
    struct.pack_into("<i", header, 0x30, dir_start)
    struct.pack_into("<I", header, 0x38, 4096)
    struct.pack_into("<i", header, 0x3C, minifat_start)
    struct.pack_into("<I", header, 0x40, minifat_count)
    struct.pack_into("<i", header, 0x44, -2)
    for step in range(fat_count):
        struct.pack_into("<i", header, 0x4C + 4 * step, fat_start + step)
    for step in range(fat_count, 109):
        struct.pack_into("<i", header, 0x4C + 4 * step, -1)
    return bytes(header) + b"".join(pieces)


def _dir_entry(name: str, typ: int, color: int, left: int, right: int, child: int, start: int, size: int) -> bytes:
    raw = (name + "\x00").encode("utf-16-le") if name else b""
    field = raw[:64].ljust(64, b"\x00")
    return struct.pack(
        "<64sHBB iii 16s I QQ i Q".replace(" ", ""),
        field,
        len(raw),
        typ,
        color,
        left,
        right,
        child,
        b"\x00" * 16,
        0,
        0,
        0,
        start,
        size,
    )
