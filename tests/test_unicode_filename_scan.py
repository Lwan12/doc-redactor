import io
import tempfile
import zipfile
from pathlib import Path

import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont
from docx import Document

from doc_redactor.config import load_rules
from doc_redactor.scan import scan_path, write_scans
from doc_redactor.tesseract_cmd import configure_tesseract
from doc_redactor.verify_output import count_leftovers
from helpers import RULES, load_make_samples, needs_tesseract

_CASES = [
    ("english", "CAPA回复1", False),
    ("中文目录", "report", False),
    ("中文 目录", "CAPA回复1", False),
    ("english", "report", True),
    ("中文 目录", "CAPA回复1", True),
]


def _engine():
    return load_rules(RULES)


def _use_temp(monkeypatch, path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)
    monkeypatch.setenv("TMPDIR", str(path))
    monkeypatch.setenv("TEMP", str(path))
    monkeypatch.setenv("TMP", str(path))
    monkeypatch.setattr(tempfile, "tempdir", None)


def _folded(text: str) -> str:
    return "".join(text.split()).lower()


def _ocr(image: Image.Image) -> str:
    configure_tesseract()
    import pytesseract

    return pytesseract.image_to_string(image, lang="chi_sim+eng")


def _assert_readable(text: str) -> None:
    folded = _folded(text)
    assert "abc-2026-001" in folded
    assert "acmepharma" in folded
    assert "fictizumab" in folded


def _assert_unreadable(text: str) -> None:
    folded = _folded(text)
    assert "abc-2026-001" not in folded
    assert "acmepharma" not in folded
    assert "fictizumab" not in folded


def _card(path: Path) -> None:
    font = ImageFont.truetype(str(load_make_samples().font_path()), 64)
    image = Image.new("RGB", (1400, 360), "white")
    draw = ImageDraw.Draw(image)
    y = 24
    for line in ("Project ABC-2026-001", "Sponsor Acme Pharma", "Drug Fictizumab"):
        draw.text((24, y), line, fill="black", font=font)
        y += 110
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG", dpi=(300, 300))


def _prepare(tmp_path: Path, monkeypatch, folder: str, chinese_temp: bool) -> Path:
    if chinese_temp:
        _use_temp(monkeypatch, tmp_path / "临时 目录")
    return tmp_path / folder


def _finish(scan, out: Path, engine, stem: str, suffix: str):
    assert scan.blocked_reason is None, scan.blocked_reason
    result = write_scans([scan], out, engine=engine)[0]
    assert result.error is None
    assert result.dest == out / f"{stem}.redacted{suffix}"
    assert result.dest.is_file()
    labels = [engine.label(category) for category in result.counts]
    assert count_leftovers(result.dest, engine=engine, known=[hit.text for hit in scan.hits], labels=labels) == 0
    return result.dest


@needs_tesseract
@pytest.mark.parametrize("folder,stem,chinese_temp", _CASES)
def test_docx_scan_write_and_recheck(tmp_path, monkeypatch, folder, stem, chinese_temp):
    card = tmp_path / "card.png"
    _card(card)
    _assert_readable(_ocr(Image.open(card)))
    directory = _prepare(tmp_path, monkeypatch, folder, chinese_temp)
    directory.mkdir()
    source = directory / f"{stem}.docx"
    document = Document()
    document.add_paragraph("Project ABC-2026-001 Sponsor Acme Pharma Drug Fictizumab")
    document.add_picture(str(card))
    document.save(source)
    assert source.name == f"{stem}.docx"

    engine = _engine()
    scan = scan_path(source, engine)
    assert any(hit.unit.startswith("media:") for hit in scan.hits)
    dest = _finish(scan, tmp_path / "out" / folder, engine, stem, ".docx")
    with zipfile.ZipFile(dest) as archive:
        name = next(item for item in archive.namelist() if item.lower().startswith("word/media/"))
        payload = archive.read(name)
    _assert_unreadable(_ocr(Image.open(io.BytesIO(payload))))
    assert source.is_file()


@needs_tesseract
@pytest.mark.parametrize("folder,stem,chinese_temp", _CASES)
def test_pdf_scan_write_and_recheck(tmp_path, monkeypatch, folder, stem, chinese_temp):
    card = tmp_path / "card.png"
    _card(card)
    directory = _prepare(tmp_path, monkeypatch, folder, chinese_temp)
    directory.mkdir()
    source = directory / f"{stem}.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Project ABC-2026-001 Sponsor Acme Pharma Drug Fictizumab")
    page.insert_image(pymupdf.Rect(72, 140, 560, 320), filename=str(card))
    document.save(source)
    document.close()

    engine = _engine()
    scan = scan_path(source, engine)
    assert any(hit.unit.startswith("pdfimg:") for hit in scan.hits)
    dest = _finish(scan, tmp_path / "out" / folder, engine, stem, ".pdf")
    opened = pymupdf.open(dest)
    text = "\n".join(item.get_text() for item in opened)
    _assert_unreadable(text)
    pixmap = pymupdf.Pixmap(opened, opened[0].get_images()[0][0])
    _assert_unreadable(_ocr(Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)))
    opened.close()


@needs_tesseract
@pytest.mark.parametrize("folder,stem,chinese_temp", _CASES)
def test_png_scan_write_and_recheck(tmp_path, monkeypatch, folder, stem, chinese_temp):
    directory = _prepare(tmp_path, monkeypatch, folder, chinese_temp)
    source = directory / f"{stem}.png"
    _card(source)
    engine = _engine()
    scan = scan_path(source, engine)
    assert any(hit.unit == "image" for hit in scan.hits)
    dest = _finish(scan, tmp_path / "out" / folder, engine, stem, ".png")
    _assert_unreadable(_ocr(Image.open(dest)))
