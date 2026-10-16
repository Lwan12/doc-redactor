import zipfile
from pathlib import Path

from docx import Document

from doc_redactor.engine import RegexRule, WordlistRule, build_engine
from doc_redactor.scan import (
    FILE_HIT_WARN,
    scan_path,
    too_many_file_hits,
    write_report,
    write_scan,
)
from doc_redactor.pattern_guess import TEST_HIT_WARN, too_many_test_hits


def _engine(pattern: str = r"ABC-\d{4}-\d{3}", name: str = "PROJECT_ID"):
    return build_engine(
        [RegexRule(name=name, pattern=pattern, flags=0)],
        [],
    )


def _docx(path: Path, lines: list[str]) -> None:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(path)


def test_file_hit_threshold_500_and_501(tmp_path: Path):
    assert too_many_file_hits(500) is False
    assert too_many_file_hits(501) is True
    assert FILE_HIT_WARN == 500
    assert too_many_test_hits(20) is False
    assert too_many_test_hits(21) is True
    assert TEST_HIT_WARN == 20

    engine = _engine()
    token = "ABC-2026-001"
    five_hundred = tmp_path / "n500.docx"
    _docx(five_hundred, [" ".join([token] * 500)])
    scan500 = scan_path(five_hundred, engine)
    assert len(scan500.hits) == 500
    assert scan500.too_many is False

    five_oh_one = tmp_path / "n501.docx"
    _docx(five_oh_one, [" ".join([token] * 501)])
    scan501 = scan_path(five_oh_one, engine)
    assert len(scan501.hits) == 501
    assert scan501.too_many is True


def test_uncheck_then_write_counts_match(tmp_path: Path):
    source = tmp_path / "notes.docx"
    _docx(
        source,
        [
            "Project ABC-2026-001 first",
            "Project ABC-2026-001 second",
            "Project ABC-2026-001 third",
        ],
    )
    engine = _engine()
    scan = scan_path(source, engine)
    assert len(scan.hits) == 3
    scan.hits[1].enabled = False
    dest = tmp_path / "notes.redacted.docx"
    counts = write_scan(scan, dest, replacements={"PROJECT_ID": "[REDACTED:PROJECT_ID]"})
    assert counts == {"PROJECT_ID": 2}
    report = tmp_path / "notes.report.json"
    write_report(report, counts)
    raw = report.read_text(encoding="utf-8")
    assert "ABC-2026-001" not in raw
    assert "start" not in raw
    assert "location" not in raw

    written = Document(str(dest))
    texts = [paragraph.text for paragraph in written.paragraphs]
    assert texts[0] == "Project [REDACTED:PROJECT_ID] first"
    assert texts[1] == "Project ABC-2026-001 second"
    assert texts[2] == "Project [REDACTED:PROJECT_ID] third"
    with zipfile.ZipFile(dest) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert xml.count("ABC-2026-001") == 1


def test_report_has_no_source_text(tmp_path: Path):
    source = tmp_path / "one.docx"
    _docx(source, ["Acme Pharma and ABC-2026-001"])
    engine = build_engine(
        [RegexRule(name="PROJECT_ID", pattern=r"ABC-\d{4}-\d{3}", flags=0)],
        [WordlistRule(category="SPONSOR", terms=["Acme Pharma"], case_sensitive=False)],
    )
    scan = scan_path(source, engine)
    dest = tmp_path / "one.redacted.docx"
    counts = write_scan(scan, dest)
    report = tmp_path / "one.report.json"
    write_report(report, counts)
    payload = report.read_text(encoding="utf-8")
    assert "Acme Pharma" not in payload
    assert "ABC-2026-001" not in payload
    assert counts["PROJECT_ID"] == 1
    assert counts["SPONSOR"] == 1
