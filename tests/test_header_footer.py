"""One-click header/footer removal: identify first, apply only confirmed candidates.

All fixtures are made up (ABC-2026-001, Acme Pharma, Fictizumab).
PDF routes: Pagination artifact tags, then repetition across pages, else nothing.
"""

from __future__ import annotations

import io
import json
import zipfile

import pymupdf
import pytest
from docx import Document
from docx.oxml import parse_xml
from PIL import Image

from doc_redactor import header_footer as hf

W, H = 595.0, 842.0
HEADER = "ABC-2026-001 Version 1.0"
SPONSOR_FOOTER = "Acme Pharma CONFIDENTIAL"
SIGN = "Docusign Envelope ID: 0000AAAA-FAKE-0000"
BODY_TOP = "Fictizumab body first line"
BODY_MID = "Fictizumab dosing schedule"
# Body lines differ by words, not only by digits (digits are ignored when comparing).
WORDS = ("alpha", "bravo", "charlie", "delta", "echo", "foxtrot", "golf", "hotel", "india", "juliet", "kilo", "lima")


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def _wrap_last(doc, page, prefix: bytes) -> None:
    xref = page.get_contents()[-1]
    doc.update_stream(xref, prefix + b"\n" + doc.xref_stream(xref) + b"\nEMC\n")


HDR_TAG = b"/Artifact <</Attached [/Top]/Type/Pagination/Subtype/Header>> BDC"
FTR_TAG = b"/Artifact <</Attached [/Bottom]/Type/Pagination/Subtype/Footer>> BDC"


def _tagged_pdf(pages: int = 4) -> bytes:
    doc = pymupdf.open()
    logo = _png()
    for i in range(pages):
        page = doc.new_page(width=W, height=H)
        page.insert_text((72, 55), HEADER, fontsize=10)
        _wrap_last(doc, page, HDR_TAG)
        page.insert_image(pymupdf.Rect(480, 40, 520, 60), stream=logo)
        _wrap_last(doc, page, HDR_TAG)
        page.draw_line((72, 82.7), (523, 82.7), width=0.5)
        _wrap_last(doc, page, HDR_TAG)
        page.insert_text((72, 760), SPONSOR_FOOTER, fontsize=9)
        _wrap_last(doc, page, FTR_TAG)
        page.insert_text((280, 785), f"Page {i + 1}/{pages}", fontsize=9)
        _wrap_last(doc, page, FTR_TAG)
        # Untagged signing line added after export, top of every page.
        page.insert_text((72, 15), SIGN, fontsize=7)
        # Body hugs the header rule.
        page.insert_text((72, 92), f"{BODY_TOP} {WORDS[i]}", fontsize=10)
        page.insert_text((72, 400), f"{BODY_MID} {i}", fontsize=10)
        # A plain untyped artifact in the body (e.g. table rule) must stay.
        page.insert_text((72, 420), f"Layout note {i}", fontsize=10)
        _wrap_last(doc, page, b"/Artifact BMC")
    doc.set_xml_metadata("")
    doc.xref_set_key(doc.pdf_catalog(), "MarkInfo", "<</Marked true>>")
    return doc.tobytes()


def _repeat_pdf(pages: int = 12) -> bytes:
    doc = pymupdf.open()
    logo = _png()
    for i in range(pages):
        landscape = i == 5
        width, height = (H, W) if landscape else (W, H)
        page = doc.new_page(width=width, height=height)
        page.insert_text((72, 55), HEADER, fontsize=10)
        page.insert_image(pymupdf.Rect(width - 115, 40, width - 75, 60), stream=logo)
        page.draw_line((72, 70), (width - 72, 70), width=0.5)
        label = f"Page {i + 1}/{pages}"
        text_w = pymupdf.get_text_length(label, fontsize=9)
        page.insert_text(((width - text_w) / 2, height - 40), label, fontsize=9)
        page.insert_text((72, height - 60), SPONSOR_FOOTER, fontsize=9)
        page.insert_text((72, 80), f"{BODY_TOP} {WORDS[i]}", fontsize=10)
        page.insert_text((72, height / 2), "Repeated body sentence in the middle", fontsize=10)
        if i == 3:
            page.insert_text((72, 30), "One-off top note", fontsize=9)
    return doc.tobytes()


def _plain_pdf(pages: int = 4) -> bytes:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page(width=W, height=H)
        page.insert_text((72, 60 + 37 * i), f"Fictizumab paragraph {i}", fontsize=10)
        page.insert_text((72, 400), f"Body text number {i} for Acme Pharma", fontsize=10)
    return doc.tobytes()


def _open(data: bytes):
    return pymupdf.open("pdf", data)


def _all_text(doc) -> str:
    return "\n".join(page.get_text() for page in doc)


def _raw_streams(doc) -> bytes:
    return b"".join(doc.xref_stream(x) for page in doc for x in page.get_contents())


def _secrets_in(blob: str) -> list[str]:
    return [s for s in (HEADER, SPONSOR_FOOTER, SIGN, BODY_TOP, BODY_MID, "ABC-2026-001", "Acme", "Fictizumab", "Docusign") if s in blob]


# ---------------------------------------------------------------- pure layer


def _layout(index, width, height, items):
    return hf.PageLayout(index=index, width=width, height=height, items=tuple(items))


def _t(text, box):
    return hf.Element(kind=hf.KIND_TEXT, bbox=box, text=text)


def test_find_repeats_tolerates_page_number_width_and_landscape():
    pages = []
    for i in range(10):
        width = 842.0 if i == 4 else 595.0
        label = f"Page {i + 1}/10"
        half = 3.0 if i < 9 else 4.0  # wider when the number has two digits
        cx = width / 2
        pages.append(
            _layout(
                i,
                width,
                842.0,
                [
                    _t(HEADER, (72, 45, 200, 57)),
                    _t(label, (cx - 10 * half, 800, cx + 10 * half, 810)),
                    _t(f"body {i}", (72, 300, 200, 312)),
                ],
            )
        )
    found = hf.find_repeats(pages)
    roles = sorted((c.role, len(c.pages)) for c in found)
    assert roles == [("footer", 10), ("header", 10)]
    assert all(c.needs_confirmation and not c.confirmed for c in found)
    assert all(c.source == hf.SOURCE_REPEAT for c in found)


def test_find_repeats_needs_three_pages_and_more_than_half():
    two = [_layout(i, W, H, [_t(HEADER, (72, 45, 200, 57))]) for i in range(2)]
    assert hf.find_repeats(two) == []
    four = [_layout(i, W, H, [_t(HEADER, (72, 45, 200, 57))] if i < 2 else []) for i in range(4)]
    assert hf.find_repeats(four) == []  # exactly half is not more than half
    five = [_layout(i, W, H, [_t(HEADER, (72, 45, 200, 57))] if i < 3 else []) for i in range(5)]
    assert len(hf.find_repeats(five)) == 1


def test_find_repeats_ignores_middle_of_page():
    pages = [_layout(i, W, H, [_t("same in the middle", (72, 400, 200, 412))]) for i in range(5)]
    assert hf.find_repeats(pages) == []


def test_signing_line_gets_its_own_role():
    pages = [_layout(i, W, H, [_t(SIGN, (72, 8, 260, 16))]) for i in range(4)]
    (only,) = hf.find_repeats(pages)
    assert only.role == hf.ROLE_SIGNATURE


def test_confirm_marks_every_candidate_and_is_pure():
    pages = [_layout(i, W, H, [_t(HEADER, (72, 45, 200, 57))]) for i in range(4)]
    plan = hf.Plan(kind=hf.KIND_PDF, route=hf.ROUTE_REPEAT, page_count=4, candidates=tuple(hf.find_repeats(pages)))
    done = hf.confirm(plan)
    assert all(c.confirmed for c in done.candidates)
    assert not any(c.confirmed for c in plan.candidates)


def test_text_is_kept_out_of_repr_summary_and_preview():
    pages = [_layout(i, W, H, [_t(HEADER, (72, 45, 200, 57))]) for i in range(4)]
    plan = hf.Plan(kind=hf.KIND_PDF, route=hf.ROUTE_REPEAT, page_count=4, candidates=tuple(hf.find_repeats(pages)))
    dumped = repr(plan) + json.dumps(hf.summary(plan), ensure_ascii=False) + json.dumps(hf.preview_rows(plan), ensure_ascii=False)
    assert _secrets_in(dumped) == []


def test_ui_strings_follow_flow_doc_and_avoid_the_banned_word():
    for text in hf.UI_TEXT.values():
        assert "正则" not in text
    assert hf.UI_TEXT["not_found"] == "ⓘ 这个 PDF 没认出页眉页脚，没有删除。页眉页脚里的词仍按规则替换。"
    assert hf.label_for(hf.ROLE_HEADER) == "将删除：页眉"
    assert hf.label_for(hf.ROLE_FOOTER) == "将删除：页脚"
    assert hf.label_for(hf.ROLE_SIGNATURE) == "签署信息"


# ---------------------------------------------------------------- PDF: tags


def test_tagged_pdf_uses_tags_and_flags_untagged_signing_line():
    doc = _open(_tagged_pdf())
    plan = hf.identify_pdf(doc)
    assert plan.route == hf.ROUTE_TAG
    tags = [c for c in plan.candidates if c.source == hf.SOURCE_TAG]
    extra = [c for c in plan.candidates if c.source == hf.SOURCE_REPEAT]
    assert {c.role for c in tags} == {hf.ROLE_HEADER, hf.ROLE_FOOTER}
    assert all(c.confirmed and not c.needs_confirmation for c in tags)
    covered = {(p, c.role) for c in tags for p in c.pages}
    assert covered == {(p, r) for p in range(4) for r in (hf.ROLE_HEADER, hf.ROLE_FOOTER)}
    header_kinds = {e.kind for c in tags if c.role == hf.ROLE_HEADER for e in c.elements}
    assert header_kinds == {hf.KIND_TEXT, hf.KIND_IMAGE, hf.KIND_LINE}
    assert [c.role for c in extra] == [hf.ROLE_SIGNATURE]
    assert extra[0].needs_confirmation and not extra[0].confirmed
    texts = " ".join(e.text for c in plan.candidates for e in c.elements)
    assert "Layout note" not in texts and BODY_TOP not in texts


def test_apply_refuses_unconfirmed_candidates():
    doc = _open(_tagged_pdf())
    plan = hf.identify_pdf(doc)
    with pytest.raises(hf.UnconfirmedCandidates):
        hf.apply_pdf(doc, plan)
    # Nothing was touched.
    assert SIGN in _all_text(doc) and HEADER in _all_text(doc)


def test_tagged_apply_really_removes_and_keeps_body():
    doc = _open(_tagged_pdf())
    before_body = [line for line in _all_text(doc).splitlines() if "body" in line or "Layout" in line or "dosing" in line]
    plan = hf.confirm(hf.identify_pdf(doc))
    result = hf.apply_pdf(doc, plan)
    out = _open(doc.tobytes(garbage=3, deflate=True))
    text = _all_text(out)
    for gone in (HEADER, SPONSOR_FOOTER, SIGN, "Page 1/4"):
        assert gone not in text
    raw = _raw_streams(out)
    assert b"ABC-2026-001" not in raw and b"Acme Pharma" not in raw and b"Docusign" not in raw
    assert all(not page.get_images() for page in out)
    assert all(not page.get_drawings() for page in out)
    after_body = [line for line in text.splitlines() if "body" in line or "Layout" in line or "dosing" in line]
    assert after_body == before_body and len(after_body) == 12
    assert result.pages == 4
    assert result.body_lost == 0 and result.left == 0
    assert result.removed[hf.KIND_TEXT] == 4 * 4  # header, sponsor line, page number, signing line


def test_plan_from_another_file_is_rejected():
    plan = hf.confirm(hf.identify_pdf(_open(_tagged_pdf())))
    other = _open(_plain_pdf())
    with pytest.raises(hf.PlanMismatch):
        hf.apply_pdf(other, plan)


# ---------------------------------------------------------------- PDF: repeats


def test_untagged_pdf_falls_back_to_repeats_and_needs_confirmation():
    doc = _open(_repeat_pdf())
    plan = hf.identify_pdf(doc)
    assert plan.route == hf.ROUTE_REPEAT
    assert plan.candidates and all(c.needs_confirmation and not c.confirmed for c in plan.candidates)
    kinds = {e.kind for c in plan.candidates for e in c.elements}
    assert kinds == {hf.KIND_TEXT, hf.KIND_IMAGE, hf.KIND_LINE}
    footer_pages = {p for c in plan.candidates if c.role == hf.ROLE_FOOTER for p in c.pages}
    assert footer_pages == set(range(12))
    texts = " ".join(e.text for c in plan.candidates for e in c.elements)
    assert "One-off" not in texts and "middle" not in texts and BODY_TOP not in texts
    with pytest.raises(hf.UnconfirmedCandidates):
        hf.apply_pdf(doc, plan)


def test_confirmed_repeats_are_removed_and_body_stays():
    doc = _open(_repeat_pdf())
    plan = hf.confirm(hf.identify_pdf(doc))
    result = hf.apply_pdf(doc, plan)
    out = _open(doc.tobytes(garbage=3, deflate=True))
    text = _all_text(out)
    assert HEADER not in text and SPONSOR_FOOTER not in text and "Page " not in text
    assert text.count("Repeated body sentence in the middle") == 12
    assert text.count(BODY_TOP) == 12 and "One-off top note" in text
    assert b"ABC-2026-001" not in _raw_streams(out)
    assert all(not page.get_images() for page in out)
    assert all(not page.get_drawings() for page in out)
    assert result.pages == 12 and result.body_lost == 0 and result.left == 0


# ---------------------------------------------------------------- PDF: nothing


def test_plain_pdf_is_left_alone():
    data = _plain_pdf()
    doc = _open(data)
    plan = hf.identify_pdf(doc)
    assert plan.route == hf.ROUTE_NONE and plan.candidates == ()
    assert hf.summary(plan)["message"] == hf.UI_TEXT["not_found"]
    result = hf.apply_pdf(doc, plan)
    assert result.pages == 0
    assert _all_text(_open(doc.tobytes())) == _all_text(_open(data))


def test_summary_counts_pages_and_flags_auto_found():
    plan = hf.identify_pdf(_open(_repeat_pdf()))
    info = hf.summary(plan)
    assert info["route"] == hf.ROUTE_REPEAT
    assert info["pages"] == 12
    assert info["needs_confirmation"] is True
    assert info["line"] == "页眉页脚　12 页 · 将删除"
    assert info["note"] == "ⓘ 自动认出，请看预览确认"
    rows = hf.preview_rows(plan)
    assert rows and set(rows[0]) == {"id", "page", "role", "label", "source", "bbox", "needs_confirmation", "confirmed", "items"}


# ---------------------------------------------------------------- Word

VML_WATERMARK = (
    '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict>'
    '<v:shape id="PowerPlusWaterMarkObject1" type="#_x0000_t136">'
    '<v:textpath string="Acme Pharma CONFIDENTIAL"/></v:shape></w:pict></w:r></w:p>'
)
TEXTBOX = (
    '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict><v:shape><v:textbox><w:txbxContent>'
    '<w:p><w:r><w:t>ABC-2026-001 in a box</w:t></w:r></w:p>'
    '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'
)
FIELD = (
    '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:fldSimple w:instr=" TITLE "><w:r><w:t>Fictizumab title field</w:t></w:r></w:fldSimple></w:p>'
)


def _docx(tmp_path, *, watermark: bool = True):
    document = Document()
    document.add_paragraph("Body mentions Fictizumab once.")
    section = document.sections[0]
    section.header.paragraphs[0].text = HEADER
    section.header._element.append(parse_xml(TEXTBOX))
    if watermark:
        section.header._element.append(parse_xml(VML_WATERMARK))
    section.footer.paragraphs[0].text = SPONSOR_FOOTER
    section.footer._element.append(parse_xml(FIELD))
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "First page ABC-2026-001"
    buf = io.BytesIO()
    Image.new("RGB", (30, 10), (0, 0, 255)).save(buf, format="PNG")
    buf.seek(0)
    section.header.paragraphs[0].add_run().add_picture(buf)
    path = tmp_path / "in.docx"
    document.save(path)
    return path


def _part_xml(path, prefix):
    with zipfile.ZipFile(path) as zf:
        return {n: zf.read(n).decode("utf-8") for n in zf.namelist() if n.startswith(prefix) and n.endswith(".xml")}


def test_word_identifies_each_owned_header_footer_and_watermark(tmp_path):
    document = Document(_docx(tmp_path))
    plan = hf.identify_docx(document)
    assert plan.route == hf.ROUTE_WORD and plan.has_watermark is True
    roles = sorted(c.role for c in plan.candidates)
    assert roles == [hf.ROLE_FOOTER, hf.ROLE_HEADER, hf.ROLE_HEADER]
    assert all(c.confirmed and not c.needs_confirmation for c in plan.candidates)
    assert hf.summary(plan)["watermark_note"] == "这个文件的水印在页眉里，会一起删除。"
    assert _secrets_in(repr(plan) + json.dumps(hf.summary(plan), ensure_ascii=False)) == []


def test_word_apply_clears_textbox_field_watermark_and_images(tmp_path):
    src = _docx(tmp_path)
    media_before = [n for n in zipfile.ZipFile(src).namelist() if n.startswith("word/media/")]
    assert media_before
    document = Document(src)
    plan = hf.identify_docx(document)
    result = hf.apply_docx(document, plan)
    dest = tmp_path / "out.docx"
    document.save(dest)
    assert hf.verify_docx_output(dest) == 0
    xml = "".join(_part_xml(dest, "word/header").values()) + "".join(_part_xml(dest, "word/footer").values())
    for gone in ("ABC-2026-001", "Acme Pharma", "Fictizumab", "PowerPlusWaterMarkObject", "fldSimple", "txbxContent"):
        assert gone not in xml
    assert not [n for n in zipfile.ZipFile(dest).namelist() if n.startswith("word/media/")]
    assert [p.text for p in Document(dest).paragraphs] == ["Body mentions Fictizumab once."]
    assert result.parts == 3


def test_word_without_headers_has_nothing_to_do(tmp_path):
    document = Document()
    document.add_paragraph("Only body.")
    path = tmp_path / "plain.docx"
    document.save(path)
    plan = hf.identify_docx(Document(path))
    assert plan.route == hf.ROUTE_NONE and plan.candidates == ()
