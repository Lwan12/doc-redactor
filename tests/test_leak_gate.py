"""出厂门禁、整份无命中、示例规则里的 \\b。全部是编造数据。"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from doc_redactor import pdf_redactor
from doc_redactor.cli import main, redact_paths
from doc_redactor.config import load_rules
from doc_redactor.engine import RegexRule, WordlistRule, build_engine
from doc_redactor.leak_check import MSG_LEAKED_HEAD, MSG_NO_HITS, count_leaks
from doc_redactor.scan import scan_path, write_scans
from helpers import RULES

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _engine(pattern=r"(?<![A-Za-z0-9])ABC-\d{4}-\d{3}(?![A-Za-z0-9])"):
    return build_engine(
        [RegexRule(name="PROJECT_ID", pattern=pattern, flags=0)],
        [
            WordlistRule(category="SPONSOR", terms=["Acme Pharma", "星辰海"], case_sensitive=False),
            WordlistRule(category="DRUG", terms=["Fictizumab"], case_sensitive=True),
        ],
    )


# --- 第 8 条：\b 在中文旁边 ----------------------------------------------------


def test_example_rule_matches_id_next_to_chinese():
    engine = load_rules(RULES)
    hits = engine.find_hits("受试者ABC-2026-001入组，方案ABC-2026-001号。XABC-2026-001 不算。")
    assert [hit.text for hit in hits] == ["ABC-2026-001", "ABC-2026-001"]


def test_word_boundary_in_user_pattern_works_next_to_chinese():
    engine = _engine(r"\bABC-\d{4}-\d{3}\b")
    assert [h.text for h in engine.find_hits("受试者ABC-2026-001入组")] == ["ABC-2026-001"]
    assert engine.find_hits("XABC-2026-001") == []
    assert engine.find_hits("ABC-2026-0012") == []
    # 中文两边都不是 ASCII 字，和词表一样不设边界。
    cjk = _engine(r"\b星辰海\b")
    assert [h.text for h in cjk.find_hits("公司星辰海医药")] == ["星辰海"]


def test_escaped_backslash_and_class_backspace_are_left_alone():
    literal = _engine(r"ABC\\b")
    assert [h.text for h in literal.find_hits("x ABC\\b y")] == ["ABC\\b"]
    in_class = _engine(r"ABC[\b]")
    assert [h.text for h in in_class.find_hits("ABC\b")] == ["ABC\b"]


# --- 整份没有命中 --------------------------------------------------------------


def _plain_docx(path: Path, text: str = "Nothing sensitive here.") -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(str(path))


def test_cli_does_not_copy_a_file_without_hits(tmp_path: Path, capsys):
    source = tmp_path / "clean.docx"
    _plain_docx(source)
    pdf = tmp_path / "clean.pdf"
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "Nothing here")
    document.save(pdf)
    document.close()
    out = tmp_path / "out"
    summaries, blocked = redact_paths([source, pdf], RULES, out)
    assert summaries == []
    assert len(blocked) == 2 and all(MSG_NO_HITS in line for line in blocked)
    assert list(out.glob("*.redacted.*")) == []
    assert list(out.glob("*.report.json")) == []
    code = main([str(source), "-r", str(RULES), "-o", str(tmp_path / "out2")])
    assert code == 1
    assert MSG_NO_HITS in capsys.readouterr().err


def test_gui_write_marks_a_file_without_hits(tmp_path: Path):
    source = tmp_path / "clean.docx"
    _plain_docx(source)
    scan = scan_path(source, load_rules(RULES))
    assert scan.status == "nohit"
    result = write_scans([scan], tmp_path / "out", engine=load_rules(RULES))[0]
    assert result.dest is None
    assert result.note == MSG_NO_HITS
    assert not (tmp_path / "out" / "clean.redacted.docx").exists()


# --- 出厂门禁 -------------------------------------------------------------------


def _numbering_leak_docx(path: Path) -> None:
    """正文有命中；编号格式（numbering.xml 的 lvlText）里还有原词，工具不处理这一处。"""
    _plain_docx(path, "Sponsor Acme Pharma.")
    numbering = (
        f'<w:numbering xmlns:w="{W}"><w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0">'
        '<w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="Fictizumab %1."/>'
        "</w:lvl></w:abstractNum><w:num w:numId=\"1\"><w:abstractNumId w:val=\"0\"/></w:num></w:numbering>"
    )
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    files["word/numbering.xml"] = numbering.encode()
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    path.write_bytes(buffer.getvalue())


def test_gate_blocks_docx_when_another_reader_still_sees_a_word(tmp_path: Path):
    source = tmp_path / "numbered.docx"
    _numbering_leak_docx(source)
    engine = _engine()
    out = tmp_path / "out"
    # 命令行
    from doc_redactor.cli import redact_path
    from doc_redactor.leak_check import LeakedText

    with pytest.raises(LeakedText) as caught:
        redact_path(source, out, engine)
    assert caught.value.count == 1
    assert "Fictizumab" not in str(caught.value)
    assert not (out / "numbered.redacted.docx").exists()
    assert not (out / "numbered.report.json").exists()
    # 界面
    scan = scan_path(source, engine)
    result = write_scans([scan], tmp_path / "gui", engine=engine)[0]
    assert result.dest is None
    assert result.error.startswith(MSG_LEAKED_HEAD.format(n=1))
    assert list((tmp_path / "gui").glob("*.docx")) == []


def test_gate_blocks_pdf_when_cleanup_is_skipped(tmp_path: Path, monkeypatch):
    source = tmp_path / "meta.pdf"
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "Sponsor Acme Pharma")
    document.set_metadata({"title": "Fictizumab report"})
    document.save(source)
    document.close()
    monkeypatch.setattr(pdf_redactor, "clear_pdf_hidden", lambda document: None)
    import doc_redactor.scan as scan_module

    monkeypatch.setattr(scan_module, "clear_pdf_hidden", lambda document: None)
    engine = _engine()
    scan = scan_path(source, engine)
    result = write_scans([scan], tmp_path / "gui", engine=engine)[0]
    assert result.dest is None
    assert MSG_LEAKED_HEAD.format(n=1) in result.error
    assert list((tmp_path / "gui").glob("*.pdf")) == []


def test_unchecked_hit_does_not_trip_the_gate(tmp_path: Path):
    source = tmp_path / "two.docx"
    _plain_docx(source, "Sponsor Acme Pharma, drug Fictizumab.")
    engine = _engine()
    scan = scan_path(source, engine)
    for hit in scan.hits:
        if hit.category == "DRUG":
            hit.enabled = False
    result = write_scans([scan], tmp_path / "gui", engine=engine)[0]
    assert result.dest is not None, result.error
    assert count_leaks(result.dest, engine, kept=["Fictizumab"]) == 0
    assert count_leaks(result.dest, engine) == 1


def test_placeholders_and_legend_page_do_not_trip_the_gate(tmp_path: Path):
    source = tmp_path / "narrow.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Acme Pharma", fontsize=4)
    page.insert_text((72, 120), "Fictizumab given")
    document.save(source)
    document.close()
    engine = _engine()
    scan = scan_path(source, engine)
    result = write_scans([scan], tmp_path / "gui", engine=engine)[0]
    assert result.dest is not None, result.error
