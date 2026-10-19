"""Only the replacement text changed: keep unchecked hits; scope changed: clear."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document

from doc_redactor.backpath import (
    allocate_batch_dir,
    apply_unchecks,
    capture_unchecks,
    session_signature,
)
from doc_redactor.engine import RegexRule, WordlistRule, build_engine
from doc_redactor.ruleset import KIND_WORDLIST, Rule
from doc_redactor.scan import scan_path, write_scans


def _doc(path: Path) -> None:
    document = Document()
    document.add_paragraph("申办方星辰海医药，合作方星辰海研究所。")
    document.save(path)


def _rules(replacement: str = "", terms=("星辰海",), enabled: bool = True) -> list[Rule]:
    return [
        Rule(name="申办方", kind=KIND_WORDLIST, terms=list(terms), enabled=enabled, replacement=replacement)
    ]


def _engine(rules: list[Rule]):
    lists = [WordlistRule(category=r.name, terms=list(r.terms), case_sensitive=False) for r in rules if r.enabled]
    return build_engine([RegexRule(name="NONE_X", pattern=r"a^", flags=0)], lists)


def _scan_once(path: Path, rules):
    scan = scan_path(path, _engine(rules))
    assert len(scan.hits) == 2
    return scan


def test_only_replacement_changed_keeps_unchecked_and_uses_new_text(tmp_path: Path):
    source = tmp_path / "memo.docx"
    _doc(source)
    before = _rules(replacement="[甲方]")
    scan = _scan_once(source, before)
    scan.hits[1].enabled = False
    memory = capture_unchecks([scan], session_signature([source], before))

    after = _rules(replacement="[SPONSOR-X]")
    signature = session_signature([source], after)
    assert signature == memory.signature

    fresh = _scan_once(source, after)
    assert apply_unchecks([fresh], memory, signature) == "kept"
    assert fresh.hits[0].enabled is True
    assert fresh.hits[1].enabled is False

    batch = allocate_batch_dir(tmp_path / "out", datetime(2026, 10, 19, 17, 30))
    result = write_scans([fresh], batch, replacements={"申办方": "[SPONSOR-X]"})
    body = "\n".join(p.text for p in Document(str(result[0].dest)).paragraphs)
    assert "[SPONSOR-X]" in body
    assert "[甲方]" not in body
    assert body.count("[SPONSOR-X]") == 1
    assert "星辰海研究所" in body  # the unchecked hit stays as is


def test_scope_changes_still_clear_unchecked(tmp_path: Path):
    source = tmp_path / "memo.docx"
    _doc(source)
    base = _rules()
    scan = _scan_once(source, base)
    scan.hits[1].enabled = False
    memory = capture_unchecks([scan], session_signature([source], base))

    for changed in (
        _rules(terms=("星辰海", "星海")),
        _rules(enabled=False) + [Rule(name="药名", kind=KIND_WORDLIST, terms=["星辰海妥单抗"])],
        [Rule(name="申办方", kind="pattern", pattern=r"星辰海")],
    ):
        assert session_signature([source], changed) != memory.signature
        fresh = _scan_once(source, base)
        assert apply_unchecks([fresh], memory, session_signature([source], changed)) == "cleared"
        assert all(hit.enabled for hit in fresh.hits)
