"""Word text inside text boxes, simple fields, content controls, and revisions.

All data is fictional (Acme Pharma, Fictizumab, ABC-2026-001). Assertions
read the XML parts of the written file directly.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from lxml import etree

from doc_redactor.docx_redactor import redact_docx
from doc_redactor.engine import RegexRule, WordlistRule, build_engine
from doc_redactor.scan import scan_path, write_scan

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = (
    f'xmlns:w="{W}" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)
LABELS = {"PROJECT_ID": "[PID]", "SPONSOR": "[SPONSOR]", "DRUG": "[DRUG]"}
SECRETS = ("Acme Pharma", "Fictizumab", "ABC-2026-001", "ABC-2026-002", "ABC-2026-003")
PART = re.compile(r"word/(document|header\d*|footer\d*)\.xml$")


def _engine():
    return build_engine(
        [RegexRule(name="PROJECT_ID", pattern=r"ABC-\d{4}-\d{3}", flags=0)],
        [
            WordlistRule(category="SPONSOR", terms=["Acme Pharma"], case_sensitive=True),
            WordlistRule(category="DRUG", terms=["Fictizumab"], case_sensitive=True),
        ],
    )


def _x(fragment: str):
    """Parse one element; the namespaces go on the first tag."""
    head, sep, rest = fragment.partition(">")
    if head.endswith("/"):
        return parse_xml(f"{head[:-1]} {NS}/>{rest}")
    return parse_xml(f"{head} {NS}>{rest}")


def _r(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _txbx(text: str) -> str:
    return f"<w:txbxContent><w:p>{_r(text)}</w:p></w:txbxContent>"


def _textbox_run(text: str):
    """A modern text box: wps copy in mc:Choice, VML copy in mc:Fallback."""
    return _x(
        "<w:r><mc:AlternateContent>"
        '<mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic>'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f"<wps:wsp><wps:txbx>{_txbx(text)}</wps:txbx></wps:wsp>"
        "</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><v:shape><v:textbox>{_txbx(text)}</v:textbox></v:shape></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r>"
    )


def _vml_box_run(text: str):
    return _x(f"<w:r><w:pict><v:shape><v:textbox>{_txbx(text)}</v:textbox></v:shape></w:pict></w:r>")


def _fld(text: str):
    return _x(f'<w:fldSimple w:instr=" DOCPROPERTY Company ">{_r(text)}</w:fldSimple>')


def _sdt_inline(text: str):
    return _x(f"<w:sdt><w:sdtPr/><w:sdtContent>{_r(text)}</w:sdtContent></w:sdt>")


def _sdt_block(text: str):
    return _x(f"<w:sdt><w:sdtPr/><w:sdtContent><w:p>{_r(text)}</w:p></w:sdtContent></w:sdt>")


def _ins(text: str):
    return _x(f'<w:ins w:id="1" w:author="Tester" w:date="2026-01-01T00:00:00Z">{_r(text)}</w:ins>')


def _del(text: str):
    return _x(
        '<w:del w:id="2" w:author="Tester" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:delText xml:space="preserve">{text}</w:delText></w:r></w:del>'
    )


def _hyperlink(text: str):
    return _x(f'<w:hyperlink w:anchor="top">{_r(text)}</w:hyperlink>')


def _plain(text: str):
    return _x(_r(text))


def _block_before_sectpr(body, element) -> None:
    sect = body.find(f"{{{W}}}sectPr")
    if sect is None:
        body.append(element)
    else:
        sect.addprevious(element)


# Every hit the fixture holds, by category. The text box counts once even
# though Word stores two copies of it.
EXPECTED = {"SPONSOR": 5, "DRUG": 7, "PROJECT_ID": 6}
EXPECTED_DELETED = {"SPONSOR": 1, "DRUG": 1, "PROJECT_ID": 2}


def _build(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]

    # Header: text box, simple field, inline and block content control, insertion.
    hp = section.header.paragraphs[0]._p
    hp.append(_plain("Head "))
    hp.append(_textbox_run("Box Acme Pharma"))
    hp.append(_fld("Fictizumab"))
    hp.append(_plain(" / "))
    hp.append(_sdt_inline("ABC-2026-001"))
    section.header._element.append(_sdt_block("Header block Fictizumab"))
    p2 = section.header.add_paragraph()._p
    p2.append(_plain("Rev "))
    p2.append(_ins("Acme Pharma"))

    # Footer: block content control and an insertion that completes a term.
    section.footer._element.append(_sdt_block("Foot ABC-2026-002"))
    fp = section.footer.add_paragraph()._p
    fp.append(_plain("Acme "))
    fp.append(_ins("Pharma"))
    # Tracked deletions in the header and footer too.
    hp.append(_del(" Fictizumab"))
    fp.append(_del(" ABC-2026-002"))

    # Body.
    p = doc.add_paragraph("Field: ")._p
    p.append(_fld("Acme Pharma"))
    p = doc.add_paragraph("Control: ")._p
    p.append(_sdt_inline("Fictizumab"))
    _block_before_sectpr(doc.element.body, _sdt_block("Block ABC-2026-003"))
    p = doc.add_paragraph("Inserted ")._p
    p.append(_ins("ABC-2026-001"))
    p = doc.add_paragraph("Old ")._p
    p.append(_del("Acme Pharma"))
    p.append(_plain(" new text"))
    p = doc.add_paragraph("VML ")._p
    p.append(_vml_box_run("Old box Fictizumab"))
    # Hyperlink: scanned text includes it, so write offsets must too.
    p = doc.add_paragraph("See ")._p
    p.append(_hyperlink("Fictizumab"))
    p.append(_plain(" and ABC-2026-002"))

    # Table: a row and a cell wrapped in content controls.
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "keep a"
    table.cell(0, 1).text = "keep b"
    # Tracked insertion and deletion inside a table cell.
    cell_p = table.cell(0, 1).add_paragraph("Cell rev ")._p
    cell_p.append(_ins("Fictizumab"))
    cell_p.append(_del("ABC-2026-003"))
    tbl = table._tbl
    tbl.append(
        _x(
            "<w:sdt><w:sdtPr/><w:sdtContent><w:tr>"
            f"<w:tc><w:p>{_r('Row Acme Pharma')}</w:p></w:tc>"
            f"<w:tc><w:p>{_r('keep c')}</w:p></w:tc>"
            "</w:tr></w:sdtContent></w:sdt>"
        )
    )
    tbl.append(
        _x(
            "<w:tr>"
            f"<w:tc><w:p>{_r('keep d')}</w:p></w:tc>"
            f"<w:sdt><w:sdtPr/><w:sdtContent><w:tc><w:p>{_r('Cell Fictizumab')}</w:p></w:tc></w:sdtContent></w:sdt>"
            "</w:tr>"
        )
    )
    doc.add_paragraph("Body ABC-2026-003")
    doc.save(str(path))


def _parts(path: Path) -> dict[str, etree._Element]:
    with zipfile.ZipFile(path) as archive:
        return {name: etree.fromstring(archive.read(name)) for name in archive.namelist() if PART.match(name)}


def _texts(path: Path, tag: str) -> str:
    chunks = []
    for root in _parts(path).values():
        chunks.append("".join(el.text or "" for el in root.iter(f"{{{W}}}{tag}")))
    return "\n".join(chunks)


def _assert_clean(path: Path) -> None:
    live = _texts(path, "t")
    deleted = _texts(path, "delText")
    for secret in SECRETS + ("Acme ", "Pharma"):
        assert secret not in live, secret
        assert secret not in deleted, secret
    # Each text box copy got its own placeholder.
    assert live.count("Box [SPONSOR]") == 2
    assert "See [DRUG] and [PID]" in live
    assert "Row [SPONSOR]" in live and "Cell [DRUG]" in live
    # Each deletion stays a deletion, with the placeholder inside.
    for label in ("[SPONSOR]", "[DRUG]", "[PID]"):
        assert label in deleted
    assert "Cell rev [DRUG]" in live
    assert "keep a" in live and "keep d" in live and "new text" in live


def _total() -> dict[str, int]:
    out = dict(EXPECTED)
    for key, value in EXPECTED_DELETED.items():
        out[key] = out.get(key, 0) + value
    return out


def test_source_fixture_holds_terms(tmp_path: Path):
    src = tmp_path / "hidden.docx"
    _build(src)
    assert "Box Acme Pharma" in _texts(src, "t")
    assert "Acme Pharma" in _texts(src, "delText")


def test_cli_redacts_hidden_containers(tmp_path: Path):
    src = tmp_path / "hidden.docx"
    _build(src)
    engine = _engine()
    base_label = engine.label
    engine.label = lambda cat: LABELS.get(cat) or base_label(cat)
    dest = tmp_path / "out.docx"
    counts = redact_docx(src, dest, engine)
    _assert_clean(dest)
    assert counts == _total()


def test_scan_then_write_hidden_containers(tmp_path: Path):
    src = tmp_path / "hidden.docx"
    _build(src)
    scan = scan_path(src, _engine())
    counts: dict[str, int] = {}
    for hit in scan.hits:
        counts[hit.category] = counts.get(hit.category, 0) + 1
    assert counts == _total()
    deleted = [hit for hit in scan.hits if "修订删除" in hit.location]
    assert sorted(hit.category for hit in deleted) == ["DRUG", "PROJECT_ID", "PROJECT_ID", "SPONSOR"]
    dest = tmp_path / "out.docx"
    write_scan(scan, dest, replacements=LABELS)
    _assert_clean(dest)


def test_unchecked_text_box_hit_keeps_both_copies(tmp_path: Path):
    src = tmp_path / "hidden.docx"
    _build(src)
    scan = scan_path(src, _engine())
    box = [hit for hit in scan.hits if hit.text == "Acme Pharma"]
    target = None
    for hit in box:
        if scan.paragraphs[int(hit.unit.split(":")[1])] == "Box Acme Pharma":
            target = hit
    assert target is not None
    target.enabled = False
    dest = tmp_path / "out.docx"
    write_scan(scan, dest, replacements=LABELS)
    live = _texts(dest, "t")
    # Both stored copies follow the one decision.
    assert live.count("Box Acme Pharma") == 2
    assert "Box [SPONSOR]" not in live


def test_drawing_run_survives_a_hit_around_it(tmp_path: Path):
    src = tmp_path / "around.docx"
    doc = Document()
    p = doc.add_paragraph()._p
    p.append(_plain("Acme "))
    p.append(_vml_box_run("inside keep"))
    p.append(_plain("Pharma tail"))
    doc.save(str(src))
    engine = _engine()
    dest = tmp_path / "out.docx"
    redact_docx(src, dest, engine)
    live = _texts(dest, "t")
    assert "inside keep" in live
    assert "Acme" not in live and "Pharma" not in live


def test_shared_walk_reads_revisions_and_finds_nothing_after_write(tmp_path: Path):
    """iter_document_paragraphs is the one walk; an after-write check can reuse it."""
    from doc_redactor.docx_redactor import iter_document_paragraphs

    src = tmp_path / "hidden.docx"
    _build(src)
    engine = _engine()
    before = list(iter_document_paragraphs(Document(str(src))))
    assert sum(len(unit.find_hits(engine)) for unit in before) == sum(_total().values())
    dest = tmp_path / "out.docx"
    write_scan(scan_path(src, engine), dest, replacements=LABELS)
    after = list(iter_document_paragraphs(Document(str(dest))))
    assert sum(len(unit.find_hits(engine)) for unit in after) == 0
    # Raw text of every unit (deleted ones included) has no term left either.
    assert sum(len(engine.find_hits(unit.text)) for unit in after) == 0


def _straddle(path: Path) -> None:
    doc = Document()
    # Before changes this read "Sponsor Acme Pharma"; now "Sponsor Acme Corp".
    p = doc.add_paragraph("Sponsor Acme ")._p
    p.append(_del("Pharma"))
    p.append(_ins("Corp"))
    # A kept hit next to a deletion counts once, in the text unit.
    p = doc.add_paragraph("Kept Fictizumab ")._p
    p.append(_del("old words"))
    doc.save(str(path))


def test_term_split_between_kept_and_deleted_text(tmp_path: Path):
    src = tmp_path / "straddle.docx"
    _straddle(src)
    scan = scan_path(src, _engine())
    pairs = sorted((hit.category, hit.location) for hit in scan.hits)
    assert pairs == [("DRUG", "第 2 段"), ("SPONSOR", "第 1 段 · 修订删除")]
    for dest, run in (
        (tmp_path / "gui.docx", lambda d: write_scan(scan, d, replacements=LABELS)),
        (tmp_path / "cli.docx", None),
    ):
        if run is None:
            engine = _engine()
            base_label = engine.label
            engine.label = lambda cat: LABELS.get(cat) or base_label(cat)
            assert redact_docx(src, dest, engine) == {"SPONSOR": 1, "DRUG": 1}
        else:
            run(dest)
        live = _texts(dest, "t")
        deleted = _texts(dest, "delText")
        assert "Pharma" not in deleted and "[SPONSOR]" in deleted
        assert "Sponsor Acme Corp" in live
        assert "Kept [DRUG] " in live and "old words" in deleted
