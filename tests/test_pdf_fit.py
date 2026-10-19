"""PDF placeholders: full label, shrink to 6pt, short code plus legend, or blank.

No half-written labels. Word keeps the full label. Figures stay black boxes.
"""

from __future__ import annotations

import json
import subprocess
from pathlib import Path

import pymupdf
from docx import Document

from doc_redactor.docx_redactor import redact_docx
from doc_redactor.engine import RegexRule, WordlistRule, build_engine, redaction_label
from doc_redactor.pdf_redactor import (
    LEGEND_FOOTER,
    LEGEND_TITLE,
    MIN_PLACEHOLDER_SIZE,
    MSG_FIT_BLANK,
    MSG_FIT_SHORT,
    MSG_PDF_WIDTH,
    MSG_REPL_HINT,
    ShortPlan,
    fill_placeholder,
    fit_notice,
    fitted_fontsize,
    legend_lines,
    pdf_fit_counts,
    placeholder_choice,
    redact_pdf,
)
from doc_redactor.backpath import hit_list_payload
from doc_redactor.scan import KIND_PDF, scan_path, write_scan, write_scans
from helpers import assert_report_has_no_secrets


def test_messages_match_the_product_copy():
    assert MIN_PLACEHOLDER_SIZE == 6.0
    assert MSG_REPL_HINT == (
        "PDF 里位置窄时，会先把字缩小；还放不下就换成短码（如 [SP]），并在 PDF 最后加一页说明。Word 里总是写完整的替换文字。"
    )
    assert MSG_FIT_SHORT == "ⓘ 这个 PDF 有 {n} 处位置太窄，已换成短码，最后一页有对照说明。"
    assert MSG_FIT_BLANK == "⚠ 另有 {m} 处太窄，连短码也放不下，只留了空白，原文已删除。"
    assert MSG_PDF_WIDTH == "PDF 只能在原来的位置写字，空位和原文一样宽，看的人能看出原文大概多长。"
    assert LEGEND_TITLE == "替换说明"
    assert LEGEND_FOOTER == "本页由 doc-redactor 生成，只说明替换文字的意思，不含任何原文。"
    assert fit_notice(2, 0) == MSG_FIT_SHORT.format(n=2)
    assert fit_notice(2, 1) == MSG_FIT_SHORT.format(n=2) + "\n" + MSG_FIT_BLANK.format(m=1)
    assert fit_notice(0, 3) == ""
    assert "正则" not in MSG_REPL_HINT + MSG_FIT_SHORT + MSG_FIT_BLANK + MSG_PDF_WIDTH


def test_builtin_short_codes_and_custom_numbers_follow_creation_order():
    plan = ShortPlan(
        [
            "项目号",
            "方案号清单",
            "字母加数字",
            "申办方",
            "药名",
            "NCT 试验登记号",
            "邮箱",
            "手机号（中国大陆）",
            "合同号",
            "中心号",
        ]
    )
    assert plan.resolve("项目号") == ("[PR]", "项目号")
    assert plan.resolve("PROJECT_ID") == ("[PR]", "项目号")
    assert plan.resolve("方案号清单")[0] == "[PN]"
    assert plan.resolve("字母加数字")[0] == "[PN]"
    assert plan.resolve("申办方") == ("[SP]", "申办方")
    assert plan.resolve("SPONSOR") == ("[SP]", "申办方")
    assert plan.resolve("药名") == ("[SD]", "药名")
    assert plan.resolve("DRUG") == ("[SD]", "药名")
    assert plan.resolve("NCT")[0] == "[NCT]"
    assert plan.resolve("NCT 试验登记号")[0] == "[NCT]"
    assert plan.resolve("邮箱")[0] == "[EM]"
    assert plan.resolve("EMAIL")[0] == "[EM]"
    assert plan.resolve("手机号") == ("[PH]", "手机")
    assert plan.resolve("手机号（中国大陆）") == ("[PH]", "手机")
    assert plan.resolve("PHONE") == ("[PH]", "手机")
    assert plan.resolve("合同号") == ("[R1]", "合同号")
    assert plan.resolve("中心号") == ("[R2]", "中心号")
    # The same code is not issued twice inside one plan.
    codes = [plan.resolve(name)[0] for name in ("申办方", "SPONSOR", "合同号", "中心号")]
    assert codes == ["[SP]", "[SP]", "[R1]", "[R2]"]
    assert legend_lines(
        [
            ("[SD]", "药名", "[STUDY DRUG]"),
            ("[SP]", "申办方", "[SPONSOR]"),
            ("[SP]", "申办方", "[CO]"),
            ("[R2]", "中心号", "[REDACTED:中心号]"),
            ("[R1]", "合同号", "[REDACTED:合同号]"),
        ]
    ) == [
        "[SP] = 申办方（[SPONSOR]）",
        "[SD] = 药名（[STUDY DRUG]）",
        "[R1] = 合同号（[REDACTED:合同号]）",
        "[R2] = 中心号（[REDACTED:中心号]）",
    ]


def test_choice_covers_four_tiers_and_never_splits_a_label():
    full_box = pymupdf.Rect(0, 0, 200, 18)
    kind, shown = placeholder_choice(full_box, "[SPONSOR]", "[SP]")
    assert (kind, shown) == ("full", "[SPONSOR]")
    assert fitted_fontsize(full_box, "[SPONSOR]") >= 8

    shrink_box = pymupdf.Rect(0, 0, 40, 18)
    kind, shown = placeholder_choice(shrink_box, "[SPONSOR]", "[SP]")
    assert (kind, shown) == ("full", "[SPONSOR]")
    size = fitted_fontsize(shrink_box, "[SPONSOR]")
    assert size is not None and MIN_PLACEHOLDER_SIZE <= size < 8

    short_box = pymupdf.Rect(0, 0, 20, 16)
    kind, shown = placeholder_choice(short_box, "[SPONSOR]", "[SP]")
    assert (kind, shown) == ("short", "[SP]")
    assert fitted_fontsize(short_box, "[SPONSOR]") is None
    assert fitted_fontsize(short_box, "[SP]") >= MIN_PLACEHOLDER_SIZE

    blank_box = pymupdf.Rect(0, 0, 6, 6)
    kind, shown = placeholder_choice(blank_box, "[SPONSOR]", "[SP]")
    assert (kind, shown) == ("blank", "")
    assert fitted_fontsize(blank_box, "[SP]") is None


def test_fill_placeholder_shrinks_or_writes_nothing(tmp_path: Path):
    document = pymupdf.open()
    page = document.new_page()
    wide = pymupdf.Rect(72, 80, 72 + 200, 80 + 18)
    assert fill_placeholder(page, wide, "[PROTOCOL]") is True
    assert "[PROTOCOL]" in page.get_text()

    page2 = document.new_page()
    shrink = pymupdf.Rect(72, 80, 72 + 50, 80 + 16)
    assert fill_placeholder(page2, shrink, "[STUDY DRUG]") is True
    assert "[STUDYDRUG]" in page2.get_text().replace(" ", "")
    sizes = _placeholder_sizes(page2)
    assert sizes and min(sizes) >= MIN_PLACEHOLDER_SIZE

    page3 = document.new_page()
    tight = pymupdf.Rect(72, 80, 96, 88)
    assert fill_placeholder(page3, tight, "[STUDY DRUG]") is False
    raw = page3.get_text()
    assert raw.strip() == ""
    assert "[STU" not in raw
    assert "[SPO" not in raw
    document.close()


def test_pdf_four_tiers_keep_each_placeholder_inside_its_rect(tmp_path: Path):
    engine = _tier_engine()
    source = tmp_path / "tiers.pdf"
    _tier_pdf(source)
    sources = _term_rects(source)
    assert set(sources) == {"Acme Pharma International", "Fictizumab", "Acme", "I"}

    cli_out = tmp_path / "cli.pdf"
    cli_counts = redact_pdf(source, cli_out, engine)
    scan = scan_path(source, engine)
    gui_out = tmp_path / "gui.pdf"
    gui_counts = write_scan(scan, gui_out, engine=engine)

    for counts in (cli_counts, gui_counts):
        assert counts["申办方"] == 2
        assert counts["药名"] == 1
        assert counts["手机号"] == 1
        assert counts.short_count == 1
        assert counts.blank_count == 1
        assert "短码" not in counts
        assert "留空" not in counts

    for output in (cli_out, gui_out):
        mupdf_text, poppler = _extract(output)
        for body in (mupdf_text, poppler):
            assert "Acme" not in body
            assert "Fictizumab" not in body
            assert "International" not in body
            _assert_no_truncated_label(body)
        compact = _compact(poppler)
        assert "[SP]=申办方（[SPONSOR]）" in compact
        assert LEGEND_TITLE in compact
        assert _compact(LEGEND_FOOTER) in compact
        assert "[PH]" not in compact
        opened = pymupdf.open(output)
        try:
            assert opened.page_count == 2
            assert LEGEND_TITLE in opened[1].get_text()
            assert "[SP] = 申办方（[SPONSOR]）" in opened[1].get_text()
            assert "药名" not in opened[1].get_text()
            page = opened[0]
            wide = _blob(page, sources["Acme Pharma International"])
            drug = _blob(page, sources["Fictizumab"])
            short = _blob(page, sources["Acme"])
            blank = _blob(page, sources["I"])
            assert "[SPONSOR]" in wide and "[SP]" not in wide.replace("[SPONSOR]", "")
            assert "STUDY" in drug and "DRUG" in drug
            assert short.replace(" ", "") == "[SP]"
            assert blank.strip() == ""
            assert _inside(page, sources["Acme Pharma International"], "[SPONSOR]")
            assert _inside(page, sources["Fictizumab"], "[STUDY DRUG]")
            assert _inside(page, sources["Acme"], "[SP]")
            sponsor_size = _size_inside(page, sources["Acme Pharma International"])
            drug_size = _size_inside(page, sources["Fictizumab"])
            short_size = _size_inside(page, sources["Acme"])
            assert sponsor_size >= 8
            assert MIN_PLACEHOLDER_SIZE <= drug_size < sponsor_size
            assert short_size >= MIN_PLACEHOLDER_SIZE
        finally:
            opened.close()


def test_legend_page_only_when_a_short_code_was_used(tmp_path: Path):
    engine = _tier_engine()
    wide = tmp_path / "wide.pdf"
    _write_lines(wide, [(11, "Sponsor Acme Pharma International")])
    dest = tmp_path / "wide.redacted.pdf"
    counts = redact_pdf(wide, dest, engine)
    opened = pymupdf.open(dest)
    try:
        assert opened.page_count == 1
        assert LEGEND_TITLE not in opened[0].get_text()
    finally:
        opened.close()
    assert counts.short_count == 0
    assert counts.blank_count == 0

    blank_only = tmp_path / "blank.pdf"
    _write_lines(blank_only, [(6, "I")])
    blank_out = tmp_path / "blank.redacted.pdf"
    blank_counts = redact_pdf(blank_only, blank_out, engine)
    opened = pymupdf.open(blank_out)
    try:
        assert opened.page_count == 1
        text = opened[0].get_text()
        assert LEGEND_TITLE not in text
        assert "I" not in text
        assert "[" not in text
    finally:
        opened.close()
    assert blank_counts.short_count == 0
    assert blank_counts.blank_count == 1
    assert blank_counts["手机号"] == 1


def test_custom_rules_use_creation_order_and_share_one_legend(tmp_path: Path):
    engine = build_engine(
        [
            RegexRule("合同号", r"HT-\d{4}", 0),
            RegexRule("中心号", r"ZX-\d{4}", 0),
        ],
        [],
    )
    assert engine.category_order == ["合同号", "中心号"]
    source = tmp_path / "custom.pdf"
    _write_lines(source, [(11, "HT-1001"), (11, "ZX-2002 later")])
    dest = tmp_path / "custom.redacted.pdf"
    counts = redact_pdf(source, dest, engine)
    assert counts.short_count == 2
    assert counts.blank_count == 0
    mupdf_text, poppler = _extract(dest)
    for body in (mupdf_text, poppler):
        assert "HT-1001" not in body
        assert "ZX-2002" not in body
        _assert_no_truncated_label(body)
    compact = _compact(mupdf_text)
    assert "[R1]=合同号（[REDACTED:合同号]）" in compact
    assert "[R2]=中心号（[REDACTED:中心号]）" in compact
    assert compact.index("[R1]") < compact.index("[R2]")

    only_second = tmp_path / "second.pdf"
    _write_lines(only_second, [(11, "ZX-2002")])
    second_out = tmp_path / "second.redacted.pdf"
    redact_pdf(only_second, second_out, engine)
    second_text = "\n".join(page.get_text() for page in pymupdf.open(second_out))
    assert "[R2]" in second_text
    assert "[R1]" not in second_text
    pymupdf.open(second_out).close()


def test_disabling_the_short_hit_drops_the_legend(tmp_path: Path):
    engine = _tier_engine()
    source = tmp_path / "tiers.pdf"
    _tier_pdf(source)
    scan = scan_path(source, engine)
    for hit in scan.hits:
        if hit.text == "Acme":
            hit.enabled = False
    dest = tmp_path / "out.pdf"
    counts = write_scan(scan, dest, engine=engine)
    assert counts.short_count == 0
    opened = pymupdf.open(dest)
    try:
        assert opened.page_count == 1
        assert LEGEND_TITLE not in opened[0].get_text()
        assert "Acme Pharma International" not in opened[0].get_text()
    finally:
        opened.close()


def test_write_result_and_hit_list_report_counts_without_source(tmp_path: Path):
    engine = _tier_engine()
    source = tmp_path / "tiers.pdf"
    _tier_pdf(source)
    scan = scan_path(source, engine)
    batch = tmp_path / "out"
    written = write_scans([scan], batch, engine=engine)
    assert written[0].short_count == 1
    assert written[0].blank_count == 1
    report = assert_report_has_no_secrets(batch / "tiers.report.json")
    assert set(report) == {"申办方", "药名", "手机号"}
    assert "短码" not in report

    payload = hit_list_payload([scan], category_order=list(engine.category_order))
    raw = json.dumps(payload, ensure_ascii=False)
    assert "Acme" not in raw
    assert "Fictizumab" not in raw
    assert "International" not in raw
    assert payload[source.name]["短码"] == 1
    assert payload[source.name]["留空"] == 1
    assert payload[source.name]["申办方"] == 2


def test_docx_always_writes_the_full_label(tmp_path: Path):
    engine = _tier_engine()
    source = tmp_path / "note.docx"
    document = Document()
    document.add_paragraph("Acme")
    document.save(source)
    dest = tmp_path / "note.redacted.docx"
    redact_docx(source, dest, engine)
    text = Document(str(dest)).paragraphs[0].text
    assert text == "[SPONSOR]"
    assert "[SP]" not in text
    assert "Acme" not in text


def test_demo_pdf_notice_counts():
    from doc_redactor.gui.demo import demo_scans

    pdf = next(scan for scan in demo_scans() if scan.kind == KIND_PDF and scan.status == "ok" and scan.hits)
    label = lambda category: redaction_label(category)
    short_count, blank_count = pdf_fit_counts(pdf.enabled_hits, label, ["项目号", "申办方", "药名"])
    assert (short_count, blank_count) == (2, 1)
    assert fit_notice(short_count, blank_count) == (
        "ⓘ 这个 PDF 有 2 处位置太窄，已换成短码，最后一页有对照说明。\n"
        "⚠ 另有 1 处太窄，连短码也放不下，只留了空白，原文已删除。"
    )


def _tier_engine():
    return build_engine(
        [],
        [
            WordlistRule("申办方", ["Acme Pharma International", "Acme"], False),
            WordlistRule("药名", ["Fictizumab"], True),
            WordlistRule("手机号", ["I"], True),
        ],
    )


def _write_lines(path: Path, lines: list[tuple[float, str]]) -> None:
    document = pymupdf.open()
    page = document.new_page()
    y = 80
    for size, text in lines:
        font = "china-s" if any(ord(ch) > 127 for ch in text) else "helv"
        page.insert_text((72, y), text, fontname=font, fontsize=size)
        y += 40
    document.save(path)
    document.close()


def _tier_pdf(path: Path) -> None:
    _write_lines(
        path,
        [
            (11, "Sponsor Acme Pharma International"),
            (11, "Drug Fictizumab"),
            (11, "Also Acme"),
            (6, "I"),
        ],
    )


def _term_rects(path: Path) -> dict[str, pymupdf.Rect]:
    document = pymupdf.open(path)
    try:
        page = document[0]
        found = {}
        for term in ("Acme Pharma International", "Fictizumab", "Acme", "I"):
            hits = page.search_for(term)
            assert hits, term
            rects = [pymupdf.Rect(item) for item in hits]
            if term == "Acme":
                # The long name also contains this word. The standalone one sits lower.
                found[term] = max(rects, key=lambda item: item.y0)
            elif term == "I":
                found[term] = min(rects, key=lambda item: (item.width, -item.y0))
            else:
                found[term] = max(rects, key=lambda item: item.width)
        return found
    finally:
        document.close()


def _extract(path: Path) -> tuple[str, str]:
    document = pymupdf.open(path)
    try:
        mupdf_text = "\n".join(page.get_text("text") for page in document)
    finally:
        document.close()
    proc = subprocess.run(["pdftotext", str(path), "-"], check=True, capture_output=True)
    return mupdf_text, proc.stdout.decode("utf-8", errors="replace")


def _compact(text: str) -> str:
    return "".join(text.split())


def _assert_no_truncated_label(text: str) -> None:
    cleaned = text
    wholes = [
        "[PROTOCOL NO]",
        "[STUDY DRUG]",
        "[REDACTED:合同号]",
        "[REDACTED:中心号]",
        "[PROTOCOL]",
        "[SPONSOR]",
        "[NCT ID]",
        "[EMAIL]",
        "[PHONE]",
        "[NCT]",
        "[PR]",
        "[PN]",
        "[SP]",
        "[SD]",
        "[EM]",
        "[PH]",
        "[R1]",
        "[R2]",
    ]
    for token in wholes:
        cleaned = cleaned.replace(token, "")
    for fragment in ("[SPO", "[PROT", "[STU", "[EMA", "[PHO", "[NCT I", "[PROTOCOL N", "[RED"):
        assert fragment not in cleaned, cleaned


def _words_in(page, rect: pymupdf.Rect):
    chosen = []
    for word in page.get_text("words"):
        box = pymupdf.Rect(word[:4])
        if (
            box.x0 >= rect.x0 - 1.5
            and box.y0 >= rect.y0 - 1.5
            and box.x1 <= rect.x1 + 1.5
            and box.y1 <= rect.y1 + 1.5
        ):
            chosen.append(word)
    return chosen


def _blob(page, rect: pymupdf.Rect) -> str:
    return "".join(word[4] for word in _words_in(page, rect))


def _inside(page, rect: pymupdf.Rect, needle: str) -> bool:
    words = _words_in(page, rect)
    blob = "".join(word[4] for word in words)
    compact = _compact(blob)
    assert _compact(needle) in compact
    assert words
    return True


def _size_inside(page, rect: pymupdf.Rect) -> float:
    sizes = []
    data = page.get_text("dict")
    for block in data.get("blocks") or []:
        if block.get("type", 0) != 0:
            continue
        for line in block.get("lines") or []:
            for span in line.get("spans") or []:
                bbox = pymupdf.Rect(span["bbox"])
                if (
                    bbox.x0 >= rect.x0 - 1.5
                    and bbox.y0 >= rect.y0 - 1.5
                    and bbox.x1 <= rect.x1 + 1.5
                    and bbox.y1 <= rect.y1 + 1.5
                    and "[" in span.get("text", "")
                ):
                    sizes.append(float(span["size"]))
    assert sizes, rect
    return min(sizes)


def _placeholder_sizes(page) -> list[float]:
    sizes = []
    data = page.get_text("dict")
    for block in data.get("blocks") or []:
        if block.get("type", 0) != 0:
            continue
        for line in block.get("lines") or []:
            for span in line.get("spans") or []:
                if "[" in span.get("text", ""):
                    sizes.append(float(span["size"]))
    return sizes


