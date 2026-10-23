import io
from pathlib import Path

from docx import Document
from PIL import Image

import pytest

from doc_redactor.docx_redactor import redact_docx
from doc_redactor.engine import RegexRule, build_engine
from doc_redactor.image_redactor import MSG_BAD_IMAGE, MSG_NO_OCR, WriteBlocked
from doc_redactor.scan import scan_path, write_scans


def _engine():
    return build_engine([RegexRule(name="PROJECT_ID", pattern=r"ABC-\d{4}-\d{3}", flags=0)], [])


def _png() -> bytes:
    image = Image.new("RGB", (80, 40), "white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _docx_with_png(path: Path) -> None:
    raw = path.parent / "pixel.png"
    raw.write_bytes(_png())
    document = Document()
    document.add_paragraph("Status normal")
    document.add_picture(str(raw))
    document.save(path)


def _docx_with_broken_png(path: Path) -> None:
    _docx_with_png(path)
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.lower().startswith("word/media/"):
                data = b"not a picture"
            target.writestr(info, data)
    path.write_bytes(buffer.getvalue())


def test_broken_embed_stays_unopenable_and_writes_nothing(tmp_path, monkeypatch):
    monkeypatch.setattr("doc_redactor.tesseract_cmd.find_tesseract", lambda: "/tesseract")
    source = tmp_path / "CAPA回复1.docx"
    _docx_with_broken_png(source)
    scan = scan_path(source, _engine())
    assert scan.blocked_reason == MSG_BAD_IMAGE.format(n=1)
    out = tmp_path / "out"
    result = write_scans([scan], out, engine=_engine())[0]
    assert result.skipped and result.dest is None
    assert result.error == MSG_BAD_IMAGE.format(n=1)
    assert list(out.iterdir()) == []


def test_ocr_failure_is_not_reported_as_unopenable(tmp_path, monkeypatch):
    monkeypatch.setattr("doc_redactor.tesseract_cmd.find_tesseract", lambda: "/tesseract")

    def boom(*_args, **_kwargs):
        raise RuntimeError("filesystem error: Cannot convert character sequence: Illegal byte sequence")

    monkeypatch.setattr("doc_redactor.scan.analyze_image_entities", boom)
    source = tmp_path / "英文目录" / "CAPA回复1.docx"
    source.parent.mkdir()
    _docx_with_png(source)
    scan = scan_path(source, _engine())
    assert scan.blocked_reason == "有 1 张图片识别文字时出错，没法检查，未写出"
    assert "打不开" not in scan.blocked_reason
    out = tmp_path / "out"
    result = write_scans([scan], out, engine=_engine())[0]
    assert result.skipped and result.dest is None
    assert result.error == "有 1 张图片识别文字时出错，没法检查，未写出"
    assert not list(out.glob("*.docx"))


def test_write_raises_the_ocr_sentence_and_deletes_the_output(tmp_path, monkeypatch):
    monkeypatch.setattr("doc_redactor.tesseract_cmd.find_tesseract", lambda: "/tesseract")

    def boom(*_args, **_kwargs):
        raise RuntimeError("filesystem error: Cannot convert character sequence: Illegal byte sequence")

    monkeypatch.setattr("doc_redactor.docx_redactor.redact_image_bytes", boom)
    source = tmp_path / "中文 目录" / "CAPA回复1.docx"
    source.parent.mkdir()
    _docx_with_png(source)
    dest = tmp_path / "out" / "CAPA回复1.redacted.docx"
    with pytest.raises(WriteBlocked, match="识别文字时出错") as raised:
        redact_docx(source, dest, _engine())
    assert "打不开" not in str(raised.value)
    assert not dest.exists()


def test_docx_missing_language_data_keeps_the_install_message(tmp_path, monkeypatch):
    monkeypatch.setattr("doc_redactor.tesseract_cmd.find_tesseract", lambda: "/tesseract")

    def boom(*_args, **_kwargs):
        raise RuntimeError("Tesseract couldn't load any languages!")

    monkeypatch.setattr("doc_redactor.scan.analyze_image_entities", boom)
    source = tmp_path / "lang.docx"
    _docx_with_png(source)
    scan = scan_path(source, _engine())
    assert scan.blocked_reason == MSG_NO_OCR.format(n=1)
