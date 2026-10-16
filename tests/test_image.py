from pathlib import Path

import io
import zipfile

import pymupdf
import pytest
from docx import Document
from PIL import Image

from doc_redactor.cli import redact_path
from doc_redactor.config import load_rules
from doc_redactor.tesseract_cmd import TESSERACT_HINT, configure_tesseract, find_tesseract
from helpers import RULES, assert_report_has_no_secrets, load_make_samples

pytestmark = pytest.mark.skipif(find_tesseract() is None, reason=TESSERACT_HINT)


def _ocr(image: Image.Image) -> str:
    configure_tesseract()
    import pytesseract

    return pytesseract.image_to_string(image, lang="chi_sim+eng")


def _compact(text: str) -> str:
    return "".join(text.split()).lower()


def _assert_readable(text: str) -> None:
    folded = _compact(text)
    assert "abc-2026-001" in folded
    assert "acmepharma" in folded
    assert "fictizumab" in folded


def _assert_unreadable(text: str) -> None:
    folded = _compact(text)
    assert "abc-2026-001" not in folded
    assert "acmepharma" not in folded
    assert "fictizumab" not in folded


def test_png_and_jpg_are_painted_over(tmp_path: Path):
    samples = load_make_samples()
    png = tmp_path / "card.png"
    samples.render_png(png)
    _assert_readable(_ocr(Image.open(png)))

    outdir = tmp_path / "out"
    counts = redact_path(png, outdir, load_rules(RULES))
    redacted = Image.open(outdir / "card.redacted.png")
    _assert_unreadable(_ocr(redacted))
    report = assert_report_has_no_secrets(outdir / "card.report.json")
    assert report["PROJECT_ID"] >= 1
    assert report["SPONSOR"] >= 1
    assert report["DRUG"] >= 1
    assert counts == report

    jpg = tmp_path / "card.jpg"
    Image.open(png).convert("RGB").save(jpg, format="JPEG", quality=95)
    _assert_readable(_ocr(Image.open(jpg)))
    redact_path(jpg, outdir, load_rules(RULES))
    _assert_unreadable(_ocr(Image.open(outdir / "card.redacted.jpg")))
    assert_report_has_no_secrets(outdir / "card.report.json")


def test_docx_embedded_image(tmp_path: Path):
    samples = load_make_samples()
    png = tmp_path / "card.png"
    samples.render_png(png)
    source = tmp_path / "with-image.docx"
    document = Document()
    document.add_paragraph("Status normal")
    document.add_picture(str(png))
    document.save(source)

    outdir = tmp_path / "out"
    redact_path(source, outdir, load_rules(RULES))
    redacted = Document(outdir / "with-image.redacted.docx")
    assert redacted.paragraphs[0].text == "Status normal"

    with zipfile.ZipFile(outdir / "with-image.redacted.docx") as archive:
        name = next(item for item in archive.namelist() if item.lower().startswith("word/media/"))
        payload = archive.read(name)
    _assert_unreadable(_ocr(Image.open(io.BytesIO(payload))))

    with zipfile.ZipFile(source) as archive:
        name = next(item for item in archive.namelist() if item.lower().startswith("word/media/"))
        original = archive.read(name)
    _assert_readable(_ocr(Image.open(io.BytesIO(original))))
    assert_report_has_no_secrets(outdir / "with-image.report.json")


def test_pdf_embedded_image(tmp_path: Path):
    samples = load_make_samples()
    png = tmp_path / "card.png"
    samples.render_png(png)
    source = tmp_path / "with-image.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 720), "Project ABC-2026-001 Sponsor Acme Pharma Drug Fictizumab")
    page.insert_image(pymupdf.Rect(72, 400, 540, 620), filename=str(png))
    document.save(source)
    document.close()

    outdir = tmp_path / "out"
    redact_path(source, outdir, load_rules(RULES))
    opened = pymupdf.open(outdir / "with-image.redacted.pdf")
    text = "\n".join(item.get_text() for item in opened)
    assert "ABC-2026-001" not in text
    assert "Acme Pharma" not in text
    assert "Fictizumab" not in text

    images = opened[0].get_images()
    assert images
    pixmap = pymupdf.Pixmap(opened, images[0][0])
    if pixmap.alpha:
        pixmap = pymupdf.Pixmap(pixmap, 0)
    image = Image.open(io.BytesIO(pixmap.tobytes("png")))
    _assert_unreadable(_ocr(image))
    opened.close()
    report = assert_report_has_no_secrets(outdir / "with-image.report.json")
    assert report["PROJECT_ID"] >= 1
    assert report["SPONSOR"] >= 1
    assert report["DRUG"] >= 1


def test_make_samples_round_trip(tmp_path: Path):
    samples = load_make_samples()
    png = tmp_path / "sample.png"
    docx = tmp_path / "sample.docx"
    pdf = tmp_path / "sample.pdf"
    samples.render_png(png)
    samples.render_docx(docx, png)
    samples.render_pdf(pdf, png)
    outdir = tmp_path / "out"
    engine = load_rules(RULES)
    for path in (png, docx, pdf):
        redact_path(path, outdir, engine)
        assert_report_has_no_secrets(outdir / f"{path.stem}.report.json")

    document = Document(outdir / "sample.redacted.docx")
    assert "fictizumab" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    opened = pymupdf.open(outdir / "sample.redacted.pdf")
    text = "\n".join(page.get_text() for page in opened)
    opened.close()
    assert "ABC-2026-001" not in text
    assert "Fictizumab" not in text
    assert "fictizumab stays" in text
    _assert_unreadable(_ocr(Image.open(outdir / "sample.redacted.png")))
