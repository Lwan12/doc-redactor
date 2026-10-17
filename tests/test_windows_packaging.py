import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
ZIP_TOOL = ROOT / "tools" / "make_windows_zip.py"
CHECK_TOOL = ROOT / "tools" / "check_redacted_output.py"


def _run(script: Path, args: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(script), *args],
        check=False,
        capture_output=True,
        text=True,
    )


def test_spec_is_onedir_and_manifest_is_as_invoker():
    source = (ROOT / "doc-redactor.spec").read_text(encoding="utf-8")
    compile(source, "doc-redactor.spec", "exec")
    assert "exclude_binaries=True" in source
    assert "uac_admin=False" in source
    assert 'contents_directory="."' in source
    assert "asInvoker.manifest" in source
    assert "MERGE(" not in source
    manifest = (ROOT / "packaging" / "asInvoker.manifest").read_text(encoding="utf-8")
    assert 'level="asInvoker"' in manifest
    assert "requireAdministrator" not in manifest


def _as_invoker_blob() -> bytes:
    return b'<requestedExecutionLevel level="asInvoker" uiAccess="false"/>'


def _stage_onedir(root: Path, *, admin: bool = False) -> None:
    (root / "tesseract" / "tessdata").mkdir(parents=True)
    if admin:
        blob = (
            b'<requestedExecutionLevel level="asInvoker" uiAccess="false"/>'
            b'<requestedExecutionLevel level="requireAdministrator"/>'
        )
    else:
        blob = _as_invoker_blob()
    (root / "doc-redactor.exe").write_bytes(blob)
    (root / "doc-redactor-gui.exe").write_bytes(blob)
    (root / "tesseract" / "tesseract.exe").write_bytes(b"tess")
    (root / "tesseract" / "libtesseract.dll").write_bytes(b"dll")
    (root / "tesseract" / "tessdata" / "eng.traineddata").write_bytes(b"eng")
    (root / "tesseract" / "tessdata" / "chi_sim.traineddata").write_bytes(b"chi")


def test_make_windows_zip_writes_hash_and_rejects_admin(tmp_path: Path):
    onedir = tmp_path / "doc-redactor"
    _stage_onedir(onedir)
    zip_path = tmp_path / "doc-redactor-windows.zip"
    sums_path = tmp_path / "SHA256SUMS.txt"
    created = _run(ZIP_TOOL, [str(onedir), str(zip_path), str(sums_path)])
    assert created.returncode == 0, created.stderr
    line = sums_path.read_text(encoding="ascii").strip()
    digest, name = line.split()
    assert name == "doc-redactor-windows.zip"
    assert len(digest) == 64
    with zipfile.ZipFile(zip_path) as archive:
        names = set(archive.namelist())
    assert "doc-redactor/doc-redactor.exe" in names
    assert "doc-redactor/doc-redactor-gui.exe" in names
    assert "doc-redactor/tesseract/tesseract.exe" in names
    assert "doc-redactor/tesseract/libtesseract.dll" in names
    assert "doc-redactor/tesseract/tessdata/chi_sim.traineddata" in names
    assert "doc-redactor/tesseract/tessdata/eng.traineddata" in names

    admin = tmp_path / "admin"
    _stage_onedir(admin, admin=True)
    rejected = _run(
        ZIP_TOOL,
        [str(admin), str(tmp_path / "no.zip"), str(tmp_path / "no.txt")],
    )
    assert rejected.returncode == 1
    assert "administrator" in rejected.stderr


def _write_docx(path: Path, *, body: str, header: str) -> None:
    document = Document()
    document.sections[0].header.paragraphs[0].text = header
    document.add_paragraph(body)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _write_pdf(path: Path, text: str) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 720), text)
    document.save(path)
    document.close()


def _fake_tesseract(directory: Path, stdout: str, code: int = 0) -> Path:
    directory.mkdir(parents=True)
    (directory / "tessdata").mkdir()
    executable = directory / "tesseract"
    executable.write_text(
        "#!/usr/bin/env python3\n"
        "import sys\n"
        f"sys.stdout.write({stdout!r})\n"
        f"raise SystemExit({code})\n",
        encoding="utf-8",
    )
    executable.chmod(0o755)
    return executable


def _write_output(directory: Path, *, docx_body: str, docx_header: str, pdf_text: str) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    _write_docx(directory / "sample.redacted.docx", body=docx_body, header=docx_header)
    _write_pdf(directory / "sample.redacted.pdf", pdf_text)
    Image.new("RGB", (8, 8), "white").save(directory / "sample.redacted.png")
    report = {"DRUG": 1, "PROJECT_ID": 2, "SPONSOR": 1}
    (directory / "sample.report.json").write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


POSIX_ONLY = pytest.mark.skipif(
    sys.platform == "win32", reason="uses fake script/symlink tesseract; real exe is checked by the verify job"
)


@POSIX_ONLY
def test_check_redacted_output_accepts_clean_files_and_rejects_leaks(tmp_path: Path):
    clean = tmp_path / "clean"
    _write_output(
        clean,
        docx_body="Keep fictizumab in lowercase. [PROTOCOL]",
        docx_header="[SPONSOR]",
        pdf_text="fictizumab stays [PROTOCOL]",
    )
    tesseract = _fake_tesseract(tmp_path / "tesseract", "")
    ok = _run(CHECK_TOOL, ["--output", str(clean), "--tesseract", str(tesseract)])
    assert ok.returncode == 0, ok.stderr

    leaked = tmp_path / "leaked"
    _write_output(
        leaked,
        docx_body="Keep fictizumab in lowercase. [PROTOCOL]",
        docx_header="Acme Pharma",
        pdf_text="fictizumab stays [PROTOCOL]",
    )
    bad_docx = _run(CHECK_TOOL, ["--output", str(leaked), "--tesseract", str(tesseract)])
    assert bad_docx.returncode == 1
    assert "Acme Pharma" in bad_docx.stderr

    readable = _fake_tesseract(tmp_path / "readable", "Project ABC-2026-001")
    bad_png = _run(CHECK_TOOL, ["--output", str(clean), "--tesseract", str(readable)])
    assert bad_png.returncode == 1
    assert "abc-2026-001" in bad_png.stderr

    report = clean / "sample.report.json"
    report.write_text('{"DRUG": 1, "note": "Fictizumab"}\n', encoding="utf-8")
    bad_report = _run(CHECK_TOOL, ["--output", str(clean), "--tesseract", str(tesseract)])
    assert bad_report.returncode == 1
    assert "Fictizumab" in bad_report.stderr


@POSIX_ONLY
def test_checker_reads_tessdata_beside_symlink(tmp_path: Path):
    output = tmp_path / "out"
    _write_output(
        output,
        docx_body="Keep fictizumab in lowercase. [PROTOCOL]",
        docx_header="[SPONSOR]",
        pdf_text="fictizumab stays [PROTOCOL]",
    )
    real_dir = tmp_path / "real"
    real_dir.mkdir()
    real = real_dir / "tesseract"
    real.write_text(
        "#!/usr/bin/env python3\nimport sys\nsys.stdout.write('')\nraise SystemExit(0)\n",
        encoding="utf-8",
    )
    real.chmod(0o755)
    package = tmp_path / "pkg"
    (package / "tessdata").mkdir(parents=True)
    link = package / "tesseract"
    link.symlink_to(real)
    checked = _run(CHECK_TOOL, ["--output", str(output), "--tesseract", str(link)])
    assert checked.returncode == 0, checked.stderr
