"""Written docx keeps no document properties. Fictional data only."""

import zipfile
from pathlib import Path

from docx import Document

from doc_redactor.docx_redactor import redact_docx
from doc_redactor.engine import RegexRule, build_engine
from doc_redactor.scan import scan_path, write_scan

SECRET = "Acme Pharma Fictizumab ABC-2026-001"

APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    f"<Application>Microsoft Office Word</Application><Company>{SECRET}</Company>"
    f"<Manager>{SECRET}</Manager><HyperlinkBase>{SECRET}</HyperlinkBase><Template>{SECRET}.dotm</Template>"
    "</Properties>"
)
CUSTOM_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Sponsor">'
    f"<vt:lpwstr>{SECRET}</vt:lpwstr></property></Properties>"
)
CUSTOM_CT = (
    '<Override PartName="/docProps/custom.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/>'
)
CUSTOM_REL = (
    '<Relationship Id="rIdCustom9" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" '
    'Target="docProps/custom.xml"/>'
)


def _build(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Project ABC-2026-001")
    cp = doc.core_properties
    cp.author = SECRET
    cp.title = SECRET
    cp.subject = SECRET
    cp.keywords = SECRET
    cp.comments = SECRET
    cp.last_modified_by = SECRET
    cp.category = SECRET
    cp.content_status = SECRET
    cp.identifier = SECRET
    cp.version = SECRET
    raw = path.with_suffix(".raw.docx")
    doc.save(str(raw))
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/app.xml":
                data = APP_XML.encode()
            elif info.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", CUSTOM_CT.encode() + b"</Types>")
            elif info.filename == "_rels/.rels":
                data = data.replace(b"</Relationships>", CUSTOM_REL.encode() + b"</Relationships>")
            dst.writestr(info, data)
        dst.writestr("docProps/custom.xml", CUSTOM_XML)
    raw.unlink()


def _assert_clean(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            assert "thumbnail" not in name.lower(), name
            if name.startswith("docProps/") and name.endswith(".xml"):
                text = archive.read(name).decode("utf-8")
                for word in ("Acme", "Fictizumab", "ABC-2026"):
                    assert word not in text, name
        app = archive.read("docProps/app.xml").decode("utf-8")
    assert "Microsoft Office Word" in app or "<Application" in app  # harmless field kept
    cp = Document(str(path)).core_properties
    for field in ("author", "title", "subject", "keywords", "comments", "last_modified_by", "category"):
        assert not getattr(cp, field)


def _engine():
    return build_engine([RegexRule(name="PROJECT_ID", pattern=r"ABC-\d{4}-\d{3}", flags=0)], [])


def test_scan_write_clears_properties(tmp_path: Path):
    src = tmp_path / "a.docx"
    _build(src)
    scan = scan_path(src, _engine())
    dest = tmp_path / "out.docx"
    write_scan(scan, dest)
    _assert_clean(dest)


def test_cli_redact_clears_properties(tmp_path: Path):
    src = tmp_path / "a.docx"
    _build(src)
    dest = tmp_path / "out.docx"
    redact_docx(src, dest, _engine())
    _assert_clean(dest)
