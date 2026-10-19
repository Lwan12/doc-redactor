"""Scan and write with 「去掉页眉页脚」: delete headers/footers first, then rules.

Fictional data only.
"""

from __future__ import annotations

import zipfile

import pymupdf
from docx import Document

from doc_redactor import header_footer as hf
from doc_redactor.engine import WordlistRule, build_engine
from doc_redactor.scan import scan_path, set_file_hf_confirmed, write_scan, write_scans
from test_header_footer import HEADER, SPONSOR_FOOTER, SIGN, _docx, _plain_pdf, _repeat_pdf, _tagged_pdf


def _engine():
    return build_engine([], [WordlistRule("项目号", ["ABC-2026-001"], True), WordlistRule("药名", ["Fictizumab"], True)])


def _save(tmp_path, name, data):
    path = tmp_path / name
    path.write_bytes(data)
    return path


def test_option_off_changes_nothing(tmp_path):
    path = _save(tmp_path, "t.pdf", _tagged_pdf())
    scan = scan_path(path, _engine())
    assert scan.hf_route == "" and not scan.hf_active
    assert all(not h.hf_role for h in scan.hits)


def test_tagged_pdf_hits_in_header_are_not_counted_and_header_is_deleted(tmp_path):
    path = _save(tmp_path, "t.pdf", _tagged_pdf())
    scan = scan_path(path, _engine(), remove_hf=True)
    assert scan.hf_route == hf.ROUTE_TAG and scan.hf_active
    in_hf = [h for h in scan.hits if h.hf_role]
    body = [h for h in scan.hits if not h.hf_role]
    assert len(in_hf) == 4  # ABC-2026-001 in each page header
    assert scan.enabled_count() == len(body) == 4 * 2  # Fictizumab, two body lines a page
    assert scan.hf_delete_pages == 4 and scan.hf_pending_pages == 0
    roles = {role for (_p, _l), role in scan.hf_lines.items()}
    assert roles == {hf.ROLE_HEADER, hf.ROLE_FOOTER, hf.ROLE_SIGNATURE}
    dest = tmp_path / "out.pdf"
    write_scan(scan, dest)
    text = "".join(p.get_text() for p in pymupdf.open(dest))
    for gone in (HEADER, SPONSOR_FOOTER, SIGN, "ABC-2026-001", "Fictizumab"):
        assert gone not in text
    assert "body first line" in text and "Layout note" in text
    out = pymupdf.open(dest)
    streams = b"".join(out.xref_stream(x) or b"" for x in range(1, out.xref_length()) if out.xref_is_stream(x))
    assert b"Version 1.0" not in streams and b"CONFIDENTIAL" not in streams and b"Envelope" not in streams


def test_file_kept_counts_header_hits_and_replaces_inside_header(tmp_path):
    path = _save(tmp_path, "t.pdf", _tagged_pdf())
    scan = scan_path(path, _engine(), remove_hf=True)
    scan.hf_keep = True
    assert not scan.hf_active
    assert scan.enabled_count() == len(scan.hits) == 12
    dest = tmp_path / "kept.pdf"
    write_scan(scan, dest)
    text = "".join(p.get_text() for p in pymupdf.open(dest))
    assert "ABC-2026-001" not in text and "Version 1.0" in text and SPONSOR_FOOTER in text


def _texts(path):
    return "".join(p.get_text() for p in pymupdf.open(path))


def test_repeat_blocks_stay_until_confirmed_and_never_block_the_write(tmp_path):
    path = _save(tmp_path, "r.pdf", _repeat_pdf())
    scan = scan_path(path, _engine(), remove_hf=True)
    assert scan.hf_route == hf.ROUTE_REPEAT and scan.hf_pending and not scan.hf_active
    pending_hits = scan.enabled_count()
    # Not confirmed: written as usual, repeated blocks kept, their hits replaced.
    dest = tmp_path / "r-pending.pdf"
    write_scan(scan, dest)
    text = _texts(dest)
    assert SPONSOR_FOOTER in text and "Page 1/12" in text and "ABC-2026-001" not in text
    # Confirmed: deleted.
    assert set_file_hf_confirmed(scan, True) and scan.hf_active
    assert scan.enabled_count() < pending_hits
    dest = tmp_path / "r-ok.pdf"
    write_scan(scan, dest)
    text = _texts(dest)
    assert SPONSOR_FOOTER not in text and "Page " not in text and "Repeated body sentence" in text
    # Undone: kept again.
    assert set_file_hf_confirmed(scan, False) and not scan.hf_active
    assert scan.enabled_count() == pending_hits
    dest = tmp_path / "r-undo.pdf"
    write_scan(scan, dest)
    assert SPONSOR_FOOTER in _texts(dest)


def test_tags_and_signing_line_go_whether_or_not_confirmed(tmp_path):
    path = _save(tmp_path, "t.pdf", _tagged_pdf())
    for confirmed in (False, True):
        scan = scan_path(path, _engine(), remove_hf=True)
        set_file_hf_confirmed(scan, confirmed)
        assert scan.hf_active and not scan.hf_pending
        dest = tmp_path / f"t-{confirmed}.pdf"
        write_scan(scan, dest)
        text = _texts(dest)
        for gone in (HEADER, SPONSOR_FOOTER, SIGN):
            assert gone not in text, (confirmed, gone)


def _signed_plain_pdf() -> bytes:
    doc = pymupdf.open()
    for i in range(3):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 15), SIGN, fontsize=7)
        page.insert_text((72, 60 + 37 * i), f"Fictizumab paragraph {i}", fontsize=10)
        page.insert_text((72, 400), f"Body text number {i} for Acme Pharma", fontsize=10)
    return doc.tobytes()


def test_not_found_pdf_still_drops_the_signing_line(tmp_path):
    """产品 ①: nothing recognised as header/footer, the signing line still goes,
    and the file's box stays usable."""
    path = _save(tmp_path, "s.pdf", _signed_plain_pdf())
    scan = scan_path(path, _engine(), remove_hf=True)
    assert scan.hf_route == hf.ROUTE_NONE and scan.hf_has_sign and scan.hf_active
    assert scan.hf_sign_pages == 3 and scan.hf_pending_pages == 0
    dest = tmp_path / "s-out.pdf"
    write_scan(scan, dest)
    text = _texts(dest)
    assert "Envelope" not in text and "Fictizumab" not in text and "Body text number 2" in text
    scan.hf_keep = True  # the user unticks this file: the line stays
    dest = tmp_path / "s-kept.pdf"
    write_scan(scan, dest)
    assert "Envelope" in _texts(dest)


def test_signing_page_with_pending_header_counts_in_both(tmp_path):
    """产品 ②: the same page is in 将删除 (signing line) and 待确认 (repeats)."""
    doc = pymupdf.open(stream=_repeat_pdf(), filetype="pdf")
    for page in doc:
        page.insert_text((72, 15), SIGN, fontsize=7)
    path = _save(tmp_path, "rs.pdf", doc.tobytes())
    scan = scan_path(path, _engine(), remove_hf=True)
    assert scan.hf_route == hf.ROUTE_REPEAT and scan.hf_pending
    assert scan.hf_delete_pages == 12  # signing line on every page
    assert scan.hf_pending_pages == 12  # repeated blocks on every page, too
    set_file_hf_confirmed(scan, True)
    assert scan.hf_delete_pages == 12 and scan.hf_pending_pages == 0


def test_plain_pdf_reports_not_found(tmp_path):
    path = _save(tmp_path, "p.pdf", _plain_pdf())
    scan = scan_path(path, _engine(), remove_hf=True)
    assert scan.hf_route == hf.ROUTE_NONE and not scan.hf_active
    assert scan.enabled_count() == len(scan.hits)


def test_file_with_hits_only_in_header_is_still_written(tmp_path):
    path = _save(tmp_path, "h.pdf", _tagged_pdf())
    engine = build_engine([], [WordlistRule("项目号", ["ABC-2026-001"], True)])
    scan = scan_path(path, engine, remove_hf=True)
    assert scan.enabled_count() == 0 and scan.hf_active
    (result,) = write_scans([scan], tmp_path / "out")
    assert result.dest is not None and result.hf_removed
    assert HEADER not in "".join(p.get_text() for p in pymupdf.open(result.dest))


def test_word_header_hits_hidden_and_header_emptied(tmp_path):
    path = _docx(tmp_path)
    scan = scan_path(path, _engine(), remove_hf=True)
    assert scan.hf_route == hf.ROUTE_WORD and scan.hf_active and scan.hf_watermark
    hidden = [h for h in scan.hits if h.hf_role]
    assert hidden and all(scan.hf_units[h.unit] == h.hf_role for h in hidden)
    assert scan.enabled_count() == 1  # Fictizumab in the body
    dest = tmp_path / "out.docx"
    write_scan(scan, dest)
    assert hf.verify_docx_output(dest) == 0
    body = [p.text for p in Document(dest).paragraphs]
    assert body == ["Body mentions [STUDY DRUG] once."] or "Fictizumab" not in body[0]
    with zipfile.ZipFile(dest) as zf:
        xml = "".join(zf.read(n).decode() for n in zf.namelist() if n.startswith("word/header"))
    assert "ABC-2026-001" not in xml


_W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"'


def _sdt(text: str) -> str:
    return (
        f'<w:sdt {_W_NS}><w:sdtPr><w:alias w:val="Field"/></w:sdtPr><w:sdtContent>'
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    )


def _textbox(text: str) -> str:
    return (
        f"<w:p {_W_NS}><w:r><w:pict><v:shape><v:textbox><w:txbxContent>"
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )


def test_word_header_controls_and_text_boxes_go_body_ones_stay(tmp_path):
    """产品第 8 条: the header/footer part is emptied as a whole, with its
    content controls and text boxes; the same things in the body are kept.
    Checked on the written XML."""
    from docx.oxml import parse_xml

    document = Document()
    document.add_paragraph("Body mentions Fictizumab once.")
    body = document.element.body
    sect = body[-1]
    sect.addprevious(parse_xml(_sdt("Body control Acme Pharma note")))
    sect.addprevious(parse_xml(_textbox("Body box star sea note")))
    header = document.sections[0].header
    header.paragraphs[0].text = "Header line"
    header._element.append(parse_xml(_sdt("Header control ABC-2026-001")))
    header._element.append(parse_xml(_textbox("Header box ABC-2026-001")))
    footer = document.sections[0].footer
    footer._element.append(parse_xml(_sdt("Footer control Acme")))
    src = tmp_path / "cc.docx"
    document.save(src)

    scan = scan_path(src, _engine(), remove_hf=True)
    assert scan.hf_active
    dest = tmp_path / "cc-out.docx"
    write_scan(scan, dest)
    with zipfile.ZipFile(dest) as zf:
        parts = {n: zf.read(n).decode() for n in zf.namelist() if n.endswith(".xml")}
    hf_xml = "".join(v for n, v in parts.items() if n.startswith(("word/header", "word/footer")))
    for gone in ("Header control", "Header box", "Footer control", "w:sdt", "txbxContent", "Header line"):
        assert gone not in hf_xml, gone
    doc_xml = parts["word/document.xml"]
    assert "Body control Acme Pharma note" in doc_xml and ("<w:sdt>" in doc_xml or "<w:sdt " in doc_xml)
    assert "Body box star sea note" in doc_xml and "txbxContent" in doc_xml
    assert "Fictizumab" not in doc_xml  # the body rule still applied


def test_word_header_hits_count_only_when_the_header_is_kept(tmp_path):
    path = _docx(tmp_path)
    scan = scan_path(path, _engine(), remove_hf=True)
    in_hf = [h for h in scan.hits if h.hf_role]
    assert in_hf and scan.enabled_count() == len(scan.hits) - len(in_hf)
    scan.hf_keep = True
    assert scan.enabled_count() == len(scan.hits)
