from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

from doc_redactor.cli import main
from doc_redactor.config import load_rules
from doc_redactor.scan import scan_path, write_scans
from doc_redactor.sheet import MSG_ENCRYPTED_SHEET, MSG_PICTURE_SHEET, extract_formulas, extract_in_scope
from tests.helpers import RULES

TOKENS = ("ABC-2026-001", "Acme Pharma", "Fictizumab")
OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")


def _engine():
    return load_rules(RULES)


def _xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Fictizumab"
    ws["A1"] = "ABC-2026-001"
    ws["B1"] = "Acme Pharma"
    ws["C1"] = "Fictizumab"
    ws["A2"] = '="ABC-2026-001"'
    ws.row_dimensions[3].hidden = True
    ws["A3"] = "ABC-2026-001"
    ws.column_dimensions["G"].hidden = True
    ws["G1"] = "Fictizumab"
    ws["D4"] = "hello"
    ws["D4"].comment = Comment("batch Fictizumab", "Reviewer")
    ws.oddHeader.left.text = "ABC-2026-001"
    ws.oddFooter.center.text = "Fictizumab"
    ws["A10"] = 1
    ws["A11"] = 2
    chart = BarChart()
    chart.title = "Fictizumab"
    chart.add_data(Reference(ws, min_col=1, min_row=10, max_row=11))
    ws.add_chart(chart, "E10")
    wb.save(path)


def _xls(path: Path) -> None:
    import xlwt
    from xlwt import Formula

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Fictizumab")
    ws.write(0, 0, "ABC-2026-001")
    ws.write(1, 1, "Acme Pharma")
    ws.write(0, 2, "Fictizumab")
    ws.write(2, 2, Formula('"ABC-2026-001"'))
    ws.header_str = b"&LABC-2026-001"
    ws.footer_str = b"&CFictizumab"
    ws.row(3).hidden = 1
    ws.write(3, 0, "ABC-2026-001")
    ws.col(6).hidden = 1
    ws.write(0, 6, "Fictizumab")
    wb.save(str(path))


def _csv(path: Path) -> None:
    path.write_text(
        "ABC-2026-001,Acme Pharma,Fictizumab\n=\"ABC-2026-001\",ok,ok\n",
        encoding="utf-8",
        newline="",
    )


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write(path: Path, outdir: Path, *, disable_text: str | None = None):
    engine = _engine()
    before = _digest(path)
    scan = scan_path(path, engine)
    if disable_text is not None:
        for hit in scan.hits:
            if hit.text == disable_text and getattr(hit.place, "row", None) == 0 and getattr(hit.place, "column", None) == 1:
                hit.enabled = False
    results = write_scans([scan], outdir, replacements=engine.replacements, overwrite=True, engine=engine)
    assert _digest(path) == before
    return scan, results[0]


def _assert_clean(dest: Path) -> None:
    blob = extract_in_scope(dest).blob()
    for token in TOKENS:
        assert blob.count(token) == 0
    formulas = "\n".join(extract_formulas(dest))
    assert "ABC-2026-001" in formulas


def test_xlsx_redacts_in_scope_and_keeps_formula_and_chart(tmp_path: Path):
    src = tmp_path / "demo.xlsx"
    _xlsx(src)
    _scan, result = _write(src, tmp_path / "out")
    assert result.dest is not None and result.dest.is_file()
    _assert_clean(result.dest)
    with zipfile.ZipFile(result.dest) as package:
        chart = package.read("xl/charts/chart1.xml")
    assert b"Fictizumab" in chart


def test_xlsx_unchecked_cell_keeps_its_token(tmp_path: Path):
    src = tmp_path / "demo.xlsx"
    _xlsx(src)
    _scan, result = _write(src, tmp_path / "out", disable_text="Acme Pharma")
    assert result.dest is not None
    blob = extract_in_scope(result.dest).blob()
    assert blob.count("Acme Pharma") == 1
    assert blob.count("ABC-2026-001") == 0
    assert blob.count("Fictizumab") == 0
    assert "ABC-2026-001" in "\n".join(extract_formulas(result.dest))


def test_xls_redacts_in_scope_and_keeps_formula(tmp_path: Path):
    src = tmp_path / "demo.xls"
    _xls(src)
    _scan, result = _write(src, tmp_path / "out")
    assert result.dest is not None and result.dest.is_file()
    assert result.dest.suffix == ".xls"
    _assert_clean(result.dest)


def test_csv_redacts_values_and_keeps_formula_field(tmp_path: Path):
    src = tmp_path / "demo.csv"
    _csv(src)
    _scan, result = _write(src, tmp_path / "out")
    assert result.dest is not None and result.dest.is_file()
    _assert_clean(result.dest)
    text = result.dest.read_text(encoding="utf-8")
    assert '="ABC-2026-001"' in text


def test_encrypted_xlsx_fails_with_its_sentence_and_writes_nothing(tmp_path: Path):
    src = tmp_path / "locked.xlsx"
    src.write_bytes(OLE_MAGIC + b"\x00" * 512)
    out = tmp_path / "out"
    scan, result = _write(src, out)
    assert scan.status == "fail"
    assert scan.error == MSG_ENCRYPTED_SHEET
    assert result.dest is None
    assert list(out.glob("*")) == []


def test_picture_xlsx_fails_with_its_sentence_and_writes_nothing(tmp_path: Path):
    src = tmp_path / "pict.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ABC-2026-001"
    png = tmp_path / "dot.png"
    PILImage.new("RGB", (8, 8), "white").save(png)
    ws.add_image(XLImage(str(png)), "C3")
    wb.save(src)
    out = tmp_path / "out"
    scan, result = _write(src, out)
    assert scan.status == "fail"
    assert scan.error == MSG_PICTURE_SHEET
    assert result.dest is None
    assert list(out.glob("*")) == []


def test_xlsm_is_rejected_and_not_written(tmp_path: Path):
    src = tmp_path / "macro.xlsm"
    good = tmp_path / "demo.xlsx"
    _xlsx(good)
    src.write_bytes(good.read_bytes())
    out = tmp_path / "out"
    scan, result = _write(src, out)
    assert scan.status == "fail"
    assert ".xlsm" not in (scan.error or "")
    assert "表格" in (scan.error or "")
    assert result.dest is None
    assert list(out.glob("*")) == []


def test_cli_csv_keeps_the_formula_and_writes_the_file(tmp_path: Path, capsys):
    src = tmp_path / "demo.csv"
    _csv(src)
    out = tmp_path / "out"
    code = main([str(src), "-r", str(RULES), "-o", str(out)])
    captured = capsys.readouterr()
    assert code == 0, captured.err
    dest = out / "demo.redacted.csv"
    assert dest.is_file()
    text = dest.read_text(encoding="utf-8")
    assert '="ABC-2026-001"' in text
    assert "ABC-2026-001" not in text.splitlines()[0]


def test_cli_failures_keep_their_sentences_and_write_nothing(tmp_path: Path, capsys):
    locked = tmp_path / "locked.xlsx"
    locked.write_bytes(OLE_MAGIC + b"\x00" * 512)
    locked_out = tmp_path / "locked-out"
    code = main([str(locked), "-r", str(RULES), "-o", str(locked_out)])
    captured = capsys.readouterr()
    assert code == 1
    assert MSG_ENCRYPTED_SHEET in captured.err
    assert list(locked_out.glob("*")) == []

    pictured = tmp_path / "pict.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ABC-2026-001"
    png = tmp_path / "dot.png"
    PILImage.new("RGB", (8, 8), "white").save(png)
    ws.add_image(XLImage(str(png)), "C3")
    wb.save(pictured)
    pictured_out = tmp_path / "pict-out"
    code = main([str(pictured), "-r", str(RULES), "-o", str(pictured_out)])
    captured = capsys.readouterr()
    assert code == 1
    assert MSG_PICTURE_SHEET in captured.err
    assert list(pictured_out.glob("*")) == []
