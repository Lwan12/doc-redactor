from pathlib import Path

from docx import Document

from doc_redactor.cli import redact_path
from doc_redactor.config import load_rules
from doc_redactor.docx_redactor import redact_docx
from helpers import RULES, assert_report_has_no_secrets


def _build(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = True
    section.first_page_header.paragraphs[0].text = "FIRST ABC-2026-001"
    section.header.paragraphs[0].text = "ODD Acme Pharma"
    section.even_page_header.paragraphs[0].text = "EVEN Fictizumab"
    section.first_page_footer.paragraphs[0].text = "FF Fictizumab"
    section.footer.paragraphs[0].text = "DF ABC-2026-001"
    section.even_page_footer.paragraphs[0].text = "EF Acme Pharma"

    document.add_paragraph("Project ABC-2026-001.")
    document.add_paragraph("Keep fictizumab in lowercase.")
    document.add_paragraph("Seen acme pharma today.")

    split = document.add_paragraph()
    split.add_run("ID ")
    bold = split.add_run("ABC-20")
    bold.bold = True
    middle = split.add_run("26")
    middle.italic = True
    tail = split.add_run("-001")
    tail.underline = True
    split.add_run(" end")

    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = "Site Acme Pharma"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Drug Fictizumab"
    document.save(path)


def _parts(document):
    chunks = []
    for section in document.sections:
        parts = (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        )
        for part in parts:
            if not part.is_linked_to_previous:
                chunks.extend(_container_text(part))
    return chunks


def _container_text(container) -> list[str]:
    chunks = [paragraph.text for paragraph in container.paragraphs]
    for table in container.tables:
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                marker = id(cell._tc)
                if marker in seen:
                    continue
                seen.add(marker)
                chunks.extend(_container_text(cell))
    return chunks


def test_docx_body_table_header_and_cross_run(tmp_path: Path):
    source = tmp_path / "notes.docx"
    _build(source)
    original = source.read_bytes()
    outdir = tmp_path / "out"
    output = outdir / "notes.redacted.docx"
    counts = redact_docx(source, output, load_rules(RULES))
    assert source.read_bytes() == original

    document = Document(output)
    body = "\n".join(_container_text(document))
    headers = "\n".join(_parts(document))
    combined = body + "\n" + headers
    assert "ABC-2026-001" not in combined
    assert "ABC-20" not in combined
    assert "26-001" not in combined
    assert "Acme Pharma" not in combined
    assert "acme pharma" not in combined
    assert "Fictizumab" not in combined
    assert "fictizumab" in combined
    assert "[PROTOCOL]" in combined
    assert "[SPONSOR]" in combined
    assert "[STUDY DRUG]" in combined
    assert "[REDACTED:PROJECT_ID]" not in combined

    split = document.paragraphs[3]
    assert split.text == "ID [PROTOCOL] end"
    label = next(run for run in split.runs if run.text == "[PROTOCOL]")
    assert label.bold
    assert any(run.text == "" for run in split.runs)

    nested_text = document.tables[0].cell(0, 0).tables[0].cell(0, 0).text
    assert nested_text == "Drug [STUDY DRUG]"
    assert document.tables[0].cell(0, 0).paragraphs[0].text == "Site [SPONSOR]"
    assert "FIRST [PROTOCOL]" in headers
    assert "ODD [SPONSOR]" in headers
    assert "EVEN [STUDY DRUG]" in headers
    assert "FF [STUDY DRUG]" in headers
    assert "DF [PROTOCOL]" in headers
    assert "EF [SPONSOR]" in headers

    assert counts == {"DRUG": 3, "PROJECT_ID": 4, "SPONSOR": 4}
    redact_path(source, outdir, load_rules(RULES))
    report = assert_report_has_no_secrets(outdir / "notes.report.json")
    assert report == counts
