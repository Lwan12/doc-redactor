"""Word tables: every cell visited once, placeholders stay in their own cell.

All data is fictional (Acme Pharma, Fictizumab, ABC-2026-001).
"""

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from doc_redactor.docx_redactor import redact_docx
from doc_redactor.engine import RegexRule, WordlistRule, build_engine
from doc_redactor.scan import scan_path, write_scan

LABELS = {"PROJECT_ID": "[PID]", "SPONSOR": "[SPONSOR]", "DRUG": "[DRUG]"}


def _engine():
    return build_engine(
        [RegexRule(name="PROJECT_ID", pattern=r"ABC-\d{4}-\d{3}", flags=0)],
        [
            WordlistRule(category="SPONSOR", terms=["Acme Pharma"], case_sensitive=True),
            WordlistRule(category="DRUG", terms=["Fictizumab"], case_sensitive=True),
        ],
    )


def _cell_texts(path: Path) -> list[str]:
    """Text of every physical w:tc (own paragraphs only), in XML order."""
    body = Document(str(path)).element.body
    out = []
    for tc in body.iter(qn("w:tc")):
        paras = tc.findall(qn("w:p"))
        out.append("\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in paras))
    return out


def _expected(texts: list[str]) -> list[str]:
    result = []
    for text in texts:
        text = text.replace("Acme Pharma", "[SPONSOR]").replace("Fictizumab", "[DRUG]")
        result.append(re.sub(r"ABC-\d{4}-\d{3}", "[PID]", text))
    return result


def _build(path: Path, big: int = 0) -> None:
    doc = Document()
    t = doc.add_table(rows=3, cols=3)
    data = [
        ["Acme Pharma", "plain a", "ABC-2026-001"],
        ["plain b", "Fictizumab x", "plain c"],
        ["pre Acme Pharma", "plain d", "ABC-2026-002"],
    ]
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            t.cell(r, c).text = value
    # gridSpan (horizontal merge)
    h = doc.add_table(rows=2, cols=3)
    merged = h.cell(0, 0).merge(h.cell(0, 1))
    merged.text = "span Acme Pharma"
    h.cell(0, 2).text = "keep h"
    h.cell(1, 0).text = "ABC-2026-003"
    h.cell(1, 1).text = "keep i"
    h.cell(1, 2).text = "Fictizumab"
    # vMerge (vertical merge)
    v = doc.add_table(rows=3, cols=2)
    vm = v.cell(0, 0).merge(v.cell(2, 0))
    vm.text = "vert Fictizumab"
    v.cell(0, 1).text = "keep j"
    v.cell(1, 1).text = "ABC-2026-004"
    v.cell(2, 1).text = "keep k"
    # nested
    n = doc.add_table(rows=1, cols=2)
    n.cell(0, 0).text = "outer Acme Pharma"
    inner = n.cell(0, 1).add_table(rows=2, cols=2)
    inner.cell(0, 0).text = "inner Fictizumab"
    inner.cell(0, 1).text = "keep l"
    inner.cell(1, 0).text = "keep m"
    inner.cell(1, 1).text = "ABC-2026-005"
    if big:
        b = doc.add_table(rows=big, cols=6)
        for r in range(big):
            for c in range(6):
                b.cell(r, c).text = f"ABC-2026-{(r * 6 + c) % 1000:03d}" if (r + c) % 3 == 0 else f"cell {r}-{c}"
    doc.save(str(path))


def _check(source: Path, dest: Path) -> None:
    before = _cell_texts(source)
    after = _cell_texts(dest)
    assert len(after) == len(before)
    assert after == _expected(before)
    joined = "\n".join(after)
    for secret in ("Acme Pharma", "Fictizumab", "ABC-2026-"):
        assert secret not in joined
    # merged cells: placeholder written exactly once
    assert after.count("span [SPONSOR]") == 1
    assert after.count("vert [DRUG]") == 1


def test_scan_then_write_tables(tmp_path: Path):
    src = tmp_path / "t.docx"
    _build(src)
    scan = scan_path(src, _engine())
    dest = tmp_path / "out.docx"
    write_scan(scan, dest, replacements=LABELS)
    _check(src, dest)


def test_scan_then_write_many_cells(tmp_path: Path):
    src = tmp_path / "big.docx"
    _build(src, big=120)
    scan = scan_path(src, _engine())
    dest = tmp_path / "out.docx"
    write_scan(scan, dest, replacements=LABELS)
    _check(src, dest)


def test_cli_redact_tables(tmp_path: Path):
    src = tmp_path / "big.docx"
    _build(src, big=120)
    engine = _engine()
    engine_label = engine.label
    engine.label = lambda cat: LABELS.get(cat) or engine_label(cat)
    dest = tmp_path / "out.docx"
    redact_docx(src, dest, engine)
    _check(src, dest)
