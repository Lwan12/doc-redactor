"""Pager order, preview scrolling, and the rules menu vs table headings."""

from __future__ import annotations

import os
import re
import sys
from pathlib import Path

import pytest

from doc_redactor.gui.result_view import PDF_PAGER, clamp_page, line_fully_inside, line_visible
from doc_redactor.gui.rules_view import place_more_menu, rects_overlap


def test_clamp_page_matches_prev_and_next():
    assert PDF_PAGER == (("‹", -1), ("›", 1))
    assert PDF_PAGER[0][1] < 0 < PDF_PAGER[-1][1]
    assert clamp_page(1, PDF_PAGER[0][1], 2) == 0
    assert clamp_page(0, PDF_PAGER[0][1], 2) == 0
    assert clamp_page(0, PDF_PAGER[-1][1], 2) == 1
    assert clamp_page(1, PDF_PAGER[-1][1], 2) == 1
    assert clamp_page(3, 1, 0) == 0


def test_place_more_menu_misses_headings_and_first_row():
    headings = [
        (20, 134, 100, 163),
        (100, 134, 230, 163),
        (230, 134, 280, 163),
        (280, 134, 710, 163),
        (710, 134, 840, 163),
        (840, 134, 950, 163),
    ]
    first_row = [
        (20, 163, 100, 191),
        (100, 163, 230, 191),
        (230, 163, 280, 191),
        (280, 163, 710, 191),
        (710, 163, 840, 191),
        (840, 163, 950, 191),
    ]
    action = (840, 163, 950, 400)
    x, y = place_more_menu(
        button=(860, 90, 80, 32),
        menu_size=(128, 30),
        window=(0, 0, 960, 640),
        headings=headings,
        avoid=[action, *first_row],
    )
    box = (x, y, x + 128, y + 30)
    for heading in headings:
        assert not rects_overlap(box, heading)
    for cell in first_row:
        assert not rects_overlap(box, cell)
    assert not rects_overlap(box, action)


def test_place_more_menu_stays_under_the_button_when_table_is_lower():
    headings = [(20, 170, 950, 199)]
    first = (20, 199, 950, 227)
    x, y = place_more_menu(
        button=(860, 90, 80, 32),
        menu_size=(128, 30),
        window=(0, 0, 960, 640),
        headings=headings,
        avoid=[first, (840, 199, 950, 400)],
    )
    assert (x, y) == (812, 122)
    box = (x, y, x + 128, y + 30)
    assert not rects_overlap(box, headings[0])
    assert not rects_overlap(box, first)


def test_place_more_menu_stays_under_the_button_when_nothing_is_below():
    x, y = place_more_menu(
        button=(100, 40, 70, 30),
        menu_size=(80, 24),
        window=(0, 0, 960, 640),
        headings=[],
    )
    assert (x, y) == (90, 70)


@pytest.fixture(scope="module")
def gui():
    pytest.importorskip("tkinter")
    if sys.platform == "win32" and not os.environ.get("DOC_REDACTOR_GUI_TESTS"):
        pytest.skip("window geometry tests run under xvfb (set DOC_REDACTOR_GUI_TESTS=1 to force)")
    from doc_redactor.gui.app import App, create_root

    try:
        root, dnd = create_root()
    except Exception as exc:
        pytest.skip(f"no display: {exc}")
    root.geometry("960x640+0+0")
    app = App(root, dnd, shoot=True)
    root.update_idletasks()
    root.deiconify()
    root.update()
    yield app
    try:
        root.destroy()
    except Exception:
        pass


def _settle(app) -> None:
    app.root.update_idletasks()
    app.root.update()


def _pager(app):
    pager = app.result_view._pdf_pager
    assert pager is not None and pager.winfo_exists()
    parts = pager.pack_slaves()
    assert len(parts) >= 3
    return parts


def test_pdf_pager_reads_prev_page_next_and_clicks_match(gui):
    app = gui
    app.apply_shoot_state("res-pdf")
    _settle(app)
    assert app.pdf_page == 1
    parts = _pager(app)
    texts = [part.cget("text") for part in parts]
    assert texts[0] == "‹"
    assert texts[-1] == "›"
    assert "第 2 / 2 页" in texts[1]
    xs = [part.winfo_rootx() for part in parts]
    assert xs[0] < xs[1] < xs[-1]

    parts[0].invoke()
    _settle(app)
    assert app.pdf_page == 0
    parts = _pager(app)
    assert "第 1 / 2 页" in parts[1].cget("text")
    parts[0].invoke()
    _settle(app)
    assert app.pdf_page == 0

    parts = _pager(app)
    assert parts[-1].cget("text") == "›"
    assert parts[-1].winfo_rootx() > parts[0].winfo_rootx()
    parts[-1].invoke()
    _settle(app)
    assert app.pdf_page == 1
    parts = _pager(app)
    parts[-1].invoke()
    _settle(app)
    assert app.pdf_page == 1
    parts = _pager(app)
    assert [part.cget("text") for part in parts][0] == "‹"
    assert parts[0].winfo_rootx() < parts[-1].winfo_rootx()


def test_pdf_preview_keeps_the_body_reachable(gui):
    app = gui
    app.apply_shoot_state("res-pdf")
    _settle(app)
    text = app.result_view._preview_text
    scroll = app.result_view._preview_scroll
    assert text is not None and text.winfo_exists()
    assert scroll is not None and scroll.winfo_ismapped()
    assert text.cget("yscrollcommand")
    assert scroll.cget("command")
    body = text.get("1.0", "end")
    assert "第 1 页" in body
    assert "第 2 页" in body
    assert "入组速度最快" in body
    root = app.root
    assert text.winfo_rooty() >= root.winfo_rooty()
    assert text.winfo_rooty() + text.winfo_height() <= root.winfo_rooty() + root.winfo_height() + 2
    # The text widget itself is not clipped by the preview footer.
    assert text.winfo_height() > 80
    first = text.search("第 1 页", "1.0")
    opened = text.search("第 2 页", "1.0")
    last = text.search("入组速度最快", "1.0")
    assert first and opened and last
    # Default view keeps the last line whole. The line at the top of the pane
    # is a full line, not a sliced glyph.
    assert line_fully_inside(text, last)
    top = text.dlineinfo(text.index("@0,0"))
    assert top is not None and top[1] >= 0
    assert line_visible(text, opened)
    text.yview_moveto(0)
    text.update_idletasks()
    assert line_visible(text, first)
    text.yview_moveto(1)
    text.update_idletasks()
    assert line_fully_inside(text, last)


def _assert_page_marker_at_top(app, page_number: int) -> None:
    text = app.result_view._preview_text
    assert text is not None and text.winfo_exists()
    label = f"第 {page_number} 页"
    idx = text.search(label, "1.0")
    assert idx
    assert line_visible(text, idx)
    assert line_fully_inside(text, idx)
    info = text.dlineinfo(idx)
    assert info is not None and info[1] <= 12
    top = text.get("@0,0 linestart", "@0,0 lineend")
    assert label in top
    if page_number == 1:
        assert text.yview()[0] == 0


def test_pdf_page_turn_scrolls_marker_to_the_top(gui):
    app = gui
    app.apply_shoot_state("res-pdf")
    _settle(app)
    assert app.pdf_page == 1
    parts = _pager(app)
    parts[0].invoke()
    _settle(app)
    assert app.pdf_page == 0
    _assert_page_marker_at_top(app, 1)

    parts = _pager(app)
    parts[-1].invoke()
    _settle(app)
    assert app.pdf_page == 1
    _assert_page_marker_at_top(app, 2)


def test_more_menu_does_not_cover_rule_headings(gui):
    app = gui
    app.apply_shoot_state("rules-more")
    _settle(app)
    menu = app.rules_view._more_menu
    tree = app.rules_view._rules_tree
    btn = app.rules_view._more_btn
    assert menu is not None and tree is not None and btn is not None
    width = menu.winfo_width()
    height = menu.winfo_height()
    assert width > 20 and height > 8
    box = (menu.winfo_rootx(), menu.winfo_rooty(), menu.winfo_rootx() + width, menu.winfo_rooty() + height)
    names = [tree.heading(col, "text") for col in tree.cget("columns")]
    rects = app.rules_view.heading_rects()
    assert names == ["启用", "规则名", "类型", "例子", "替换成", "操作"]
    assert len(rects) == len(names)
    for name, rect in zip(names, rects):
        assert not rects_overlap(box, rect), f"{name} {rect} overlaps menu {box}"
    kids = tree.get_children()
    assert kids
    first_rects = app.rules_view.first_row_rects()
    assert len(first_rects) == len(names)
    for name, rect in zip(names, first_rects):
        assert not rects_overlap(box, rect), f"first-row {name} {rect} overlaps menu {box}"
    values = tree.item(kids[0], "values")
    assert "修改" in values[-1]
    assert "删除" in values[-1]
    assert values[names.index("替换成")]
    for rect in app.rules_view.action_rects():
        assert not rects_overlap(box, rect), f"操作 {rect} overlaps menu {box}"
    assert box[1] >= btn.winfo_rooty() + btn.winfo_height()


def test_write_files_keeps_existing_output_and_source(gui, tmp_path: Path):
    """A same-named output already on disk is left alone. The write makes a new folder."""
    from docx import Document

    from doc_redactor.scan import KIND_WORD, FileScan, LocatedHit

    source = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("ABC-2026-001")
    document.save(source)
    original = source.read_bytes()
    out = tmp_path / "out"
    out.mkdir()
    prior_dir = out / "20200101-0000"
    prior_dir.mkdir()
    prior = prior_dir / "notes.redacted.docx"
    prior.write_bytes(b"prior-output-do-not-touch")

    app = gui
    app.show("main")
    _settle(app)
    app.output_var.set(str(out))
    app.scans = [
        FileScan(
            path=source,
            kind=KIND_WORD,
            status="ok",
            hits=[
                LocatedHit(
                    hit_id=1,
                    category="项目号",
                    text="ABC-2026-001",
                    location="第 1 段",
                    unit="p:0",
                    start=0,
                    end=12,
                )
            ],
            paragraphs=["ABC-2026-001"],
        )
    ]
    app.running = False
    app.written = False
    app.write_errors = {}
    app.batch_dir = None
    app.kept_outputs = []

    app.write_files()
    _settle(app)
    first_dir = app.batch_dir
    assert first_dir is not None
    assert first_dir.parent == out
    assert first_dir != prior_dir
    assert re.fullmatch(r"\d{8}-\d{4}(-\d+)?", first_dir.name)
    first = first_dir / "notes.redacted.docx"
    assert first.is_file()
    assert first.resolve() != source.resolve()
    assert first.resolve() != prior.resolve()
    assert prior.read_bytes() == b"prior-output-do-not-touch"
    assert source.read_bytes() == original
    first_bytes = first.read_bytes()

    app.write_files()
    _settle(app)
    second_dir = app.batch_dir
    assert second_dir is not None
    assert second_dir != first_dir
    assert second_dir.parent == out
    assert re.fullmatch(r"\d{8}-\d{4}(-\d+)?", second_dir.name)
    second = second_dir / "notes.redacted.docx"
    assert second.is_file()
    assert second.resolve() != first.resolve()
    assert first.read_bytes() == first_bytes
    assert prior.read_bytes() == b"prior-output-do-not-touch"
    assert source.read_bytes() == original


def test_res_skipped_shoot_is_written_not_conflict(gui):
    from doc_redactor.gui.result_view import _scan_status

    app = gui
    app.apply_shoot_state("res-skipped")
    _settle(app)
    assert app.written is True
    assert app.write_failure_rows() == []
    assert getattr(app, "_overlay_frame", None) is None
    notes = list(app.write_notes.values())
    assert notes
    assert all(note == "这个文件之前已写出，已跳过" for note in notes)
    kind, text, tag = _scan_status(app.scans[0], app)
    assert kind == "skipped"
    assert tag == "ok"
    assert "✕" not in text


def test_retry_existing_file_is_skipped_as_written(gui, tmp_path: Path):
    """Retry that meets an already-written dest skips it. Not an error, not red."""
    from docx import Document

    from doc_redactor.gui.result_view import _scan_status
    from doc_redactor.gui.theme import OK, status_style
    from doc_redactor.scan import KIND_WORD, FileScan, LocatedHit, MSG_ALREADY_WRITTEN

    source = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("ABC-2026-001")
    document.save(source)
    out = tmp_path / "out"
    out.mkdir()

    app = gui
    app.show("main")
    _settle(app)
    app.output_var.set(str(out))
    app.scans = [
        FileScan(
            path=source,
            kind=KIND_WORD,
            status="ok",
            hits=[
                LocatedHit(
                    hit_id=1,
                    category="项目号",
                    text="ABC-2026-001",
                    location="第 1 段",
                    unit="p:0",
                    start=0,
                    end=12,
                )
            ],
            paragraphs=["ABC-2026-001"],
        )
    ]
    app.running = False
    app.written = False
    app.write_errors = {}
    app.write_notes = {}
    app.batch_dir = None
    app.kept_outputs = []

    app.write_files()
    _settle(app)
    dest = app.batch_dir / "notes.redacted.docx"
    assert dest.is_file()
    kept = dest.read_bytes()
    assert app.written is True
    assert str(source) not in app.write_errors

    app.write_errors[str(source)] = "输出文件夹里这个文件正被别的程序占用"
    app.written = False
    app.write_files(retry=True)
    _settle(app)
    assert dest.read_bytes() == kept
    assert str(source) not in app.write_errors
    assert app.write_notes[str(source)] == MSG_ALREADY_WRITTEN
    assert app.written is True
    kind, text, tag = _scan_status(app.scans[0], app)
    assert kind == "skipped"
    assert text == MSG_ALREADY_WRITTEN
    assert tag == "ok"
    icon, label, color = status_style(kind)
    assert icon == "✓"
    assert label == MSG_ALREADY_WRITTEN
    assert color == OK
    assert tag != "bad"
    assert app.write_failure_rows() == []


def _walk(widget):
    yield widget
    for child in widget.winfo_children():
        yield from _walk(child)


def _visible_copy(widget) -> str:
    chunks = []
    for node in _walk(widget):
        try:
            value = node.cget("text")
        except Exception:
            continue
        # Hand-wrapped labels keep the unwrapped sentence here.
        value = getattr(node, "logical_text", value)
        if isinstance(value, str) and value:
            chunks.append(value)
    return "\n".join(chunks)


def test_blank_only_result_notice_and_column_wrap(gui):
    import tkinter as tk

    from doc_redactor.pdf_redactor import MSG_FIT_BLANK
    from doc_redactor.scan import KIND_PDF, FileScan, LocatedHit

    app = gui
    blank = FileScan(
        path=Path("blank_only.pdf"),
        kind=KIND_PDF,
        status="ok",
        hits=[
            LocatedHit(
                hit_id=1,
                category="申办方",
                text="I",
                location="第 1 页 · 第 1 行",
                unit="page:0",
                start=0,
                end=1,
                rects=[(72.0, 80.0, 82.0, 88.0)],
                page=0,
            )
        ],
        pages=[["I"]],
    )
    lost = FileScan(
        path=Path("lost.pdf"),
        kind=KIND_PDF,
        status="ok",
        hits=[
            LocatedHit(
                hit_id=2,
                category="申办方",
                text="Acme",
                location="第 1 页 · 第 1 行",
                unit="page:0",
                start=0,
                end=4,
                rects=[(72.0, 80.0, 280.0, 110.0)],
                page=0,
            )
        ],
        pages=[["Acme"]],
        unlocated_count=2,
    )
    app.scans = [blank, lost]
    app.selected = 0
    app.written = False
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)
    app.root.update_idletasks()
    app.root.update()

    notices = []
    for node in _walk(app.root):
        if not isinstance(node, tk.Label):
            continue
        try:
            text = node.cget("text")
        except tk.TclError:
            continue
        if "只留了空白" in text or "无法定位" in text:
            notices.append(node)
    assert notices
    assert MSG_FIT_BLANK.format(m=1) in _visible_copy(app.root)
    assert "2\u00a0处无法定位，未写出" in _visible_copy(app.root)
    assert "换成短码" not in _visible_copy(app.root)
    import tkinter.font as tkfont

    from doc_redactor.gui.result_view import NO_LINE_START

    for node in notices:
        # The label is laid out by hand so a comma never starts a line.
        lines = node.cget("text").split("\n")
        parent_w = node.master.winfo_width()
        assert parent_w > 80
        assert int(node.cget("wraplength")) == 0
        measure = tkfont.Font(font=node.cget("font")).measure
        for line in lines:
            assert line and line[0] not in NO_LINE_START
            assert measure(line) <= parent_w - 8
        assert "".join(lines).replace(" ", "") == node.logical_text.replace(" ", "")


def test_mixed_unlocated_row_warns_and_write_button_skips_it(gui):
    """A file with both placed and unplaced hits is warned before writing.

    The write button counts only files that will actually be written.
    """
    from tkinter import ttk

    from doc_redactor.gui.result_view import _scan_status
    from doc_redactor.gui.theme import status_style
    from doc_redactor.scan import KIND_PDF, FileScan, LocatedHit

    def _hit(hit_id: int, text: str) -> LocatedHit:
        return LocatedHit(
            hit_id=hit_id,
            category="申办方",
            text=text,
            location="第 1 页 · 第 1 行",
            unit="page:0",
            start=0,
            end=len(text),
            rects=[(72.0, 80.0, 140.0, 94.0)],
            page=0,
        )

    kept = FileScan(
        path=Path("kept.pdf"),
        kind=KIND_PDF,
        status="ok",
        hits=[_hit(1, "Acme")],
        pages=[["Acme"]],
    )
    mixed = FileScan(
        path=Path("mixed.pdf"),
        kind=KIND_PDF,
        status="ok",
        hits=[_hit(2, "星辰海")],
        pages=[["星辰海"]],
        unlocated_count=1,
    )
    app = gui
    app.scans = [kept, mixed]
    app.selected = 0
    app.written = False
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)

    kind, text, tag = _scan_status(mixed, app)
    assert kind == "unlocated"
    assert tag == "warn"
    assert text == "1 处无法定位"
    icon, _label, _color = status_style(kind)
    assert icon == "⚠"

    labels = []
    for node in _walk(app.root):
        if isinstance(node, ttk.Treeview) and tuple(node.cget("columns")) == ("name", "n"):
            labels += [node.item(item, "values")[0] for item in node.get_children()]
    mixed_row = next(row for row in labels if "mixed.pdf" in row)
    kept_row = next(row for row in labels if "kept.pdf" in row)
    assert mixed_row.startswith("⚠")
    assert "mixed.pdf" in mixed_row
    assert "1 处无法定位" in mixed_row
    assert mixed_row.index("mixed.pdf") < mixed_row.index("无法定位")
    assert kept_row.startswith("✓")
    assert "无法定位" not in kept_row

    buttons = [node.cget("text") for node in _walk(app.root) if isinstance(node, ttk.Button)]
    assert any(label.startswith("写出 1 个脱敏文件") for label in buttons)
    assert not any("写出 2 个" in label for label in buttons)
    write_label = next(label for label in buttons if label.startswith("写出 "))
    assert "替换 1 处" in write_label
    assert "替换 2 处" not in write_label


def test_pdf_fit_notices_use_the_product_copy(gui):
    from doc_redactor.pdf_redactor import MSG_FIT_BLANK, MSG_FIT_SHORT, MSG_PDF_WIDTH, MSG_REPL_HINT

    app = gui
    app.apply_shoot_state("res-pdf")
    _settle(app)
    copy = _visible_copy(app.root)
    assert MSG_FIT_SHORT.format(n=2) in copy
    assert MSG_FIT_BLANK.format(m=1) in copy
    assert "有\u00a02\u00a0处" in copy
    assert "另有\u00a01\u00a0处" in copy
    assert MSG_PDF_WIDTH in copy
    assert "正则" not in copy

    app.apply_shoot_state("rule-placeholder")
    _settle(app)
    rules = _visible_copy(app.root)
    assert MSG_REPL_HINT in rules
    assert "正则" not in rules

    app.apply_shoot_state("words-list")
    _settle(app)
    words = _visible_copy(app.root)
    assert MSG_REPL_HINT in words
    assert "正则" not in words

    app.apply_shoot_state("plan-list")
    _settle(app)
    plan = _visible_copy(app.root)
    assert MSG_REPL_HINT in plan
    assert "正则" not in plan

    app.apply_shoot_state("main-rules")
    _settle(app)
    main = _visible_copy(app.root)
    assert "截断" not in main
    assert "换成短码" in main
    assert "正则" not in main


def test_skipped_row_puts_the_file_name_before_the_note(gui):
    from tkinter import ttk

    from doc_redactor.scan import MSG_ALREADY_WRITTEN

    app = gui
    app.apply_shoot_state("res-skipped")
    _settle(app)
    labels = []
    for w in _walk(app.root):
        if isinstance(w, ttk.Treeview) and tuple(w.cget("columns")) == ("name", "n"):
            labels += [w.item(i, "values")[0] for i in w.get_children()]
    row = next(x for x in labels if MSG_ALREADY_WRITTEN in x)
    assert row.index(app.scans[0].path.name) < row.index(MSG_ALREADY_WRITTEN)


def _table_top(app) -> int:
    app.root.update_idletasks()
    app.root.update()
    return app.rules_view._rules_tree.winfo_rooty()


def _open_menu(app):
    rv = app.rules_view
    rv.post_more_menu()
    app.root.update()
    return rv


@pytest.mark.parametrize("how", ["unpost", "escape", "item", "close_gap"])
def test_more_menu_open_then_close_restores_table_top(gui, how):
    app = gui
    app.show_rules_page("list")
    _settle(app)
    base = _table_top(app)
    rv = _open_menu(app)
    assert _table_top(app) > base
    menu = rv._more_menu
    if how == "unpost":
        menu.unpost()
        menu.event_generate("<Unmap>")
    elif how == "escape":
        menu.event_generate("<Escape>")
        menu.unpost()
        menu.event_generate("<Unmap>")
    elif how == "item":
        called = []
        app.ask_restore_defaults = lambda: called.append(1)
        menu.unpost()
        menu.event_generate("<Unmap>")
    else:
        rv._close_menu_gap()
    app.root.update()
    assert _table_top(app) == base
    assert not rv._menu_gap.winfo_manager()
    # and it can open again after closing
    _open_menu(app)
    assert _table_top(app) > base
    rv._close_menu_gap()
    assert _table_top(app) == base


def _wait_scan(app, limit: float = 20.0) -> None:
    import time

    deadline = time.time() + limit
    app.root.update()
    while app.running and time.time() < deadline:
        app.root.update()
        time.sleep(0.02)
    for _ in range(5):
        app.root.update()
    assert not app.running


def test_relabel_only_keeps_unchecked_and_redraws_preview(gui, tmp_path, monkeypatch):
    from docx import Document

    import doc_redactor.gui.app as app_mod

    class _Inline:
        """Run the scan worker on the Tk thread; there is no mainloop in tests."""

        def __init__(self, target, daemon=None):
            self._target = target

        def start(self):
            self._target()

    monkeypatch.setattr(app_mod, "Thread", _Inline)

    from doc_redactor.ruleset import KIND_WORDLIST, Rule

    app = gui
    source = tmp_path / "memo.docx"
    document = Document()
    document.add_paragraph("申办方星辰海医药，合作方星辰海研究所。")
    document.save(source)
    app.ruleset.rules = [Rule(name="申办方", kind=KIND_WORDLIST, terms=["星辰海"], replacement="[甲方]")]
    app.files = [source]
    (tmp_path / "out").mkdir()
    app.output_var.set(str(tmp_path / "out"))
    app.output_error = None
    app.start_scan()
    _wait_scan(app)
    hits = app.scans[0].hits
    assert len(hits) == 2
    hits[1].enabled = False
    app._clear_preview()
    app.ruleset.rules[0].replacement = "[SPONSOR-X]"
    app.start_scan()
    _wait_scan(app)
    assert app.note_cleared is False
    assert [h.enabled for h in app.scans[0].hits] == [True, False]
    app.preview_mode = "after"
    app.show("result")
    _settle(app)
    body = app.result_view._preview_text.get("1.0", "end")
    assert "[SPONSOR-X]" in body
    assert "[甲方]" not in body
    shots = os.environ.get("DOC_REDACTOR_RELABEL_SHOT")
    if shots:
        from PIL import ImageGrab

        root = app.root
        menu = getattr(app.rules_view, "_more_menu", None)
        if menu is not None:
            menu.unpost()  # left open by an earlier test in this module
        root.update()
        x, y = root.winfo_rootx(), root.winfo_rooty()
        ImageGrab.grab(bbox=(x, y, x + root.winfo_width(), y + root.winfo_height())).save(shots)

    app.scans[0].hits[1].enabled = False
    app._clear_preview()
    app.ruleset.rules[0].terms = ["星辰海", "星海"]
    app.start_scan()
    _wait_scan(app)
    assert app.note_cleared is True
    assert all(h.enabled for h in app.scans[0].hits)


def _s2_scans():
    from doc_redactor.scan import KIND_PDF, FileScan, LocatedHit

    def _hit(hit_id: int, text: str) -> LocatedHit:
        return LocatedHit(
            hit_id=hit_id,
            category="申办方",
            text=text,
            location="第 1 页 · 第 1 行",
            unit="page:0",
            start=0,
            end=len(text),
            rects=[(72.0, 80.0, 140.0, 94.0)],
            page=0,
        )

    kept = FileScan(
        path=Path("kept.pdf"),
        kind=KIND_PDF,
        status="ok",
        hits=[_hit(1, "Acme"), _hit(2, "Acme"), _hit(3, "Acme")],
        pages=[["Acme Acme Acme"]],
    )
    lost = FileScan(
        path=Path("lost.pdf"),
        kind=KIND_PDF,
        status="ok",
        hits=[_hit(4, "星辰海"), _hit(5, "星辰海"), _hit(6, "星辰海")],
        pages=[["星辰海 星辰海 星辰海"]],
        unlocated_count=2,
    )
    return kept, lost


def test_writable_total_skips_blocked_and_failed_files():
    from doc_redactor.gui.result_view import writable_hit_total

    kept, lost = _s2_scans()
    assert writable_hit_total([kept, lost], written=False, write_errors={}) == 3
    kept.hits[0].enabled = False
    assert writable_hit_total([kept, lost], written=False, write_errors={}) == 2
    # After a write, a file the write refused no longer counts.
    errors = {str(kept.path): "占用"}
    assert writable_hit_total([kept, lost], written=True, write_errors=errors) == 0
    # Before writing, stale errors are ignored.
    assert writable_hit_total([kept, lost], written=False, write_errors=errors) == 2


def _summary_number(app, verb: str) -> str:
    import tkinter as tk

    for node in _walk(app.root):
        if isinstance(node, tk.Label) and node.cget("text").strip() == verb:
            siblings = node.master.pack_slaves()
            return siblings[siblings.index(node) + 1].cget("text")
    raise AssertionError(f"no {verb} label")


def test_summary_count_matches_the_write_button(gui):
    """S2: the top 将替换 N 处 used to include hits of a blocked file."""
    from tkinter import ttk

    kept, lost = _s2_scans()
    app = gui
    app.scans = [kept, lost]
    app.selected = 0
    app.written = False
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)

    buttons = [node.cget("text") for node in _walk(app.root) if isinstance(node, ttk.Button)]
    write_label = next(label for label in buttons if label.startswith("写出 "))
    assert "替换 3 处" in write_label
    assert _summary_number(app, "将替换") == "3"
    rows = []
    for node in _walk(app.root):
        if isinstance(node, ttk.Treeview) and tuple(node.cget("columns")) == ("name", "n"):
            rows += [node.item(item, "values") for item in node.get_children()]
    all_row = next(row for row in rows if "全部文件" in str(row[0]))
    assert str(all_row[1]) == "3"
    lost_row = next(row for row in rows if "lost.pdf" in str(row[0]))
    assert str(lost_row[1]) == "—"
    # Blocked hits are not hits the user cancelled.
    assert "你取消了" not in _visible_copy(app.root)

    # After a write the blocked file is a 没写出 row; it still adds nothing.
    app.written = True
    app.write_errors = {str(lost.path): "2 处无法定位，未写出"}
    app.show("result")
    _settle(app)
    rows = []
    for node in _walk(app.root):
        if isinstance(node, ttk.Treeview) and tuple(node.cget("columns")) == ("name", "n"):
            rows += [node.item(item, "values") for item in node.get_children()]
    assert str(next(row for row in rows if "全部文件" in str(row[0]))[1]) == "3"
    assert str(next(row for row in rows if "lost.pdf" in str(row[0]))[1]) == "—"
    assert _summary_number(app, "已替换") == "3"


_S3_TEXT = "⚠ 另有\u00a02\u00a0处太窄，连短码也放不下，只留了空白，原文已删除。 2\u00a0处无法定位，未写出"


@pytest.mark.parametrize("width", range(4, 60))
def test_wrapped_lines_never_start_with_a_comma(width):
    """S3: Tk breaks a long CJK run anywhere, so 「，」 could start a line."""
    from doc_redactor.gui.result_view import NO_LINE_START, wrap_lines

    lines = wrap_lines(_S3_TEXT, width, len)
    for line in lines:
        assert line, lines
        assert line[0] not in NO_LINE_START, lines
        assert not line.startswith(" ") and not line.endswith(" ")
    # Only normal spaces at a break are dropped; nothing else is lost.
    assert "".join(lines).replace(" ", "") == _S3_TEXT.replace(" ", "")
    for line in lines:
        assert not line.endswith("\u00a0") and not line.startswith("\u00a0")
    # The count and 处 stay together.
    assert all("2\u00a0处" in line or "2" not in line for line in lines)


def test_wrap_lines_fits_the_width_when_it_can():
    from doc_redactor.gui.result_view import wrap_lines

    for width in range(8, 60):
        for line in wrap_lines(_S3_TEXT, width, len):
            # A glued unit (count + 处, or a char plus its comma) may hang by 1.
            assert len(line) <= width + 1


@pytest.mark.parametrize("padx", [6, 8])
@pytest.mark.parametrize("width", [180, 233, 260, 297, 320, 341])
def test_wrap_label_fits_inside_its_pack_padding(gui, padx, width):
    """The middle column packs the notice with padx; no glyph may be clipped."""
    import tkinter as tk

    from doc_redactor.gui.result_view import _wrap_label
    from doc_redactor.pdf_redactor import MSG_FIT_BLANK, MSG_FIT_SHORT

    app = gui
    host = tk.Toplevel(app.root)
    try:
        frame = tk.Frame(host, width=width, height=200)
        frame.pack_propagate(False)
        frame.pack()
        text = MSG_FIT_SHORT.format(n=3) + " " + MSG_FIT_BLANK.format(m=2)
        label = _wrap_label(frame, text, bg="white", fg="black", font=app.fonts["small"])
        label.pack(fill="x", padx=padx)
        for _ in range(3):
            host.update_idletasks()
            host.update()
        assert frame.winfo_width() == width
        # Everything Tk wants to draw fits in the space the pack padding leaves.
        assert label.winfo_reqwidth() <= width - 2 * padx, (label.cget("text"), label.winfo_reqwidth())
    finally:
        host.destroy()


def test_unwritten_file_is_not_ticked_after_write(gui):
    """After a write, hits of a file that was not written say 未写出, not ✓."""
    import tkinter as tk
    from tkinter import ttk

    kept, lost = _s2_scans()
    app = gui
    app.scans = [kept, lost]
    app.selected = "all"
    app.selected_cat = None
    app.written = True
    app.write_errors = {str(lost.path): "2 处无法定位，未写出"}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)

    marks = {}
    cat_rows = []
    for node in _walk(app.root):
        if isinstance(node, ttk.Treeview) and tuple(node.cget("columns")) == ("on", "loc", "cat", "text"):
            for item in node.get_children():
                marks[int(item)] = str(node.item(item, "values")[0])
        if isinstance(node, ttk.Treeview) and tuple(node.cget("columns")) == ("cat", "n"):
            cat_rows += [tuple(str(v) for v in node.item(item, "values")) for item in node.get_children()]
    lost_ids = {hit.hit_id for hit in lost.hits}
    assert marks and all(marks[i] == "未写出" for i in lost_ids), marks
    assert all(marks[h.hit_id] == "✓" for h in kept.hits), marks
    # The category total matches 已替换; the unwritten hits are listed apart.
    assert ("合计", "3") in cat_rows, cat_rows
    assert ("未写出", "3") in cat_rows, cat_rows
    # The middle matrix names the file that was not written and leaves it out.
    copy = [getattr(n, "logical_text", n.cget("text")) for n in _walk(app.root) if isinstance(n, tk.Label)]
    assert "lost.pdf 未写出，不计入合计" in copy



def _r15_scans():
    """Three PDFs like the #26 review: one blocked, one refused, one written."""
    from doc_redactor.scan import KIND_PDF, FileScan, LocatedHit

    def _hit(hit_id: int, text: str, line: int) -> LocatedHit:
        return LocatedHit(
            hit_id=hit_id,
            category="申办方",
            text=text,
            location=f"第 1 页 · 第 {line} 行",
            unit="page:0",
            start=0,
            end=len(text),
            rects=[(72.0, 80.0, 140.0, 94.0)],
            page=0,
        )

    lost = FileScan(
        path=Path("lost.pdf"),
        kind=KIND_PDF,
        status="ok",
        hits=[_hit(11, "星辰海", 1), _hit(12, "星辰海", 2), _hit(13, "星辰海", 3)],
        pages=[["星辰海 星辰海 星辰海"]],
        unlocated_count=2,
    )
    swap = FileScan(
        path=Path("swap.pdf"),
        kind=KIND_PDF,
        status="ok",
        hits=[_hit(21, "Acme", 1), _hit(22, "Acme", 2)],
        pages=[["Acme Pharma", "Acme Pharma"]],
    )
    other = FileScan(
        path=Path("other.pdf"),
        kind=KIND_PDF,
        status="ok",
        hits=[_hit(31, "Acme", 1), _hit(32, "Acme", 2)],
        pages=[["Acme Pharma", "Acme Pharma"]],
    )
    return lost, swap, other


def _show_r15(app, *, written: bool, selected="all"):
    lost, swap, other = _r15_scans()
    swap.hits[1].enabled = False
    app.scans = [lost, swap, other]
    app.selected = selected
    app.selected_cat = None
    app.current_hit = None
    app.written = written
    app.write_errors = (
        {str(lost.path): "2 处无法定位，未写出", str(other.path): "模拟输出文件被占用"} if written else {}
    )
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)
    _settle(app)
    return lost, swap, other


def _loc_tree(app):
    from tkinter import ttk

    for node in _walk(app.root):
        if isinstance(node, ttk.Treeview) and tuple(node.cget("columns")) == ("on", "loc", "cat", "text"):
            return node
    raise AssertionError("no location list")


def _loc_header(app) -> str:
    import tkinter as tk

    for node in _walk(app.root):
        if isinstance(node, tk.Label) and str(node.cget("text")).startswith("被替换的位置"):
            return node.cget("text")
    raise AssertionError("no 被替换的位置 header")


def test_location_counts_share_the_top_count():
    """#26 S1: the header counted every listed hit, written or not."""
    from doc_redactor.gui.result_view import location_header, unwritten_hit_total, writable_hit_total

    lost, swap, other = _r15_scans()
    swap.hits[1].enabled = False
    scans = [lost, swap, other]
    errors = {str(other.path): "占用"}
    assert writable_hit_total(scans, written=True, write_errors=errors) == 1
    assert unwritten_hit_total(scans, written=True, write_errors=errors) == 5
    assert location_header(scans, written=True, write_errors=errors) == "被替换的位置 1（另 5 处未写出）"
    # Category filter: the same count, narrowed.
    assert writable_hit_total(scans, written=True, write_errors=errors, category="药名") == 0
    assert location_header(scans, written=True, write_errors=errors, category="药名") == "被替换的位置 0"
    # Nothing left out: no note.
    assert location_header([swap], written=True, write_errors=errors) == "被替换的位置 1"
    # One file that was not written.
    assert location_header([other], written=True, write_errors=errors) == "被替换的位置 0（另 2 处未写出）"
    # Before a write only the blocked file is left out; stale errors are ignored.
    assert location_header(scans, written=False, write_errors=errors) == "被替换的位置 3（另 3 处未写出）"


@pytest.mark.parametrize("written", [False, True])
def test_location_header_matches_the_top_summary(gui, written):
    app = gui
    _show_r15(app, written=written)
    verb = "已替换" if written else "将替换"
    top = _summary_number(app, verb)
    header = _loc_header(app)
    assert header.startswith(f"被替换的位置 {top}"), (top, header)
    if written:
        assert header == "被替换的位置 1（另 5 处未写出）"
    else:
        assert header == "被替换的位置 3（另 3 处未写出）"


def _full_rows(tree) -> int:
    tree.update_idletasks()
    height = tree.winfo_height()
    full = 0
    for item in tree.get_children():
        box = tree.bbox(item)
        if box and box[1] >= 0 and box[1] + box[3] <= height:
            full += 1
    return full


@pytest.mark.parametrize("written", [False, True])
def test_location_list_shows_three_full_rows_at_960x640(gui, written):
    """#26 S2: at 960×640 the category table pushed the list down to 1½ rows."""
    app = gui
    app.root.geometry("960x640+0+0")
    _settle(app)
    _show_r15(app, written=written)
    assert app.root.winfo_height() == 640
    assert _full_rows(_loc_tree(app)) >= 3


def test_category_table_keeps_all_rows_reachable_when_short(gui):
    """A shortened category table scrolls; no row is dropped."""
    from tkinter import ttk

    app = gui
    app.root.geometry("960x640+0+0")
    _settle(app)
    _show_r15(app, written=True)
    cat_tree = next(
        n for n in _walk(app.root) if isinstance(n, ttk.Treeview) and tuple(n.cget("columns")) == ("cat", "n")
    )
    rows = [tuple(str(v) for v in cat_tree.item(i, "values")) for i in cat_tree.get_children()]
    assert ("合计", "1") in rows and ("未写出", "5") in rows
    if int(cat_tree.cget("height")) < len(rows):
        bars = [n for n in cat_tree.master.winfo_children() if isinstance(n, ttk.Scrollbar)]
        assert bars and bars[0].winfo_ismapped()


def test_category_rows_that_fit_leaves_room_for_the_list():
    from doc_redactor.gui.result_view import CAT_MIN_ROWS, LOC_MIN_ROWS, category_rows_that_fit

    # Tall pane: every wanted row.
    assert category_rows_that_fit(1000, row_h=26, overhead=28, wanted=7) == 7
    # Exactly enough for 5 category rows plus LOC_MIN_ROWS list rows.
    space = 2 * 28 + (5 + LOC_MIN_ROWS) * 26
    assert category_rows_that_fit(space, row_h=26, overhead=28, wanted=7) == 5
    assert category_rows_that_fit(space - 1, row_h=26, overhead=28, wanted=7) == 4
    # Very short: the table keeps its floor and scrolls.
    assert category_rows_that_fit(50, row_h=26, overhead=28, wanted=7) == CAT_MIN_ROWS
    # Never more than wanted, even when wanted is below the floor.
    assert category_rows_that_fit(50, row_h=26, overhead=28, wanted=2) == 2


@pytest.mark.parametrize("mode", ["tag", "auto", "mixed", "none"])
def test_location_list_shows_three_full_rows_with_header_footer_on_at_960x640(gui, mode):
    """#28 + #30: with 「去掉页眉页脚」 on and after writing, the header/footer
    lines in the right pane share the height and the list keeps 3 full rows."""
    from doc_redactor.gui.demo import demo_hf_scans

    app = gui
    app.root.geometry("960x640+0+0")
    _settle(app)
    scans = demo_hf_scans(mode)
    app.scans = scans
    app.remove_hf = True
    app.output_var.set("/tmp/out")
    app.selected_cat = None
    app.current_hit = None
    app.written = True
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    for selected in ["all"] + [i for i, s in enumerate(scans) if s.status != "fail" and s.visible_hits]:
        app.selected = selected
        app.show("result")
        _settle(app)
        _settle(app)
        assert app.root.winfo_height() == 640
        assert _full_rows(_loc_tree(app)) >= 3, (mode, selected)


def test_header_footer_lines_fold_then_hide_only_when_needed():
    from doc_redactor.gui.result_view import HF_FOLDED, HF_FULL, HF_HIDDEN, hf_level

    heights = [(56, 31, 0), (31, 31, 0)]  # two-line value, one-line value
    need = 2 * 30 + (3 + 3) * 24
    assert hf_level(need + 87, heights, row_h=24, overhead=30, wanted=7) == HF_FULL
    assert hf_level(need + 62, heights, row_h=24, overhead=30, wanted=7) == HF_FOLDED
    assert hf_level(need + 40, heights, row_h=24, overhead=30, wanted=7) == HF_HIDDEN
    assert hf_level(need, [], row_h=24, overhead=30, wanted=7) == HF_FULL


def _inside(widget, outer) -> bool:
    """Widget is drawn and lies wholly inside ``outer`` (on screen)."""
    if not widget.winfo_ismapped() or widget.winfo_height() <= 1 or widget.winfo_width() <= 1:
        return False
    x, y = widget.winfo_rootx(), widget.winfo_rooty()
    ox, oy = outer.winfo_rootx(), outer.winfo_rooty()
    return (
        x >= ox
        and y >= oy
        and x + widget.winfo_width() <= ox + outer.winfo_width()
        and y + widget.winfo_height() <= oy + outer.winfo_height()
    )


@pytest.mark.parametrize(
    "key",
    ["res-hf-auto", "res-hf-auto-ok", "res-hf-auto-undo", "res-hf-mixed", "res-hf-mixed-ok", "res-hf-fixed", "res-hf-fixed-ok"],
)
def test_confirm_bar_keeps_preview_pager_and_radios_visible_at_960x640(gui, key):
    """#28 r2 M3: at 960×640 the 「待确认」 strip squeezed its text into a column
    three characters wide and pushed the radios, the pager and the preview
    out of the middle column. Seeing the preview is what makes 确认删除 safe."""
    import tkinter.font as tkfont
    from tkinter import ttk

    app = gui
    app.root.geometry("960x640+0+0")
    _settle(app)
    app.apply_shoot_state(key)
    _settle(app)
    _settle(app)
    assert app.root.winfo_height() == 640
    view = app.result_view
    text = view._preview_text
    assert text is not None
    middle = text
    while middle.master is not None and middle.master.winfo_width() < 600:
        middle = middle.master
    # The preview shows at least five full lines.
    line = tkfont.Font(font=text.cget("font")).metrics("linespace")
    assert text.winfo_ismapped() and text.winfo_height() >= 5 * line, (key, text.winfo_height(), line)
    assert _inside(text, app.root), key
    # Radios and the pager are on screen.
    radios = [w for w in _walk(middle) if isinstance(w, ttk.Radiobutton)]
    assert len(radios) == 2
    for widget in radios + list(view._pdf_pager.pack_slaves()):
        assert _inside(widget, middle), (key, widget)
    # The strip text wraps across the column: no line squeezed to a few characters.
    bars = [w for w in _walk(middle) if getattr(w, "hf_bar", None)]
    assert len(bars) == 1
    label = bars[0]
    assert _inside(label, middle), key
    lines = label.cget("text").split("\n")
    assert all(len(row) >= 10 for row in lines[:-1]), (key, [len(row) for row in lines])
    assert label.winfo_width() >= middle.winfo_width() * 0.7, key


def test_output_hint_is_not_covered_by_the_buttons_at_960x640(gui):
    """#28 r2 m5: 「每次写出都放进新的文件夹」 ran under the buttons and lost a character."""
    from tkinter import ttk

    app = gui
    app.root.geometry("960x640+0+0")
    _settle(app)
    app.apply_shoot_state("res-hf-auto")
    _settle(app)
    _settle(app)
    hints = [
        w
        for w in _walk(app.result_view.frame)
        if w.winfo_class() == "Label" and "不会盖掉上一次" in str(getattr(w, "logical_text", w.cget("text")))
    ]
    assert len(hints) == 1
    hint = hints[0]
    assert "".join(hint.cget("text").split("\n")).endswith("不会盖掉上一次。")
    assert hint.winfo_reqwidth() <= hint.winfo_width()
    buttons = [
        w for w in _walk(app.result_view.frame) if isinstance(w, ttk.Button) and w.cget("text").startswith("返回修改")
    ]
    assert len(buttons) == 1
    assert hint.winfo_rootx() + hint.winfo_width() <= buttons[0].winfo_rootx()
    assert _inside(hint, app.root)


def _m34_hit(hit_id: int, text: str, unit: str):
    from doc_redactor.scan import LocatedHit

    return LocatedHit(
        hit_id=hit_id, category="申办方", text=text, location="第 1 页", unit=unit,
        start=0, end=len(text), rects=[(72.0, 80.0, 140.0, 94.0)], page=0,
    )


def test_no_hit_file_after_write_is_not_shown_as_written(gui):
    """M34-1：0 命中的文件写出后，列表里不能是绿色「✓ 已写出，已跳过」。"""
    from tkinter import ttk

    from doc_redactor.gui.result_view import _scan_status
    from doc_redactor.leak_check import MSG_NO_HITS
    from doc_redactor.scan import KIND_PDF, FileScan

    line = "没有命中，没有写出脱敏件。不要把原件当成脱敏件使用"
    written = FileScan(
        path=Path("kept.pdf"), kind=KIND_PDF, status="ok", hits=[_m34_hit(1, "Acme", "page:0")], pages=[["Acme"]]
    )
    clean = FileScan(path=Path("appendix.pdf"), kind=KIND_PDF, status="nohit", hits=[], pages=[["nothing"]])
    app = gui
    app.scans = [written, clean]
    app.selected = 1
    app.written = True
    app.write_errors = {}
    app.write_notes = {str(clean.path): MSG_NO_HITS}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)

    kind, _text, tag = _scan_status(clean, app)
    assert kind != "skipped" and tag != "ok"
    rows = []
    for node in _walk(app.root):
        if isinstance(node, ttk.Treeview) and tuple(node.cget("columns")) == ("name", "n"):
            rows += [(node.item(item, "values"), node.item(item, "tags")) for item in node.get_children()]
    values, tags = next((v, t) for v, t in rows if "appendix.pdf" in v[0])
    assert not values[0].startswith("✓"), values
    assert "已写出" not in values[0] and "已跳过" not in values[0]
    assert "ok" not in tags
    copy = _visible_copy(app.root)
    assert line in copy
    assert "之前已写出" not in copy


def _tree_values(root) -> str:
    from tkinter import ttk

    out = []
    for node in _walk(root):
        if isinstance(node, ttk.Treeview):
            out += [" | ".join(str(v) for v in node.item(item, "values")) for item in node.get_children()]
    return "\n".join(out)


def test_approximate_ocr_hit_is_marked_and_shows_what_ocr_read(gui):
    """M34-2：近似命中标「近似」并显示 OCR 读出的字；这几个字不进命中清单。"""
    import json

    from doc_redactor.backpath import hit_list_payload
    from doc_redactor.scan import KIND_PDF, FileScan, LocatedHit

    hit = LocatedHit(
        hit_id=1, category="DRUG", text="Fictlzumab", location="第 1 页图 1 · 图内 第 1 行",
        unit="pdfimg:5", boxes=[(10, 10, 100, 20)], page=0, approx=True,
    )
    exact = LocatedHit(
        hit_id=2, category="DRUG", text="Fictizumab", location="第 1 页图 1 · 图内 第 2 行",
        unit="pdfimg:5", boxes=[(10, 40, 100, 20)], page=0,
    )
    scan = FileScan(path=Path("scan.pdf"), kind=KIND_PDF, status="ok", hits=[hit, exact], pages=[[""]])
    app = gui
    app.scans = [scan]
    app.selected = 0
    app.written = False
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)
    copy = _tree_values(app.root)
    assert copy.count("近似：读成") == 1
    assert "近似：读成 Fictlzumab" in _hover_tip(app, _hit_table(app.root), "1", "text")
    payload = json.dumps(hit_list_payload([scan]), ensure_ascii=False)
    assert "Fictlzumab" not in payload and "近似" not in payload
    app.hide_text = True
    app.show("result")
    _settle(app)
    hidden = _tree_values(app.root)
    assert "近似" in hidden and "Fictlzumab" not in hidden
    assert "Fictlzumab" not in _hover_tip(app, _hit_table(app.root), "1", "text")


def test_zero_hit_file_written_for_header_footer_is_not_shown_as_unwritten(gui, tmp_path):
    """M1（#34 r3）：0 命中但开了去页眉页脚、文件已写出时，不能说「无命中」「没有写出」。"""
    from doc_redactor.engine import WordlistRule, build_engine
    from doc_redactor.gui.result_view import _scan_status
    from doc_redactor.leak_check import MSG_NO_HITS_LINE
    from doc_redactor.scan import scan_path, write_scans

    sys.path.insert(0, str(Path(__file__).parent))
    from test_header_footer import _tagged_pdf

    src = tmp_path / "zero-hf.pdf"
    src.write_bytes(_tagged_pdf())
    engine = build_engine([], [WordlistRule("药名", ["Nothingmab"], True)])
    scan = scan_path(src, engine, remove_hf=True)
    assert not scan.hits and scan.hf_active
    [res] = write_scans([scan], tmp_path / "out", engine=engine)
    assert res.dest is not None and not res.note

    app = gui
    app.scans = [scan]
    app.selected = 0
    app.written = True
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)

    kind, text, _tag = _scan_status(scan, app)
    assert kind not in ("nohit", "nohit_unwritten"), (kind, text)
    assert "写出" in text and "没有写出" not in text and "未写出" not in text
    copy = _visible_copy(app.root)
    assert MSG_NO_HITS_LINE not in copy
    assert "没有写出" not in copy
    assert "已写出" in copy


@pytest.mark.parametrize("size", ["960x640", "1100x720"])
@pytest.mark.parametrize("key", ["res-hf-auto", "res-pdf", "res-word"])
def test_locate_hint_under_the_preview_is_not_clipped(gui, key, size):
    """#28 r3 m11: the hint under the preview ended flush with the column's
    right edge and had no closing 「。」, so it looked cut off. Measured: it was
    whole. Now it ends with 「。」 and wraps; this guards that it stays whole."""
    from doc_redactor.gui.result_view import LOCATE_HINT

    app = gui
    app.root.geometry(f"{size}+0+0")
    _settle(app)
    app.current_hit = None
    app.apply_shoot_state(key)
    _settle(app)
    _settle(app)
    found = [w for w in _walk(app.result_view.frame) if w.winfo_class() == "Message" and w.cget("text") == LOCATE_HINT]
    assert len(found) == 1, key
    message = found[0]
    frame = message.master
    assert message.winfo_reqwidth() <= frame.winfo_width(), (key, size, message.winfo_reqwidth(), frame.winfo_width())
    assert message.winfo_rootx() + message.winfo_width() <= frame.winfo_rootx() + frame.winfo_width()
    assert _inside(message, app.root)
    # Every character is drawn: the wrapped lines need no more than its height.
    import tkinter.font as tkfont

    font = tkfont.Font(font=message.cget("font"))
    lines = -(-font.measure(LOCATE_HINT) // max(1, int(message.cget("width"))))
    assert message.winfo_height() >= lines * font.metrics("linespace"), (key, size)
    # One line at 960 and up: no lone 「。」 wrapped onto a second line.
    assert font.measure(LOCATE_HINT) <= int(message.cget("width")), (key, size)
    app.root.geometry("960x640+0+0")
    _settle(app)


def test_fixed_rows_hint_is_on_the_confirm_strip(gui):
    app = gui
    app.root.geometry("960x640+0+0")
    _settle(app)
    app.apply_shoot_state("res-hf-fixed")
    _settle(app)
    bars = [w for w in _walk(app.result_view.frame) if getattr(w, "hf_bar", None)]
    assert len(bars) == 1
    assert "含 2 行固定文字" in bars[0].logical_text
    # Wrapped, nothing lost: the shown lines join back to the sentence.
    assert "".join(bars[0].cget("text").split("\n")).replace(" ", "") == bars[0].logical_text.replace(" ", "")


def _plain_hit(hit_id: int, text: str, unit: str):
    from doc_redactor.scan import LocatedHit

    return LocatedHit(
        hit_id=hit_id,
        category="申办方",
        text=text,
        location="第 1 段",
        unit=unit,
        start=0,
        end=len(text),
        rects=[(72.0, 80.0, 140.0, 94.0)],
        page=0,
    )


def test_blocked_file_warns_and_write_button_skips_it(gui):
    """A file with an image the tool cannot check is warned and not counted."""
    from tkinter import ttk

    from doc_redactor.gui.result_view import _scan_status
    from doc_redactor.image_redactor import MSG_NO_OCR
    from doc_redactor.scan import KIND_PDF, KIND_WORD, FileScan

    reason = MSG_NO_OCR.format(n=1)
    kept = FileScan(
        path=Path("kept.pdf"), kind=KIND_PDF, status="ok", hits=[_plain_hit(1, "Acme", "page:0")], pages=[["Acme"]]
    )
    blocked = FileScan(
        path=Path("blocked.docx"),
        kind=KIND_WORD,
        status="ok",
        hits=[_plain_hit(2, "星辰海", "para:0")],
        paragraphs=["星辰海"],
        blocked_reason=reason,
    )
    app = gui
    app.scans = [kept, blocked]
    app.selected = 1
    app.written = False
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)

    kind, _text, tag = _scan_status(blocked, app)
    assert (kind, tag) == ("blocked", "warn")
    labels = []
    for node in _walk(app.root):
        if isinstance(node, ttk.Treeview) and tuple(node.cget("columns")) == ("name", "n"):
            labels += [node.item(item, "values") for item in node.get_children()]
    row = next(values for values in labels if "blocked.docx" in values[0])
    assert row[0].startswith("⚠")
    assert str(row[1]) == "—"
    copy = _visible_copy(app.root)
    assert reason in copy
    assert "正则" not in copy
    buttons = [node.cget("text") for node in _walk(app.root) if isinstance(node, ttk.Button)]
    write_label = next(label for label in buttons if label.startswith("写出 "))
    assert write_label.startswith("写出 1 个脱敏文件")
    assert "替换 1 处" in write_label


def test_word_comment_notice_before_and_after_write(gui):
    from doc_redactor.scan import KIND_WORD, FileScan

    scan = FileScan(
        path=Path("notes.docx"),
        kind=KIND_WORD,
        status="ok",
        hits=[_plain_hit(1, "星辰海", "para:0")],
        paragraphs=["星辰海"],
        comment_count=2,
    )
    app = gui
    app.scans = [scan]
    app.selected = 0
    app.written = False
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)
    assert "写出时会删除 2 条批注" in _visible_copy(app.root)

    app.written = True
    app.show("result")
    _settle(app)
    copy = _visible_copy(app.root)
    assert "已删除 2 条批注" in copy
    assert "写出时会删除" not in copy


def test_blocked_no_tesseract_says_why_and_what_to_do(gui):
    """没装 Tesseract：空状态写清为什么不写出、怎么办，不叫用户把图换成普通图片。"""
    from doc_redactor.image_redactor import MSG_NO_OCR
    from doc_redactor.scan import KIND_WORD, FileScan

    blocked = FileScan(
        path=Path("needs-ocr.docx"),
        kind=KIND_WORD,
        status="ok",
        hits=[_plain_hit(2, "星辰海", "para:0")],
        paragraphs=["星辰海"],
        blocked_reason=MSG_NO_OCR.format(n=2),
    )
    app = gui
    app.scans = [blocked]
    app.selected = 0
    app.written = False
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)
    copy = _visible_copy(app.root)
    assert "图里的字" in copy
    assert "装好 Tesseract" in copy and "TESSERACT_CMD" in copy
    assert "换成普通图片" not in copy


def test_blocked_zero_hit_header_footer_file_is_not_shown_as_will_write(gui, tmp_path):
    """merge-r5 M1：被拦下、0 命中、有页眉页脚的文件，写出前不能显示绿 ✓「将删除页眉页脚并写出」。"""
    from tkinter import ttk

    from doc_redactor.engine import WordlistRule, build_engine
    from doc_redactor.gui.result_view import _scan_status
    from doc_redactor.image_redactor import MSG_NO_OCR
    from doc_redactor.scan import scan_path

    sys.path.insert(0, str(Path(__file__).parent))
    from test_header_footer import _tagged_pdf

    src = tmp_path / "blocked-zero-hf.pdf"
    src.write_bytes(_tagged_pdf())
    engine = build_engine([], [WordlistRule("药名", ["Nothingmab"], True)])
    scan = scan_path(src, engine, remove_hf=True)
    assert not scan.hits and scan.hf_active
    scan.blocked_reason = MSG_NO_OCR.format(n=1)

    app = gui
    app.scans = [scan]
    app.selected = 0
    app.written = False
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)

    kind, text, tag = _scan_status(scan, app)
    assert kind == "blocked" and tag != "ok", (kind, text, tag)
    rows = []
    for node in _walk(app.root):
        if isinstance(node, ttk.Treeview) and tuple(node.cget("columns")) == ("name", "n"):
            rows += [node.item(item, "values") for item in node.get_children()]
    row = next(v for v in rows if src.name in v[0])
    assert not row[0].startswith("✓"), row
    assert "将删除页眉页脚并写出" not in _visible_copy(app.root)


# ------------------------------------------------------------- 收尾小 PR

def _file_tree(root):
    from tkinter import ttk

    return next(
        node for node in _walk(root)
        if isinstance(node, ttk.Treeview) and tuple(node.cget("columns")) == ("name", "n")
    )


def _show_one(app, scans, written):
    app.scans = scans
    app.selected = 0
    app.written = written
    app.write_errors = {}
    app.write_notes = {}
    app.preview_mode = "orig"
    app.hide_text = False
    app.show("result")
    _settle(app)


def _hf_zero_scan(tmp_path, name="zero-hf.pdf"):
    from doc_redactor.engine import WordlistRule, build_engine
    from doc_redactor.scan import scan_path

    sys.path.insert(0, str(Path(__file__).parent))
    from test_header_footer import _tagged_pdf

    src = tmp_path / name
    src.write_bytes(_tagged_pdf())
    scan = scan_path(src, build_engine([], [WordlistRule("药名", ["Nothingmab"], True)]), remove_hf=True)
    assert not scan.hits and scan.hf_active
    return scan


def test_selected_file_row_keeps_its_status_colour(gui, tmp_path):
    """#34 次项：选中的「只删页眉页脚」那行 ✓ 不再变黑，和其他 ✓ 行同为绿色。"""
    from tkinter import ttk

    scan = _hf_zero_scan(tmp_path)
    _show_one(gui, [scan], True)
    tree = _file_tree(gui.root)
    [row] = [i for i in tree.get_children() if scan.path.name in tree.item(i, "values")[0]]
    assert tree.item(row, "values")[0].startswith("✓")
    assert "ok" in tree.item(row, "tags")
    assert row in tree.selection()
    style = tree.cget("style") or "Treeview"
    forced = [state for *state, _colour in ttk.Style().map(style, "foreground") if "selected" in state]
    assert forced == [], f"selected rows lose their tag colour: {forced}"


def test_top_bar_names_header_footer_only_files(gui, tmp_path):
    """#34 次项：顶部「无命中」旁写明其中几份只删页眉页脚。"""
    scan = _hf_zero_scan(tmp_path)
    _show_one(gui, [scan], True)
    copy = _visible_copy(gui.root)
    assert "无命中 1" in copy
    assert "其中 1 份只删页眉页脚已写出" in copy
    _show_one(gui, [scan], False)
    assert "其中 1 份只删页眉页脚将写出" in _visible_copy(gui.root)


def test_blocked_file_is_not_said_to_lose_header_footer(gui, tmp_path):
    """merge-r5b S3：被拦下的文件不写出，右侧汇总栏不能写「页眉页脚 N 页 · 将删除」。"""
    from doc_redactor.image_redactor import MSG_NO_OCR

    scan = _hf_zero_scan(tmp_path, "blocked-hf.pdf")
    scan.blocked_reason = MSG_NO_OCR.format(n=1)
    for written in (False, True):
        _show_one(gui, [scan], written)
        copy = _visible_copy(gui.root)
        assert "· 将删除" not in copy and "· 已删除" not in copy, copy
        assert "页眉页脚将删除" not in copy and "页眉页脚已删除" not in copy
        assert "未写出" in copy
        assert "只删页眉页脚" not in copy


def test_blocked_zero_hit_file_is_not_counted_as_hit(gui, tmp_path):
    """merge-r5 S2：被拦下且 0 命中的文件，顶栏不算进「有命中」。"""
    from doc_redactor.image_redactor import MSG_NO_OCR

    scan = _hf_zero_scan(tmp_path, "blocked-zero.pdf")
    scan.blocked_reason = MSG_NO_OCR.format(n=1)
    _show_one(gui, [scan], False)
    copy = _visible_copy(gui.root)
    assert "有命中 0" in copy and "无命中 0" in copy and "没法检查 1" in copy, copy


def test_inline_image_near_match_shows_approx_in_hit_table(gui, tmp_path, monkeypatch):
    """merge-r5 S1 / r5b S1-t：内联图 OCR 近似命中，界面命中表里标「近似」。编造数据。"""
    import doc_redactor.scan as scan_mod
    from doc_redactor.engine import WordlistRule, build_engine

    sys.path.insert(0, str(Path(__file__).parent))
    from test_image_metadata import _inline_pdf

    def fake_entities(image, engine):
        return [{"category": "药名", "text": "Fictlzumab", "boxes": [(10, 10, 100, 20)], "approx": True}]

    monkeypatch.setattr(scan_mod, "analyze_image_entities", fake_entities)
    monkeypatch.setattr(scan_mod, "require_tesseract", lambda *a, **k: None)
    scan = scan_mod.scan_path(_inline_pdf(tmp_path), build_engine([], [WordlistRule("药名", ["Fictizumab"], True)]))
    assert any(h.unit.startswith("pdfinline:") for h in scan.hits)
    _show_one(gui, [scan], False)
    tree = _hit_table(gui.root)
    row = tree.get_children()[0]
    assert str(tree.set(row, "text")).startswith("近似：读成")
    assert "近似：读成 Fictlzumab" in _hover_tip(gui, tree, row, "text")


def test_top_bar_unchecked_count_matches_blocked_rows(gui, tmp_path):
    """#51 M1：被拦下的文件不论有没有命中，都只算顶栏「没法检查」，数字等于列表里 ⚠ 没法检查 的份数。编造数据。"""
    import re as _re

    from doc_redactor.image_redactor import MSG_NO_OCR
    from doc_redactor.scan import KIND_PDF, KIND_WORD, FileScan

    reason = MSG_NO_OCR.format(n=1)
    kept = FileScan(path=Path("C.pdf"), kind=KIND_PDF, status="ok", hits=[_plain_hit(1, "Acme", "page:0")], pages=[["Acme"]])
    blocked_hit = FileScan(
        path=Path("A.docx"), kind=KIND_WORD, status="ok", hits=[_plain_hit(2, "星辰海", "para:0")],
        paragraphs=["星辰海"], blocked_reason=reason,
    )
    blocked_zero = _hf_zero_scan(tmp_path, "E.pdf")
    blocked_zero.blocked_reason = reason
    _show_one(gui, [kept, blocked_hit, blocked_zero], False)
    tree = _file_tree(gui.root)
    rows = [tree.item(i, "values")[0] for i in tree.get_children()]
    warned = sum(1 for r in rows if r.startswith("⚠") and "没法检查" in r)
    assert warned == 2, rows
    copy = _visible_copy(gui.root)
    m = _re.search(r"没法检查 (\d+)", copy)
    assert m and int(m.group(1)) == warned, copy
    assert "有命中 1" in copy and "无命中 0" in copy, copy


def test_top_bar_counts_each_file_once(gui, tmp_path):
    """#51：顶栏每份文件只算一类，各类之和等于总份数；写出失败单算，不再同时算「有命中」。编造数据。"""
    import re as _re

    from doc_redactor.image_redactor import MSG_NO_OCR
    from doc_redactor.scan import KIND_PDF, KIND_WORD, FileScan

    reason = MSG_NO_OCR.format(n=1)
    kept = FileScan(path=Path("C.pdf"), kind=KIND_PDF, status="ok", hits=[_plain_hit(1, "Acme", "page:0")], pages=[["Acme"]])
    wfail = FileScan(path=Path("W.pdf"), kind=KIND_PDF, status="ok", hits=[_plain_hit(2, "Acme", "page:0")], pages=[["Acme"]])
    nohit = FileScan(path=Path("D.pdf"), kind=KIND_PDF, status="ok", hits=[], pages=[["plain"]])
    blocked_hit = FileScan(
        path=Path("A.docx"), kind=KIND_WORD, status="ok", hits=[_plain_hit(3, "星辰海", "para:0")],
        paragraphs=["星辰海"], blocked_reason=reason,
    )
    scans = [kept, wfail, nohit, blocked_hit]

    def counts():
        # r2-3：只从顶栏取数，底栏的「失败 N 个」不算。
        copy = _visible_copy(_top_counts_bar(gui.root))
        got = {}
        for name in ("有命中", "无命中", "没法检查", "失败"):
            m = _re.search(name + r" (\d+)", copy)
            got[name] = int(m.group(1)) if m else 0
        return got

    _show_one(gui, scans, False)
    before = counts()
    assert sum(before.values()) == len(scans), before
    assert before == {"有命中": 2, "无命中": 1, "没法检查": 1, "失败": 0}, before

    gui.scans = scans
    gui.written = True
    gui.write_errors = {str(wfail.path): "磁盘已满", str(blocked_hit.path): reason}
    gui.write_notes = {}
    gui.show("result")
    _settle(gui)
    after = counts()
    assert sum(after.values()) == len(scans), after
    assert after["失败"] == 2 and after["有命中"] == 1, after


def _top_counts_bar(root):
    found = [node for node in _walk(root) if getattr(node, "top_counts", False)]
    assert len(found) == 1, "顶栏计数控件应恰好一个"
    return found[0]


def test_nohit_unwritten_row_does_not_share_blocked_icon(gui):
    """r2-1：「无命中，未写出」和「没法检查」不能用同一个符号。编造数据。"""
    from doc_redactor.gui.result_view import MSG_NO_HITS
    from doc_redactor.gui.theme import status_style
    from doc_redactor.image_redactor import MSG_NO_OCR
    from doc_redactor.scan import KIND_PDF, KIND_WORD, FileScan

    reason = MSG_NO_OCR.format(n=1)
    nohit = FileScan(path=Path("D.pdf"), kind=KIND_PDF, status="ok", hits=[], pages=[["plain"]])
    blocked = FileScan(
        path=Path("A.docx"), kind=KIND_WORD, status="ok", hits=[], paragraphs=["Acme Pharma"], blocked_reason=reason,
    )
    gui.scans = [nohit, blocked]
    gui.selected = 0
    gui.written = True
    gui.write_errors = {str(blocked.path): reason}
    gui.write_notes = {str(nohit.path): MSG_NO_HITS}
    gui.preview_mode = "orig"
    gui.hide_text = False
    gui.show("result")
    _settle(gui)
    assert status_style("nohit_unwritten")[0] != status_style("blocked")[0]
    line = next(r for r in _tree_values(gui.root).splitlines() if "无命中，未写出" in r)
    assert "⚠" not in line, line


def test_approx_hit_is_on_by_default_and_hidden_text_still_says_approx(gui):
    """方案 A：近似命中默认勾选；隐藏原文时只标「近似」，不露读到的字。编造数据。"""
    from doc_redactor.scan import KIND_PDF, FileScan, LocatedHit

    hit = LocatedHit(
        hit_id=1, category="DRUG", text="菲蔡珠单抗", location="第 1 页图 1 · 图内 第 1 行",
        unit="pdfimg:5", boxes=[(10, 10, 100, 20)], page=0, approx=True,
    )
    assert hit.enabled is True
    scan = FileScan(path=Path("scan.pdf"), kind=KIND_PDF, status="ok", hits=[hit], pages=[[""]])
    _show_one(gui, [scan], False)
    assert "☑" in _tree_values(gui.root)
    assert "近似：读成 菲蔡珠单抗" in _hover_tip(gui, _hit_table(gui.root), "1", "text")
    gui.hide_text = True
    gui.show("result")
    _settle(gui)
    hidden = _tree_values(gui.root)
    assert "近似" in hidden and "菲蔡珠单抗" not in hidden
    assert "菲蔡珠单抗" not in _hover_tip(gui, _hit_table(gui.root), "1", "text")


def _hit_table(root):
    from tkinter import ttk

    return next(n for n in _walk(root) if isinstance(n, ttk.Treeview) and "text" in tuple(n.cget("columns")))


def _hover_tip(gui, tree, iid, column):
    """Move the pointer onto a cell and return the tooltip text the user sees ("" if none)."""
    x, y, w, h = tree.bbox(iid, column)
    tree.event_generate("<Motion>", x=x + w // 2, y=y + h // 2)
    _settle(gui)
    tip = getattr(gui, "_loc_tip", None)
    return _tip_label_text(tip)


def _tip_label_text(tip) -> str:
    """s7：读提示框里 Label 实际显示的文字，不读内部属性。"""
    if tip is None or not tip.winfo_exists():
        return ""
    return "\n".join(str(w.cget("text")) for w in tip.winfo_children())


def _approx_scan(location, text):
    from doc_redactor.scan import KIND_PDF, FileScan, LocatedHit

    hit = LocatedHit(
        hit_id=1, category="DRUG", text=text, location=location,
        unit="pdfimg:5", boxes=[(10, 10, 100, 20)], page=500, approx=True,
    )
    return FileScan(path=Path("scan.pdf"), kind=KIND_PDF, status="ok", hits=[hit], pages=[[""]] * 501)


LONG_SPACED = "菲 蔡 珠 单 抗 注 射 液 冻 干 粉 针 剂"


@pytest.mark.parametrize("size", ["960x640", "1100x720"])
def test_long_spaced_approx_text_ellipsized_and_full_on_hover_and_select(gui, size):
    """#54 M1：一字一空格的长串放不下时末尾「…」，悬停、选中都能看到全文；全文不进命中清单。编造数据。"""
    import json

    from doc_redactor.backpath import hit_list_payload

    gui.root.geometry(f"{size}+0+0")
    scan = _approx_scan("第 1 页图 1 · 图内 第 1 行", LONG_SPACED)
    _show_one(gui, [scan], False)
    tree = _hit_table(gui.root)
    shown = str(tree.set("1", "text"))
    font = gui.fonts["small"]
    assert shown.startswith("近似：读成 菲") and shown.endswith("…"), shown
    assert font.measure(shown) <= tree.bbox("1", "text")[2], shown
    full = f"近似：读成 {LONG_SPACED}"
    assert full in _hover_tip(gui, tree, "1", "text")
    tree.event_generate("<Leave>")
    _settle(gui)
    gui.jump_to_hit(1)
    gui.show("result")
    _settle(gui)
    tip = getattr(gui, "_loc_tip", None)
    assert full in _tip_label_text(tip)
    payload = json.dumps(hit_list_payload([scan]), ensure_ascii=False)
    assert "蔡" not in payload and "近似" not in payload


def test_three_digit_page_location_visible_with_approx_row(gui):
    """#54 M1：有近似行时位置列不再被压窄；三位数页码的页号、图号看得到，全文在悬停提示里。编造数据。"""
    gui.root.geometry("960x640+0+0")
    loc_text = "第 501 页图 12 · 图内 第 3 行"
    scan = _approx_scan(loc_text, "菲蔡珠单抗")
    _show_one(gui, [scan], False)
    tree = _hit_table(gui.root)
    shown = str(tree.set("1", "loc"))
    assert shown.startswith("第 501 页图 12"), shown
    assert gui.fonts["small"].measure(shown) <= tree.bbox("1", "loc")[2], shown
    assert loc_text in _hover_tip(gui, tree, "1", "loc")


def test_three_digit_page_text_location_shown_in_full(gui):
    """#54 M1：正文命中「第 501 页 · 第 1 行」在位置列里完整显示，不加「…」。编造数据。"""
    from doc_redactor.scan import KIND_PDF, FileScan, LocatedHit

    gui.root.geometry("960x640+0+0")
    hit = LocatedHit(
        hit_id=1, category="申办方", text="Acme Pharma", location="第 501 页 · 第 1 行",
        unit="page:500", boxes=[(10, 10, 100, 20)], page=500,
    )
    approx = LocatedHit(
        hit_id=2, category="DRUG", text="菲蔡珠单抗", location="第 501 页图 1 · 图内 第 1 行",
        unit="pdfimg:5", boxes=[(10, 40, 100, 20)], page=500, approx=True,
    )
    scan = FileScan(path=Path("scan.pdf"), kind=KIND_PDF, status="ok", hits=[hit, approx], pages=[["Acme Pharma"]] * 501)
    _show_one(gui, [scan], False)
    assert str(_hit_table(gui.root).set("1", "loc")) == "第 501 页 · 第 1 行"


def test_selected_hit_tip_closes_when_leaving_result_page(gui):
    """#54 自审必改：选中后弹出的全文提示，离开结果页就关掉，不浮在别的界面上。编造数据。"""
    gui.root.geometry("960x640+0+0")
    scan = _approx_scan("第 1 页图 1 · 图内 第 1 行", LONG_SPACED)
    _show_one(gui, [scan], False)
    gui.jump_to_hit(1)
    gui.show("result")
    _settle(gui)
    tip = getattr(gui, "_loc_tip", None)
    assert tip is not None and tip.winfo_exists()
    gui.show("main")
    _settle(gui)
    assert not tip.winfo_exists()



def _many_approx(gui, n=12, location="第 1 页图 1 · 图内 第 1 行", text=LONG_SPACED):
    from doc_redactor.scan import KIND_PDF, FileScan, LocatedHit

    hits = [
        LocatedHit(hit_id=i, category="DRUG", text=text, location=location,
                   unit="pdfimg:5", boxes=[(10, 10 * i, 100, 20)], page=0, approx=True)
        for i in range(1, n + 1)
    ]
    scan = FileScan(path=Path("scan.pdf"), kind=KIND_PDF, status="ok", hits=hits, pages=[[""]])
    gui.root.geometry("960x640+0+0")
    _show_one(gui, [scan], False)
    return _hit_table(gui.root)


def test_tip_closes_when_window_loses_focus(gui):
    """s1：主窗口失去焦点时提示关掉，不盖在别的程序上。编造数据。"""
    tree = _many_approx(gui, n=1)
    assert _hover_tip(gui, tree, "1", "text")
    gui.root.event_generate("<FocusOut>")
    _settle(gui)
    assert _tip_label_text(getattr(gui, "_loc_tip", None)) == ""


def test_tip_closes_when_table_scrolls_without_wheel(gui):
    """s2：拖滚动条（yview_moveto / yview_scroll）后提示关掉。编造数据。"""
    tree = _many_approx(gui)
    assert _hover_tip(gui, tree, "1", "text")
    tree.yview_moveto(1.0)
    _settle(gui)
    assert _tip_label_text(getattr(gui, "_loc_tip", None)) == ""


def test_tip_stays_inside_window_right_edge(gui):
    """s5：长位置加长串时提示框不超出窗口右边。编造数据。"""
    long_loc = "第 501 页图 12 · 图内 第 3 行 · 附件二 扫描件 Acme Pharma 合同附页"
    tree = _many_approx(gui, n=1, location=long_loc, text=LONG_SPACED + " " + LONG_SPACED)
    assert _hover_tip(gui, tree, "1", "text")
    tip = gui._loc_tip
    tip.update_idletasks()
    right = gui.root.winfo_rootx() + gui.root.winfo_width()
    assert tip.winfo_rootx() + tip.winfo_reqwidth() <= right


def test_tip_stays_when_pointer_moves_onto_it(gui):
    """s3：鼠标从这一行移到提示上，提示不关。编造数据。"""
    tree = _many_approx(gui, n=1)
    assert _hover_tip(gui, tree, "1", "text")
    tip = gui._loc_tip
    tip.update_idletasks()
    gui.root.event_generate(
        "<Motion>", warp=True,
        x=tip.winfo_rootx() - gui.root.winfo_rootx() + 10, y=tip.winfo_rooty() - gui.root.winfo_rooty() + 5,
    )
    tree.event_generate("<Leave>")
    _settle(gui)
    assert _tip_label_text(getattr(gui, "_loc_tip", None))


def test_tip_closes_when_app_loses_focus_from_child_widget(gui, monkeypatch):
    """自审 r4 必改 1：真实失焦时 FocusOut 落在子控件（位置表）上，提示也要关。编造数据。"""
    tree = _many_approx(gui, n=1)
    assert _hover_tip(gui, tree, "1", "text")
    monkeypatch.setattr(gui.root, "focus_get", lambda: None)
    tree.event_generate("<FocusOut>")
    _settle(gui)
    assert _tip_label_text(getattr(gui, "_loc_tip", None)) == ""


def test_tip_touches_its_row_so_pointer_can_reach_it(gui):
    """自审 r4 必改 2：提示和这一行之间没有空隙，鼠标往下移不会先落到下一行。编造数据。"""
    tree = _many_approx(gui, n=3)
    assert _hover_tip(gui, tree, "1", "text")
    tip = gui._loc_tip
    tip.update_idletasks()
    x, y, w, h = tree.bbox("1")
    row_bottom = tree.winfo_rooty() + y + h
    assert tip.winfo_rooty() <= row_bottom
