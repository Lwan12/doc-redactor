from pathlib import Path

import pymupdf
from docx import Document

from doc_redactor.config import load_rules
from doc_redactor.docx_redactor import redact_docx
from doc_redactor.engine import (
    DEFAULT_REPLACEMENTS,
    RegexRule,
    build_engine,
    redaction_label,
)
from doc_redactor.pdf_redactor import fill_placeholder, redact_pdf
from doc_redactor.ruleset import (
    LETTER_DIGIT_NAME,
    PROTOCOL_LIST_NAME,
    common_library,
    library_item_to_rule,
    load_ruleset,
)
from doc_redactor.scan import scan_path, write_report, write_scan
from helpers import RULES, assert_report_has_no_secrets


def test_default_table_covers_library_and_yaml_names():
    assert redaction_label("DRUG") == "[STUDY DRUG]"
    assert redaction_label("药名") == "[STUDY DRUG]"
    assert redaction_label("SPONSOR") == "[SPONSOR]"
    assert redaction_label("申办方") == "[SPONSOR]"
    assert redaction_label("PROJECT_ID") == "[PROTOCOL]"
    assert redaction_label("项目号") == "[PROTOCOL]"
    assert redaction_label("NCT") == "[NCT ID]"
    assert redaction_label("NCT 试验登记号") == "[NCT ID]"
    assert redaction_label(PROTOCOL_LIST_NAME) == "[PROTOCOL NO]"
    assert redaction_label(LETTER_DIGIT_NAME) == "[PROTOCOL NO]"
    assert redaction_label("邮箱") == "[EMAIL]"
    assert redaction_label("EMAIL") == "[EMAIL]"
    assert redaction_label("手机号") == "[PHONE]"
    assert redaction_label("手机号（中国大陆）") == "[PHONE]"
    assert redaction_label("PHONE") == "[PHONE]"
    assert redaction_label("批次号") == "[REDACTED:批次号]"
    assert DEFAULT_REPLACEMENTS["DRUG"] == "[STUDY DRUG]"


def test_custom_replacement_overrides_default(tmp_path: Path):
    rules = tmp_path / "rules.yaml"
    rules.write_text(
        "regex:\n"
        "  - name: PROJECT_ID\n"
        "    pattern: 'ABC-\\d{4}-\\d{3}'\n"
        "    replacement: '[PROJ]'\n"
        "wordlists:\n"
        "  - category: SPONSOR\n"
        "    case_sensitive: false\n"
        "    terms: [Acme Pharma]\n"
        "    replacement: '[CO]'\n",
        encoding="utf-8",
    )
    engine = load_rules(rules)
    assert engine.label("PROJECT_ID") == "[PROJ]"
    assert engine.label("SPONSOR") == "[CO]"
    source = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("Project ABC-2026-001 by Acme Pharma")
    document.save(source)
    dest = tmp_path / "notes.redacted.docx"
    redact_docx(source, dest, engine)
    text = "\n".join(p.text for p in Document(str(dest)).paragraphs)
    assert text == "Project [PROJ] by [CO]"
    assert "ABC-2026-001" not in text
    assert "Acme Pharma" not in text


def test_old_yaml_without_replacement_uses_defaults(tmp_path: Path):
    rules = tmp_path / "old.yaml"
    rules.write_text(
        "regex:\n"
        "  - name: PROJECT_ID\n"
        "    pattern: '\\bABC-\\d{4}-\\d{3}\\b'\n"
        "wordlists:\n"
        "  - category: SPONSOR\n"
        "    case_sensitive: false\n"
        "    terms: [Acme Pharma]\n"
        "  - category: DRUG\n"
        "    case_sensitive: true\n"
        "    terms: [Fictizumab]\n",
        encoding="utf-8",
    )
    engine = load_rules(rules)
    assert engine.label("PROJECT_ID") == "[PROTOCOL]"
    assert engine.label("SPONSOR") == "[SPONSOR]"
    assert engine.label("DRUG") == "[STUDY DRUG]"
    loaded = load_ruleset(rules)
    assert loaded.rules[0].replacement == ""
    assert loaded.replacements()["PROJECT_ID"] == "[PROTOCOL]"


def test_same_category_reuses_placeholder_without_numbering(tmp_path: Path):
    engine = build_engine(
        [RegexRule(name="PROJECT_ID", pattern=r"ABC-\d{4}-\d{3}", flags=0)],
        [],
    )
    source = tmp_path / "two.docx"
    document = Document()
    document.add_paragraph("See ABC-2026-001 and ABC-2026-002.")
    document.save(source)
    dest = tmp_path / "two.redacted.docx"
    redact_docx(source, dest, engine)
    text = Document(str(dest)).paragraphs[0].text
    assert text == "See [PROTOCOL] and [PROTOCOL]."
    assert "[PROTOCOL-1]" not in text
    assert "[PROTOCOL-2]" not in text
    assert "ABC-2026-001" not in text
    assert "ABC-2026-002" not in text


def test_docx_cross_run_uses_category_placeholder(tmp_path: Path):
    engine = build_engine(
        [RegexRule(name="PROJECT_ID", pattern=r"ABC-\d{4}-\d{3}", flags=0)],
        [],
    )
    source = tmp_path / "split.docx"
    document = Document()
    para = document.add_paragraph()
    para.add_run("ID ")
    bold = para.add_run("ABC-20")
    bold.bold = True
    para.add_run("26")
    para.add_run("-001")
    para.add_run(" end")
    document.save(source)
    dest = tmp_path / "split.redacted.docx"
    redact_docx(source, dest, engine)
    written = Document(str(dest)).paragraphs[0]
    assert written.text == "ID [PROTOCOL] end"
    label = next(run for run in written.runs if run.text == "[PROTOCOL]")
    assert label.bold
    assert "ABC-20" not in written.text


def test_pdf_writes_placeholder_and_drops_source(tmp_path: Path):
    engine = load_rules(RULES)
    source = tmp_path / "brief.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 720), "Project ABC-2026-001")
    page.insert_text((72, 700), "Sponsor Acme Pharma")
    page.insert_text((72, 680), "Drug Fictizumab")
    page.insert_text((72, 660), "Again ABC-2026-001")
    document.save(source)
    document.close()
    dest = tmp_path / "brief.redacted.pdf"
    counts = redact_pdf(source, dest, engine)
    opened = pymupdf.open(dest)
    text = "\n".join(page.get_text() for page in opened)
    opened.close()
    assert "ABC-2026-001" not in text
    assert "Acme Pharma" not in text
    assert "Fictizumab" not in text
    assert "[PROTOCOL]" in text
    assert "[SPONSOR]" in text
    assert "[STUDY DRUG]" in text
    assert text.count("[PROTOCOL]") == 2
    assert counts["PROJECT_ID"] == 2


def test_pdf_placeholder_shrinks_or_stays_blank():
    document = pymupdf.open()
    page = document.new_page()
    wide = pymupdf.Rect(72, 700, 72 + 73, 700 + 15)
    assert fill_placeholder(page, wide, "[PROTOCOL]") is True
    assert "[PROTOCOL]" in page.get_text()

    page2 = document.new_page()
    shrink = pymupdf.Rect(72, 200, 72 + 50, 216)
    assert fill_placeholder(page2, shrink, "[STUDY DRUG]") is True
    compact = "".join(page2.get_text().split())
    assert "[STUDYDRUG]" in compact
    assert "[STU" not in compact.replace("[STUDYDRUG]", "")

    page3 = document.new_page()
    tight = pymupdf.Rect(72, 200, 96, 208)
    assert fill_placeholder(page3, tight, "[STUDY DRUG]") is False
    assert page3.get_text().strip() == ""
    assert "[STU" not in page3.get_text()
    document.close()


def test_report_has_no_source_and_no_placeholders(tmp_path: Path):
    source = tmp_path / "one.docx"
    document = Document()
    document.add_paragraph("Acme Pharma and ABC-2026-001 and Fictizumab")
    document.save(source)
    engine = load_rules(RULES)
    scan = scan_path(source, engine)
    dest = tmp_path / "one.redacted.docx"
    counts = write_scan(scan, dest)
    report = tmp_path / "one.report.json"
    write_report(report, counts)
    raw = report.read_text(encoding="utf-8")
    payload = assert_report_has_no_secrets(report)
    assert payload == counts
    for marker in ("[PROTOCOL]", "[SPONSOR]", "[STUDY DRUG]", "[REDACTED:", "ABC-2026-001"):
        assert marker not in raw
    assert set(payload) <= {"DRUG", "PROJECT_ID", "SPONSOR"}
    assert all(isinstance(value, int) and value > 0 for value in payload.values())


def test_library_items_carry_default_placeholders():
    items = {item.name: item for item in common_library()}
    assert items["项目号"].replacement == "[PROTOCOL]"
    assert items["申办方"].replacement == "[SPONSOR]"
    assert items["药名"].replacement == "[STUDY DRUG]"
    assert items["NCT 试验登记号"].replacement == "[NCT ID]"
    assert items[PROTOCOL_LIST_NAME].replacement == "[PROTOCOL NO]"
    assert items[LETTER_DIGIT_NAME].replacement == "[PROTOCOL NO]"
    assert items["邮箱"].replacement == "[EMAIL]"
    assert items["手机号（中国大陆）"].replacement == "[PHONE]"
    rule = library_item_to_rule(items["药名"])
    assert rule.replacement_text() == "[STUDY DRUG]"
