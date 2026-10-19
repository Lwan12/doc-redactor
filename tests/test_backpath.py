"""Back path: signature, leave/close, batch folders, retry, hit list."""

from __future__ import annotations

import errno
import json
from datetime import datetime
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from doc_redactor.backpath import (
    MSG_BACK_BODY,
    MSG_BACK_KEEP,
    MSG_BACK_TITLE,
    MSG_CLEARED_TOP,
    MSG_KEPT,
    MSG_STALE,
    MSG_STALE_DROP,
    ScanTicket,
    UncheckMemory,
    allocate_batch_dir,
    apply_unchecks,
    assert_not_source,
    back_prompt_body,
    bump_generation,
    capture_unchecks,
    clear_file_list,
    decide_close,
    decide_leave,
    decide_written,
    discard_scan_results,
    files_signature,
    hit_list_payload,
    is_current,
    is_write_failure,
    merge_kept,
    protect_written,
    rescan_hint,
    retain_editor_draft,
    unwritten_failure_banner,
    rules_signature,
    session_signature,
    step_clickable,
    without_paths,
    write_allowed,
)
from doc_redactor.engine import RegexRule, build_engine
from doc_redactor.scan import (
    MSG_ALREADY_WRITTEN,
    FileScan,
    KIND_PDF,
    KIND_WORD,
    LocatedHit,
    WriteResult,
    scan_path,
    write_scan,
    write_scans,
)


def _rule(name="项目号", *, enabled=True, pattern=r"ABC-\d{4}-\d{3}", replacement="", terms=None):
    class Rule:
        pass

    rule = Rule()
    rule.name = name
    rule.kind = "wordlist" if terms is not None else "pattern"
    rule.enabled = enabled
    rule.pattern = "" if terms is not None else pattern
    rule.placeholder = ""
    rule.replacement = replacement
    rule.case_sensitive = False
    rule.terms = list(terms or [])
    rule.examples = [] if terms is not None else ["ABC-2026-001"]
    return rule


def _hit(enabled=True, category="项目号", start=0, end=12, unit="p:0", location="第 1 段"):
    return LocatedHit(
        hit_id=1,
        category=category,
        text="ABC-2026-001",
        location=location,
        enabled=enabled,
        unit=unit,
        start=start,
        end=end,
    )


def _scan(path, hits, kind=KIND_WORD):
    return FileScan(path=Path(path), kind=kind, status="ok", hits=list(hits))


def _engine():
    return build_engine([RegexRule(name="PROJECT_ID", pattern=r"ABC-\d{4}-\d{3}", flags=0)], [])


def _docx(path: Path, line: str) -> None:
    document = Document()
    document.add_paragraph(line)
    document.save(path)


def test_signature_tracks_files_and_checked_rule_content():
    files = [Path("a.docx")]
    rules = [_rule()]
    base = session_signature(files, rules)
    assert session_signature(files, rules) == base
    assert session_signature([Path("a.docx"), Path("b.docx")], rules) != base
    assert session_signature([], rules) != base

    changed = [_rule(replacement="[PROTOCOL]")]
    assert rules_signature(changed) == rules_signature(rules)
    changed[0].replacement = ""
    changed[0].pattern = r"ZZZ-\d+"
    assert rules_signature(changed) != rules_signature(rules)

    wordlist = [_rule(terms=["华曜医药", "北辰制药"])]
    other = [_rule(terms=["华曜医药"])]
    assert rules_signature(wordlist) != rules_signature(other)

    disabled = [_rule(enabled=False), _rule(name="药名", pattern=r"drug")]
    only_drug = [_rule(name="药名", pattern=r"drug")]
    assert rules_signature(disabled) == rules_signature(only_drug)
    assert files_signature([Path("a.docx")]) == files_signature([Path("a.docx")])


def test_unchecks_kept_only_when_signature_matches():
    path = Path("notes.docx")
    hits = [_hit(True, start=0), _hit(False, start=20, end=32, location="第 2 段")]
    scan = _scan(path, hits)
    rules = [_rule()]
    signature = session_signature([path], rules)
    memory = capture_unchecks([scan], signature)
    assert len(memory.keys) == 1

    fresh = _scan(path, [_hit(True, start=0), _hit(True, start=20, end=32, location="第 2 段")])
    assert apply_unchecks([fresh], memory, signature) == "kept"
    assert fresh.hits[0].enabled is True
    assert fresh.hits[1].enabled is False

    moved = session_signature([path, Path("extra.docx")], rules)
    again = _scan(path, [_hit(True, start=0), _hit(True, start=20, end=32, location="第 2 段")])
    assert apply_unchecks([again], memory, moved) == "cleared"
    assert all(hit.enabled for hit in again.hits)

    relabelled = session_signature([path], [_rule(replacement="[其他]")])
    assert relabelled == signature
    edited = session_signature([path], [_rule(pattern=r"ZZZ-\d+")])
    third = _scan(path, [_hit(True), _hit(True, start=20, end=32, location="第 2 段")])
    assert apply_unchecks([third], memory, edited) == "cleared"
    assert all(hit.enabled for hit in third.hits)
    assert MSG_CLEARED_TOP.startswith("文件或规则改过了")


def test_leave_and_step_and_close_share_the_result_rule():
    unwritten = decide_leave(on_result=True, written=False, has_preview=True, unchecked=1)
    assert unwritten.prompt is True
    assert unwritten.clear_preview is True
    assert unwritten.title == MSG_BACK_TITLE
    assert unwritten.body == MSG_BACK_BODY + MSG_BACK_KEEP.format(n=1)
    assert unwritten.confirm == "返回修改"
    assert unwritten.stay == "留在这里"

    zero = decide_leave(on_result=True, written=False, has_preview=True, unchecked=0)
    assert zero.body == MSG_BACK_BODY
    assert MSG_BACK_KEEP.format(n=0) not in zero.body

    written = decide_leave(on_result=True, written=True, has_preview=True, unchecked=2)
    assert written.prompt is False
    assert written.clear_preview is True

    rules_page = decide_leave(on_result=False, written=False, has_preview=False, unchecked=0)
    assert rules_page.prompt is False
    assert rules_page.clear_preview is False

    # Stay is the caller ignoring the decision. The decision itself does not clear anything.
    stayed = decide_leave(on_result=True, written=False, has_preview=True, unchecked=1)
    assert stayed is not None and stayed.prompt is True

    for target in (0, 1, 2):
        assert step_clickable(target=target, current=3, reached=3, running=False)
        decision = decide_leave(on_result=True, written=False, has_preview=True, unchecked=1)
        assert decision.prompt is True
    assert step_clickable(target=3, current=3, reached=3, running=False) is False
    assert step_clickable(target=1, current=1, reached=3, running=False) is False
    assert step_clickable(target=0, current=1, reached=1, running=False) is True
    assert step_clickable(target=0, current=2, reached=2, running=True) is False

    assert decide_close(on_result=True, written=False, has_preview=True) is True
    assert decide_close(on_result=True, written=True, has_preview=True) is False
    assert decide_close(on_result=False, written=False, has_preview=True) is False
    assert decide_close(on_result=True, written=False, has_preview=False) is False


def test_prompt_copy_and_rescan_hint():
    assert back_prompt_body(0) == "返回后，这次的扫描结果会清除，要重新扫描。"
    assert "你取消勾选的 2 处" in back_prompt_body(2)
    assert rescan_hint(from_result=True, cancelled=False, signature_matches=True, unchecked=2) == MSG_KEPT.format(n=2)
    assert rescan_hint(from_result=True, cancelled=False, signature_matches=False, unchecked=2) == MSG_STALE_DROP
    assert rescan_hint(from_result=True, cancelled=False, signature_matches=False, unchecked=0) == MSG_STALE
    assert rescan_hint(from_result=True, cancelled=True, signature_matches=False, unchecked=2) is None
    assert rescan_hint(from_result=False, cancelled=False, signature_matches=True, unchecked=2) is None


def test_cancel_discards_results_and_write_needs_a_fresh_scan():
    scan = _scan("a.docx", [_hit()])
    assert discard_scan_results([scan]) == []
    assert write_allowed(has_fresh_scans=False) is False
    assert write_allowed(has_fresh_scans=True) is True


def test_remove_and_clear_keep_rules_and_output():
    files = [Path("a.docx"), Path("b.docx"), Path("c.docx")]
    sizes = ["1 KB", "2 KB", "3 KB"]
    status = ["pending", "pending", "pending"]
    kept_f, kept_s, kept_st = without_paths(files, sizes, status, {"b.docx"})
    assert [path.name for path in kept_f] == ["a.docx", "c.docx"]
    assert kept_s == ["1 KB", "3 KB"]
    assert kept_st == ["pending", "pending"]

    rules = [_rule()]
    output = Path("脱敏输出")
    cleared, same_rules, same_output = clear_file_list(rules, output)
    assert cleared == []
    assert same_rules is rules
    assert same_output is output
    assert session_signature(files, rules) != session_signature(cleared, rules)


def test_editor_draft_is_not_cleared():
    draft = {"examples": "AAA-0000-000", "test": "方案 AAA-0000-000", "name": "访视项目号"}
    kept = retain_editor_draft(draft)
    assert kept["examples"] == "AAA-0000-000"
    assert kept["test"] == "方案 AAA-0000-000"
    assert kept is draft


def test_batch_dir_sequence_does_not_reuse(tmp_path: Path):
    when = datetime(2026, 10, 17, 21, 40)
    first = allocate_batch_dir(tmp_path, when)
    assert first.name == "20261017-2140"
    first.mkdir()
    second = allocate_batch_dir(tmp_path, when)
    assert second.name == "20261017-2140-2"
    second.mkdir()
    third = allocate_batch_dir(tmp_path, when)
    assert third.name == "20261017-2140-3"
    assert first != second != third


def test_two_writes_do_not_overwrite_previous_or_source(tmp_path: Path):
    source = tmp_path / "notes.docx"
    _docx(source, "Project ABC-2026-001 stays in the source")
    original = source.read_bytes()
    engine = _engine()
    scan = scan_path(source, engine)
    assert scan.status != "fail"
    when = datetime(2026, 10, 17, 21, 40)
    root = tmp_path / "out"
    root.mkdir()
    first_dir = allocate_batch_dir(root, when)
    first = write_scans([scan], first_dir, overwrite=False, sequence=False)
    assert first[0].dest is not None
    assert first[0].dest.parent == first_dir
    assert first[0].dest.resolve() != source.resolve()
    first_bytes = first[0].dest.read_bytes()
    assert b"ABC-2026-001" not in first_bytes
    report = first[0].report.read_text(encoding="utf-8")
    assert "ABC-2026-001" not in report

    second_dir = allocate_batch_dir(root, when)
    assert second_dir != first_dir
    second = write_scans([scan], second_dir, overwrite=False, sequence=False)
    assert second[0].dest is not None
    assert second[0].dest.parent == second_dir
    assert first[0].dest.read_bytes() == first_bytes
    assert source.read_bytes() == original

    # Same folder, same name: sequence keeps the first file.
    again = write_scans([scan], first_dir, overwrite=False, sequence=True)
    assert again[0].dest is not None
    assert again[0].dest != first[0].dest
    assert again[0].dest.exists()
    assert first[0].dest.read_bytes() == first_bytes
    protected = protect_written([first[0].dest, second[0].dest])
    assert first[0].dest in protected and first[0].dest.exists()


def test_refuses_to_overwrite_the_source_path(tmp_path: Path):
    source = tmp_path / "notes.docx"
    _docx(source, "ABC-2026-001")
    try:
        assert_not_source(source, source)
    except ValueError as exc:
        assert "原文件" in str(exc)
    else:
        raise AssertionError("source path was accepted")
    scan = scan_path(source, _engine())
    try:
        write_scans_dest = source
        from doc_redactor.scan import write_scan

        write_scan(scan, write_scans_dest)
    except ValueError as exc:
        assert "原文件" in str(exc)
    else:
        raise AssertionError("write_scan overwrote the source")
    text = "\n".join(paragraph.text for paragraph in Document(str(source)).paragraphs)
    assert "ABC-2026-001" in text


def test_existing_dest_on_retry_is_skipped_not_an_error(tmp_path: Path):
    source = tmp_path / "notes.docx"
    _docx(source, "ABC-2026-001")
    scan = scan_path(source, _engine())
    batch = allocate_batch_dir(tmp_path / "out", datetime(2026, 10, 17, 21, 43))
    first = write_scans([scan], batch, overwrite=False, sequence=False)
    dest = first[0].dest
    assert dest is not None and dest.is_file()
    kept = dest.read_bytes()
    report = first[0].report.read_bytes() if first[0].report else None

    again = write_scans([scan], batch, overwrite=False, sequence=False)
    assert len(again) == 1
    item = again[0]
    assert item.skipped is True
    assert item.error is None
    assert item.note == MSG_ALREADY_WRITTEN
    assert item.dest == dest
    assert dest.read_bytes() == kept
    if report is not None:
        assert first[0].report.read_bytes() == report
    assert is_write_failure(scan, item) is False
    assert decide_written(False, again) is True
    assert decide_written(False, again * 2) is True


def test_retry_does_not_delete_successful_files(tmp_path: Path):
    good = tmp_path / "good.docx"
    bad_path = tmp_path / "missing.docx"
    _docx(good, "ABC-2026-001")
    engine = _engine()
    good_scan = scan_path(good, engine)
    bad_scan = FileScan(
        path=bad_path,
        kind=KIND_WORD,
        status="ok",
        hits=[_hit()],
        paragraphs=["ABC-2026-001"],
    )
    batch = allocate_batch_dir(tmp_path / "out", datetime(2026, 10, 17, 21, 41))
    results = write_scans([good_scan, bad_scan], batch, overwrite=False, sequence=False)
    by_path = {item.path: item for item in results}
    assert is_write_failure(good_scan, by_path[good]) is False
    assert by_path[good].dest is not None
    assert is_write_failure(bad_scan, by_path[bad_path]) is True
    kept_bytes = by_path[good].dest.read_bytes()

    retry_only = [bad_scan]
    retried = write_scans(retry_only, batch, overwrite=False, sequence=True)
    assert retried[0].path == bad_path
    assert by_path[good].dest.read_bytes() == kept_bytes
    merged = merge_kept([by_path[good].dest], [item.dest for item in retried if item.dest])
    assert by_path[good].dest in merged
    assert by_path[good].dest.exists()


def test_hit_list_has_no_source_text(tmp_path: Path):
    source = tmp_path / "one.docx"
    _docx(source, "Acme Pharma ABC-2026-001")
    scan = scan_path(source, _engine())
    scan.hits.append(
        LocatedHit(
            hit_id=99,
            category="SPONSOR",
            text="Acme Pharma",
            location="第 1 段",
            enabled=False,
            unit="p:0",
            start=0,
            end=11,
        )
    )
    payload = hit_list_payload([scan])
    raw = json.dumps(payload, ensure_ascii=False)
    assert "ABC-2026-001" not in raw
    assert "Acme Pharma" not in raw
    assert "第 1 段" not in raw
    assert payload[source.name]["PROJECT_ID"] == 1
    assert "SPONSOR" not in payload[source.name]


def test_written_pdf_is_not_restored_by_going_back(tmp_path: Path):
    source = tmp_path / "page.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "ABC-2026-001")
    document.save(source)
    document.close()
    source_bytes = source.read_bytes()
    scan = scan_path(source, _engine())
    assert scan.kind == KIND_PDF
    assert scan.hits
    batch = allocate_batch_dir(tmp_path / "pdf-out", datetime(2026, 10, 17, 21, 42))
    written = write_scans([scan], batch)
    dest = written[0].dest
    assert dest is not None
    frozen = dest.read_bytes()
    opened = pymupdf.open(dest)
    text = "\n".join(item.get_text() for item in opened)
    opened.close()
    assert "ABC-2026-001" not in text
    protected = protect_written([dest])
    assert discard_scan_results([scan]) == []
    assert protected[0].read_bytes() == frozen
    assert source.read_bytes() == source_bytes


def test_empty_uncheck_memory_is_none_outcome():
    scan = _scan("a.docx", [_hit(True)])
    memory = UncheckMemory(signature="x", keys=())
    assert apply_unchecks([scan], memory, "x") == "none"
    assert apply_unchecks([scan], None, "x") == "none"


def _write_result(name: str, *, dest: Path | None = None, error: str | None = None) -> WriteResult:
    return WriteResult(path=Path(name), dest=dest, report=None, counts={}, error=error)


def test_all_failed_writes_stay_unwritten_until_one_succeeds():
    failed = (
        _write_result("a.docx", error="输出文件夹里这个文件正被别的程序占用"),
        _write_result("b.docx", error="没写出"),
    )
    partial = (
        _write_result("a.docx", dest=Path("batch/a.redacted.docx")),
        _write_result("b.docx", error="输出文件夹里这个文件正被别的程序占用"),
    )
    assert decide_written(False, failed) is False
    assert decide_close(on_result=True, written=decide_written(False, failed), has_preview=True) is True
    assert decide_written(False, partial) is True
    assert decide_close(on_result=True, written=decide_written(False, partial), has_preview=True) is False
    # A retry that fails again must not clear a file that already landed.
    assert decide_written(True, failed) is True


@pytest.mark.parametrize(
    ("error", "expected"),
    [
        (PermissionError(13, "Permission denied", "SECRET-PATH"), "输出文件夹不能写入：没有权限"),
        (OSError(errno.ENOSPC, "No space left on device", "SECRET-PATH"), "输出文件夹不能写入：磁盘已满"),
        (OSError(errno.EIO, "Input/output error", "SECRET-PATH"), "输出文件夹不能写入：无法创建"),
    ],
)
def test_write_scans_mkdir_oserror_stays_unwritten(tmp_path: Path, monkeypatch, error, expected):
    source = tmp_path / "notes.docx"
    _docx(source, "ABC-2026-001")
    scan = scan_path(source, _engine())
    out = tmp_path / "blocked" / "batch"

    def boom(self, mode=0o777, parents=False, exist_ok=False):
        raise error

    monkeypatch.setattr(Path, "mkdir", boom)
    results = write_scans([scan], out)
    assert len(results) == 1
    assert results[0].dest is None
    assert results[0].error == expected
    assert "ABC-2026-001" not in results[0].error
    assert "SECRET-PATH" not in results[0].error
    assert "Permission denied" not in (results[0].error or "")
    assert "No space" not in (results[0].error or "")
    assert decide_written(False, results) is False
    assert unwritten_failure_banner([(source.name, results[0].error or "")]) == f"✕ {expected}"
    kept = "\n".join(paragraph.text for paragraph in Document(str(source)).paragraphs)
    assert "ABC-2026-001" in kept
    assert not out.exists()


def test_write_scan_mkdir_permission_error(tmp_path: Path, monkeypatch):
    source = tmp_path / "notes.docx"
    _docx(source, "ABC-2026-001")
    scan = scan_path(source, _engine())

    def boom(self, mode=0o777, parents=False, exist_ok=False):
        raise PermissionError(13, "Permission denied", "SECRET-PATH")

    monkeypatch.setattr(Path, "mkdir", boom)
    try:
        write_scan(scan, tmp_path / "out" / "notes.redacted.docx")
    except ValueError as exc:
        message = str(exc)
    else:
        raise AssertionError("mkdir PermissionError was not reported")
    assert message == "输出文件夹不能写入：没有权限"
    assert "SECRET-PATH" not in message
    assert "ABC-2026-001" not in message


def test_cancel_then_rescan_drops_the_previous_generation():
    generation = 0
    generation = bump_generation(generation)
    first = ScanTicket(generation)
    generation = bump_generation(generation)  # cancel retires the in-flight scan
    generation = bump_generation(generation)
    second = ScanTicket(generation)

    progress = []
    adopted = None

    def on_progress(ticket: ScanTicket, payload):
        if not is_current(ticket, generation):
            return
        progress.append(payload)

    def on_done(ticket: ScanTicket, payload):
        nonlocal adopted
        if not is_current(ticket, generation):
            return
        adopted = payload

    on_progress(first, (0, "scanned"))
    on_done(first, ["old-scan"])
    assert progress == []
    assert adopted is None
    assert is_current(first, generation) is False

    on_progress(second, (0, "scanned"))
    on_done(second, ["new-scan"])
    assert progress == [(0, "scanned")]
    assert adopted == ["new-scan"]
    assert is_current(second, generation) is True


def test_failed_write_leaves_no_half_file_for_retry_to_skip(tmp_path: Path, monkeypatch):
    import doc_redactor.scan as scan_mod

    source = tmp_path / "notes.docx"
    _docx(source, "ABC-2026-001")
    scan = scan_path(source, _engine())
    batch = allocate_batch_dir(tmp_path / "out", datetime(2026, 10, 17, 21, 44))

    def boom(scan_obj, dest, **kwargs):
        dest.write_bytes(b"half")
        raise OSError(errno.ENOSPC, "disk full")

    monkeypatch.setattr(scan_mod, "write_scan", boom)
    failed = write_scans([scan], batch, overwrite=False, sequence=False)
    assert failed[0].error
    assert not any(batch.glob("*.redacted.*"))
    monkeypatch.undo()
    retry = write_scans([scan], batch, overwrite=False, sequence=False)
    assert retry[0].note is None and retry[0].dest is not None and retry[0].dest.stat().st_size > 4
