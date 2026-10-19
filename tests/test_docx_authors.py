"""修订作者、批注作者一律改成固定值「Author」（幕僚长定：人名本身就是敏感信息）。

不管作者名命中没命中规则都改。覆盖：正文、页眉、脚注里的 w:ins / w:del / w:rPrChange，
styles.xml 里的格式修订，批注的 w:author / w:initials，people.xml 的 w15:author
和 w15:presenceInfo（userId 常是邮箱或账号）。界面和命令行两条路，
写出后用 python-docx 和 lxml（装了 pandoc 再加 pandoc --track-changes=all）读回。全部是编造数据。
"""

from __future__ import annotations

import io
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from doc_redactor.docx_redactor import redact_docx
from doc_redactor.engine import WordlistRule, build_engine
from doc_redactor.scan import scan_path, write_scans

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
# 一个命中规则（星辰海），两个不命中任何规则的编造名字。
HIT_NAME = "星辰海 审阅"
PLAIN_NAME = "Reviewer Zeta"
OTHER_NAME = "Editor Omega"
FIXED_AUTHOR = "Author"  # 幕僚长定的固定值
NAMES = (HIT_NAME, PLAIN_NAME, OTHER_NAME, "zeta@example.com", "RZ", "EO")


def _engine():
    return build_engine([], [WordlistRule(category="SPONSOR", terms=["星辰海"], case_sensitive=False)])


def _rewrite(path: Path, edit) -> None:
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    edit(files)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    path.write_bytes(buffer.getvalue())


def _add_rel(files, rels_name, rid, reltype, target):
    rel = f'<Relationship Id="{rid}" Type="{reltype}" Target="{target}"/>'.encode()
    files[rels_name] = files[rels_name].replace(b"</Relationships>", rel + b"</Relationships>")


def _add_type(files, part, ctype):
    item = f'<Override PartName="/{part}" ContentType="{ctype}"/>'.encode()
    files["[Content_Types].xml"] = files["[Content_Types].xml"].replace(b"</Types>", item + b"</Types>")


def _ins(author, text, rid):
    return (
        f'<w:ins w:id="{rid}" w:author="{author}" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:ins>'
    )


def _authors_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Body text.")
    document.sections[0].header.add_paragraph("Header text.")
    document.save(str(path))

    def edit(files):
        body = (
            f'<w:p>{_ins(PLAIN_NAME, "added ", 41)}'
            f'<w:del w:id="42" w:author="{HIT_NAME}" w:date="2026-01-01T00:00:00Z">'
            '<w:r><w:delText>gone</w:delText></w:r></w:del>'
            '<w:commentRangeStart w:id="0"/>'
            f'<w:r><w:rPr><w:b/><w:rPrChange w:id="43" w:author="{OTHER_NAME}" w:date="2026-01-01T00:00:00Z">'
            '<w:rPr/></w:rPrChange></w:rPr><w:t>bold</w:t></w:r>'
            '<w:commentRangeEnd w:id="0"/><w:r><w:commentReference w:id="0"/></w:r></w:p>'
            f'<w:p><w:pPr><w:pPrChange w:id="44" w:author="{PLAIN_NAME}" w:date="2026-01-01T00:00:00Z">'
            '<w:pPr/></w:pPrChange></w:pPr><w:r><w:t>moved</w:t></w:r></w:p>'
        )
        files["word/document.xml"] = files["word/document.xml"].replace(b"<w:sectPr", body.encode() + b"<w:sectPr", 1)
        header = next(n for n in files if n.startswith("word/header") and n.endswith(".xml"))
        files[header] = files[header].replace(b"</w:hdr>", f"<w:p>{_ins(OTHER_NAME, 'hdr ', 45)}</w:p></w:hdr>".encode())
        files["word/comments.xml"] = (
            f'<w:comments xmlns:w="{W}">'
            f'<w:comment w:id="0" w:author="{PLAIN_NAME}" w:initials="RZ" w:date="2026-01-01T00:00:00Z">'
            '<w:p><w:r><w:t>note one</w:t></w:r></w:p></w:comment>'
            f'<w:comment w:id="1" w:author="{OTHER_NAME}" w:initials="EO" w:date="2026-01-01T00:00:00Z">'
            '<w:p><w:r><w:t>note two</w:t></w:r></w:p></w:comment></w:comments>'
        ).encode()
        _add_type(files, "word/comments.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml")
        _add_rel(files, "word/_rels/document.xml.rels", "rIdC1", RT + "comments", "comments.xml")
        files["word/footnotes.xml"] = (
            f'<w:footnotes xmlns:w="{W}">'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
            f'<w:footnote w:id="1"><w:p>{_ins(HIT_NAME, "fn", 46)}</w:p></w:footnote></w:footnotes>'
        ).encode()
        _add_type(files, "word/footnotes.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml")
        _add_rel(files, "word/_rels/document.xml.rels", "rIdF1", RT + "footnotes", "footnotes.xml")
        files["word/people.xml"] = (
            f'<w15:people xmlns:w15="{W15}">'
            f'<w15:person w15:author="{PLAIN_NAME}"><w15:presenceInfo w15:providerId="Windows Live" '
            'w15:userId="zeta@example.com"/></w15:person>'
            f'<w15:person w15:author="{HIT_NAME}"/></w15:people>'
        ).encode()
        _add_type(files, "word/people.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.people+xml")
        _add_rel(files, "word/_rels/document.xml.rels", "rIdP1", "http://schemas.microsoft.com/office/2011/relationships/people", "people.xml")
        styles = files["word/styles.xml"]
        change = (
            f'<w:style w:type="character" w:customStyle="1" w:styleId="Zeta"><w:name w:val="Zeta"/>'
            f'<w:rPr><w:i/><w:rPrChange w:id="47" w:author="{OTHER_NAME}" w:date="2026-01-01T00:00:00Z"><w:rPr/></w:rPrChange></w:rPr></w:style>'
        ).encode()
        files["word/styles.xml"] = styles.replace(b"</w:styles>", change + b"</w:styles>")

    _rewrite(path, edit)


def _write_both_ways(source: Path, tmp_path: Path) -> list[Path]:
    engine = _engine()
    cli_out = tmp_path / "cli" / f"{source.stem}.redacted.docx"
    redact_docx(source, cli_out, engine)
    scan = scan_path(source, engine)
    assert scan.status == "ok", scan.error
    results = write_scans([scan], tmp_path / "gui", engine=engine)
    assert results[0].dest is not None, results[0].error
    return [cli_out, results[0].dest]


def _person_attributes(path: Path) -> list[tuple[str, str, str]]:
    """(部件, 属性名, 值)：所有 XML 部件里本地名是 author / initials / userId / providerId 的属性。"""
    found = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith(".xml"):
                continue
            root = etree.fromstring(archive.read(name))
            for element in root.iter():
                if not isinstance(element.tag, str):
                    continue
                for key, value in element.attrib.items():
                    local = key.rpartition("}")[2]
                    if local in ("author", "initials", "userId", "providerId"):
                        found.append((name, local, value))
    return found


def _all_text(path: Path) -> str:
    out = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                for element in etree.fromstring(archive.read(name)).iter():
                    out.append(element.text or "")
                    out.extend(element.attrib.values())
    return "\n".join(out)


def test_fixture_carries_every_author_name(tmp_path: Path):
    source = tmp_path / "authors.docx"
    _authors_docx(source)
    values = {value for _part, _local, value in _person_attributes(source)}
    assert {HIT_NAME, PLAIN_NAME, OTHER_NAME, "RZ", "EO", "zeta@example.com"} <= values
    assert len(Document(str(source)).comments) == 2


def test_every_revision_and_comment_author_becomes_author(tmp_path: Path):
    source = tmp_path / "authors.docx"
    _authors_docx(source)
    for output in _write_both_ways(source, tmp_path):
        found = _person_attributes(output)
        parts = {part for part, _local, _value in found}
        # 每处都还在（修订标记和批注本身不删），只是名字换了。
        assert {"word/document.xml", "word/comments.xml", "word/footnotes.xml", "word/styles.xml", "word/people.xml"} <= parts
        assert any(part.startswith("word/header") for part in parts)
        wrong = [(part, local) for part, local, value in found if local in ("author", "initials") and value != FIXED_AUTHOR]
        assert not wrong, wrong
        assert not [item for item in found if item[1] in ("userId", "providerId")]
        text = _all_text(output)
        for number, name in enumerate(NAMES, 1):
            assert name not in text, f"{output.parent.name} 读出了作者名 #{number}"
        document = Document(str(output))
        comments = list(document.comments)
        assert len(comments) == 2
        assert {(c.author, c.initials) for c in comments} == {(FIXED_AUTHOR, FIXED_AUTHOR)}


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="没装 pandoc")
def test_pandoc_track_changes_shows_only_author(tmp_path: Path):
    source = tmp_path / "authors.docx"
    _authors_docx(source)
    for output in _write_both_ways(source, tmp_path):
        shown = subprocess.run(
            ["pandoc", "--track-changes=all", "-t", "native", str(output)],
            capture_output=True, text=True, check=True,
        ).stdout
        for number, name in enumerate(NAMES[:3], 1):
            assert name not in shown, f"pandoc 读出了作者名 #{number}"
        assert f'"{FIXED_AUTHOR}"' in shown
