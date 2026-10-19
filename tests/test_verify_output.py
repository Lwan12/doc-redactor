"""Write then re-check: leftovers are reported, the output is deleted. Fictional data only."""

from pathlib import Path

import pymupdf
import pytest
from docx import Document

import doc_redactor.docx_redactor as docx_redactor
import doc_redactor.scan as scan_mod
from doc_redactor.cli import main
from doc_redactor.engine import RegexRule, WordlistRule, build_engine
from doc_redactor.scan import scan_path, write_scans
from helpers import RULES


def _engine():
    return build_engine(
        [RegexRule(name="PROJECT_ID", pattern=r"ABC-\d{4}-\d{3}", flags=0)],
        [WordlistRule(category="SPONSOR", terms=["Acme Pharma"], case_sensitive=True)],
    )


def _docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Project ABC-2026-001 by Acme Pharma")
    doc.add_paragraph("Second ABC-2026-002")
    doc.sections[0].header.paragraphs[0].text = "Header Acme Pharma"
    doc.save(str(path))


def _pdf(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Project ABC-2026-001 by Acme Pharma")
    page.insert_text((72, 100), "Second ABC-2026-002")
    doc.save(str(path))
    doc.close()


def _skip_first_paragraph_hit(monkeypatch):
    real = docx_redactor.apply_hits_to_paragraph
    state = {"n": 0}

    def leaky(paragraph, hits, replacement_for):
        state["n"] += 1
        if state["n"] == 1:
            return [hit.category for hit in hits]  # pretend, write nothing
        return real(paragraph, hits, replacement_for)

    monkeypatch.setattr(scan_mod, "apply_hits_to_paragraph", leaky)


@pytest.mark.parametrize("maker,name", [(_docx, "a.docx"), (_pdf, "a.pdf")])
def test_clean_write_passes(tmp_path: Path, maker, name):
    src = tmp_path / name
    maker(src)
    engine = _engine()
    scan = scan_path(src, engine)
    [res] = write_scans([scan], tmp_path / "out", engine=engine)
    assert res.error is None and res.dest is not None and res.dest.is_file()


def test_docx_leak_is_reported_and_deleted(tmp_path: Path, monkeypatch):
    src = tmp_path / "a.docx"
    _docx(src)
    engine = _engine()
    scan = scan_path(src, engine)
    _skip_first_paragraph_hit(monkeypatch)
    [res] = write_scans([scan], tmp_path / "out", engine=engine)
    assert res.error and "2 处写出后仍可见" in res.error
    assert res.dest is None
    assert not any((tmp_path / "out").glob("*.docx"))
    assert not any((tmp_path / "out").glob("*.report.json"))


def test_pdf_leak_is_reported_and_deleted(tmp_path: Path, monkeypatch):
    src = tmp_path / "a.pdf"
    _pdf(src)
    engine = _engine()
    scan = scan_path(src, engine)
    for hit in scan.hits:
        if hit.text == "ABC-2026-002":
            hit.rects = []  # writer has nowhere to redact
    [res] = write_scans([scan], tmp_path / "out", engine=engine)
    assert res.error and "1 处写出后仍可见" in res.error
    assert not any((tmp_path / "out").glob("*.pdf"))


def test_unchecked_hit_is_not_a_leftover(tmp_path: Path):
    src = tmp_path / "a.docx"
    _docx(src)
    engine = _engine()
    scan = scan_path(src, engine)
    scan.hits[0].enabled = False
    [res] = write_scans([scan], tmp_path / "out", engine=engine)
    assert res.error is None and res.dest.is_file()


def test_without_engine_uses_hit_texts(tmp_path: Path, monkeypatch):
    src = tmp_path / "a.docx"
    _docx(src)
    scan = scan_path(src, _engine())
    _skip_first_paragraph_hit(monkeypatch)
    [res] = write_scans([scan], tmp_path / "out")
    assert res.error and "写出后仍可见" in res.error


def test_cli_leak_exits_nonzero_and_deletes(tmp_path: Path, monkeypatch, capsys):
    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_paragraph("Project ABC-2026-001")
    doc.save(str(src))
    monkeypatch.setattr(
        docx_redactor, "_redact_paragraph", lambda p, e: [h.category for h in e.find_hits(p.text)]
    )
    outdir = tmp_path / "out"
    code = main([str(src), "-r", str(RULES), "-o", str(outdir)])
    assert code == 1
    assert "写出后仍可见" in capsys.readouterr().err
    assert not (outdir / "in.redacted.docx").exists()
    assert not (outdir / "in.report.json").exists()


def _short_code_pdf(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Sponsor SP and Fictizumab", fontsize=11)
    doc.save(str(path))
    doc.close()


def _short_code_engine():
    # "SP" is narrow, so the writer falls back to the short code [SP] and adds
    # a legend page. Neither the code nor the legend is original text.
    return build_engine([], [WordlistRule(category="SPONSOR", terms=["SP"], case_sensitive=True)])


def test_short_code_and_legend_are_not_leftovers(tmp_path: Path):
    src = tmp_path / "s.pdf"
    _short_code_pdf(src)
    engine = _short_code_engine()
    scan = scan_path(src, engine)
    [res] = write_scans([scan], tmp_path / "out", engine=engine)
    assert res.error is None and res.dest is not None and res.dest.is_file()
    assert res.short_count == 1


def test_cli_short_code_is_not_a_leftover(tmp_path: Path):
    from doc_redactor.verify_output import count_leftovers
    from doc_redactor.pdf_redactor import redact_pdf

    src = tmp_path / "s.pdf"
    _short_code_pdf(src)
    engine = _short_code_engine()
    out = tmp_path / "s.redacted.pdf"
    counts = redact_pdf(src, out, engine)
    assert counts.short_count == 1
    labels = {engine.label(category) for category in counts}
    assert count_leftovers(out, engine=engine, labels=labels, short_codes=True) == 0
    # A real leftover is still found on a page that also has short codes.
    assert count_leftovers(src, engine=engine, labels=labels, short_codes=True) == 1


def test_cli_leak_blocks_only_that_file(tmp_path: Path, monkeypatch, capsys):
    leak = tmp_path / "leak.docx"
    doc = Document()
    doc.add_paragraph("Project ABC-2026-001")
    doc.save(str(leak))
    ok = tmp_path / "ok.pdf"
    _pdf(ok)
    monkeypatch.setattr(
        docx_redactor, "_redact_paragraph", lambda p, e: [h.category for h in e.find_hits(p.text)]
    )
    outdir = tmp_path / "out"
    code = main([str(leak), str(ok), "-r", str(RULES), "-o", str(outdir)])
    assert code == 1
    err = capsys.readouterr().err
    assert "leak.docx: 1 处写出后仍可见" in err
    assert "错误" not in err
    assert not (outdir / "leak.redacted.docx").exists()
    assert (outdir / "ok.redacted.pdf").is_file()


def test_unlocated_file_is_blocked_once_not_rechecked(tmp_path: Path, monkeypatch):
    src = tmp_path / "a.pdf"
    _pdf(src)
    engine = _engine()
    scan = scan_path(src, engine)
    scan.unlocated_count = 2
    calls = []
    monkeypatch.setattr(scan_mod, "count_leftovers", lambda *a, **k: calls.append(1) or 0)
    [res] = write_scans([scan], tmp_path / "out", engine=engine)
    assert res.error and "2 处无法定位" in res.error
    assert "写出后仍可见" not in res.error
    assert calls == []
    assert not any((tmp_path / "out").glob("*.pdf"))


def test_cli_unlocated_reports_only_unlocated(tmp_path: Path, monkeypatch, capsys):
    import doc_redactor.pdf_redactor as pdf_redactor

    src = tmp_path / "in.pdf"
    _pdf(src)
    monkeypatch.setattr(pdf_redactor, "locate_text_hits", lambda page, found, tp: [[] for _ in found])
    outdir = tmp_path / "out"
    code = main([str(src), "-r", str(RULES), "-o", str(outdir)])
    err = capsys.readouterr().err
    assert code == 1
    assert "无法定位" in err and "写出后仍可见" not in err
    assert not (outdir / "in.redacted.pdf").exists()
